#!/usr/bin/env python3
"""Update the V3 thesis with official post-Phase2 A0R/A1-DINO results."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


BODY_STYLES = {"Normal", "Normal (Web)"}


def find_exact(doc: Document, text: str):
    matches = [p for p in doc.paragraphs if p.text.strip() == text]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph {text!r}, found {len(matches)}")
    return matches[0]


def replace_exact(doc: Document, old: str, new: str) -> None:
    paragraph = find_exact(doc, old)
    paragraph.text = new
    if paragraph.style.name in BODY_STYLES:
        format_body(paragraph)


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def set_run_font(run, size: float, bold: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")


def format_body(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Cm(1.27)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    for run in paragraph.runs:
        set_run_font(run, 13)


def format_heading(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.keep_together = True


def insert_body(target, text: str):
    paragraph = target.insert_paragraph_before(text, style="Normal")
    format_body(paragraph)
    return paragraph


def insert_heading(target, text: str, style: str):
    paragraph = target.insert_paragraph_before(text, style=style)
    format_heading(paragraph)
    return paragraph


def insert_caption(target, text: str, *, page_break_before: bool = False):
    paragraph = target.insert_paragraph_before(text, style="Table Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.page_break_before = page_break_before
    for run in paragraph.runs:
        set_run_font(run, 12)
    return paragraph


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def set_no_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def shade_header(table) -> None:
    for cell in table.rows[0].cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        shade = tc_pr.find(qn("w:shd"))
        if shade is None:
            shade = OxmlElement("w:shd")
            tc_pr.append(shade)
        shade.set(qn("w:fill"), "D9EAF7")


def format_cell(cell, text: str, *, bold: bool, size: float, align) -> None:
    cell.text = text
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.alignment = align
        paragraph.paragraph_format.first_line_indent = None
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            set_run_font(run, size, bold)


def make_table(
    doc: Document,
    target,
    rows: list[list[str]],
    *,
    widths_cm: list[float],
    font_size: float = 10,
    bold_cells: set[tuple[int, int]] | None = None,
):
    bold_cells = bold_cells or set()
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = doc.tables[13].style
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, values in enumerate(rows):
        for j, value in enumerate(values):
            align = WD_ALIGN_PARAGRAPH.CENTER if j < 2 or i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            format_cell(
                table.cell(i, j),
                value,
                bold=i == 0 or (i, j) in bold_cells,
                size=font_size,
                align=align,
            )
            table.cell(i, j).width = Cm(widths_cm[j])
        set_no_split(table.rows[i])
    set_repeat_header(table.rows[0])
    shade_header(table)
    target._p.addprevious(table._element)
    return table


def append_table_row(table, values: list[str], *, size: float = 9.5) -> None:
    row = table.add_row()
    for j, value in enumerate(values):
        format_cell(row.cells[j], value, bold=False, size=size, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_no_split(row)


def update_intro(doc: Document) -> None:
    replacements = {
        'Khóa luận thực hiện đề tài "Thiết kế và phát triển mô hình học sâu phát hiện người và phương tiện trong điều kiện thời tiết bất lợi phục vụ giám sát giao thông thông minh". Nghiên cứu xây dựng quy trình học chuyển giao nhiều giai đoạn, so sánh YOLOv8n, YOLO11n, Faster R-CNN và RT-DETR-L, sau đó lựa chọn RT-DETR-L bằng XWOD validation để tiếp tục fine-tuning đa miền ở Phase 2. RT-DETR-L Phase 2 là mô hình cuối cùng trong phạm vi thực nghiệm của khóa luận.':
            'Khóa luận thực hiện đề tài "Thiết kế và phát triển mô hình học sâu phát hiện người và phương tiện trong điều kiện thời tiết bất lợi phục vụ giám sát giao thông thông minh". Nghiên cứu xây dựng quy trình học chuyển giao nhiều giai đoạn, so sánh YOLOv8n, YOLO11n, Faster R-CNN và RT-DETR-L, sau đó lựa chọn RT-DETR-L bằng XWOD validation để fine-tuning đa miền ở Phase 2. Từ cùng checkpoint Phase 2, hai ablation A0R và A1-DINO được huấn luyện với cùng ngân sách 5.000 ảnh BDD100K train nhằm đánh giá tác động của chiến lược lựa chọn dữ liệu.',
        "Xuất phát từ thực tiễn đó, khóa luận xây dựng một quy trình thực nghiệm có kiểm soát cho bài toán phát hiện người và phương tiện trong điều kiện thời tiết bất lợi. Bốn kiến trúc YOLOv8n, YOLO11n, Faster R-CNN và RT-DETR-L được so sánh trên cùng không gian sáu lớp; mô hình được chọn bằng XWOD validation trước khi bước vào Phase 2.":
            "Xuất phát từ thực tiễn đó, khóa luận xây dựng một quy trình thực nghiệm có kiểm soát cho bài toán phát hiện người và phương tiện trong điều kiện thời tiết bất lợi. Bốn kiến trúc YOLOv8n, YOLO11n, Faster R-CNN và RT-DETR-L được so sánh trên cùng không gian sáu lớp; RT-DETR-L được chọn bằng XWOD validation trước Phase 2 và tiếp tục được dùng làm nền cho hai ablation lựa chọn dữ liệu sau Phase 2.",
        "Quy trình nghiên cứu gồm COCO pretrained → Stage 1 trên BDD100K → Stage 2 trên XWOD → lựa chọn bằng XWOD validation → Phase 2 multi-domain fine-tuning trên XWOD, ACDC và BDD100K. Sau khi cấu hình được cố định, mô hình được đánh giá trên XWOD, BDD, DAWN và ACDC.":
            "Quy trình nghiên cứu gồm COCO pretrained → Stage 1 trên BDD100K → Stage 2 trên XWOD → lựa chọn bằng XWOD validation → Phase 2 multi-domain fine-tuning trên XWOD, ACDC và BDD100K → A0R/A1-DINO post-Phase2 data-selection ablations. Phase 2 và hai ablation được đánh giá frozen trên XWOD, ACDC, DAWN và BDD.",
        "Lựa chọn mô hình bằng XWOD validation và tiếp tục tối ưu mô hình được chọn bằng Phase 2 multi-domain fine-tuning trên XWOD, ACDC và BDD100K.":
            "Lựa chọn RT-DETR-L bằng XWOD validation, thực hiện Phase 2 multi-domain fine-tuning và đánh giá hai ablation lựa chọn dữ liệu sau Phase 2 gồm A0R và A1-DINO.",
        "Phân tích mô hình cuối theo bộ dữ liệu, lớp đối tượng, điều kiện thời tiết, đánh đổi độ chính xác–tốc độ và khả năng duy trì tri thức giữa các miền.":
            "Phân tích Phase 2 và các biến thể sau Phase 2 theo bộ dữ liệu, lớp đối tượng, điều kiện thời tiết, đánh đổi độ chính xác–tốc độ và khả năng duy trì tri thức giữa các miền.",
        "Về mô hình, đề tài khảo sát YOLOv8n, YOLO11n, Faster R-CNN và RT-DETR-L, đại diện cho các hướng một giai đoạn, hai giai đoạn và Transformer detector. RT-DETR-L được lựa chọn bằng XWOD validation và RT-DETR-L Phase 2 được xác định là mô hình cuối cùng của đề tài.":
            "Về mô hình, đề tài khảo sát YOLOv8n, YOLO11n, Faster R-CNN và RT-DETR-L, đại diện cho các hướng một giai đoạn, hai giai đoạn và Transformer detector. RT-DETR-L được lựa chọn bằng XWOD validation; Phase 2 là baseline đa miền, còn A0R và A1-DINO là các ablation tiếp tục huấn luyện từ cùng checkpoint Phase 2.",
        "Về huấn luyện, các mô hình bắt đầu từ trọng số pretrained, tinh chỉnh trên BDD100K ở Stage 1 và tiếp tục trên XWOD ở Stage 2. XWOD validation được dùng để theo dõi hội tụ, chọn best checkpoint và làm căn cứ lựa chọn RT-DETR-L cho Phase 2; XWOD, BDD, DAWN và ACDC test chỉ được dùng cho đánh giá cuối sau khi cấu hình đã cố định.":
            "Về huấn luyện, các mô hình bắt đầu từ trọng số pretrained, tinh chỉnh trên BDD100K ở Stage 1 và tiếp tục trên XWOD ở Stage 2. XWOD validation được dùng để theo dõi hội tụ, chọn best checkpoint và lựa chọn RT-DETR-L cho Phase 2. Từ cùng checkpoint Phase 2, A0R và A1-DINO bổ sung 5.000 ảnh từ phần BDD100K train chưa sử dụng theo hai chiến lược lựa chọn khác nhau; các tập test chỉ được dùng cho frozen evaluation sau khi cấu hình đã cố định.",
        "Cuối cùng, các mô hình được đánh giá bằng cùng hệ thống độ đo, gồm Precision, Recall, mAP50, mAP50-95, tốc độ suy luận và số lượng tham số. Kết quả không chỉ được phân tích ở mức tổng thể mà còn được xem xét theo từng lớp đối tượng và từng điều kiện thời tiết. Cách tiếp cận này giúp làm rõ điểm mạnh, điểm hạn chế của từng mô hình, đồng thời cung cấp cơ sở thực nghiệm cho việc lựa chọn mô hình phù hợp với bài toán giám sát giao thông trong điều kiện thời tiết bất lợi.":
            "Cuối cùng, các mô hình được đánh giá bằng Precision, Recall, mAP50 và mAP50-95 trên các split đã cố định; tốc độ suy luận và số lượng tham số được dùng khi điều kiện đo cho phép. Với các ablation sau Phase 2, khóa luận so sánh kết quả tổng thể trên bốn tập kiểm thử và kết quả theo lớp trên XWOD để tách lợi ích của việc bổ sung dữ liệu khỏi lợi ích của chiến lược truy hồi.",
        "Chương 3 trình bày yêu cầu bài toán, kiến trúc hệ thống, chuẩn hóa dữ liệu, bốn mô hình benchmark, chiến lược huấn luyện nhiều giai đoạn và giao thức thực nghiệm.":
            "Chương 3 trình bày yêu cầu bài toán, kiến trúc hệ thống, chuẩn hóa dữ liệu, bốn mô hình benchmark, chiến lược huấn luyện nhiều giai đoạn, hai ablation lựa chọn dữ liệu sau Phase 2 và giao thức thực nghiệm.",
        "Chương 4 trình bày môi trường, cấu hình huấn luyện, benchmark bốn mô hình, Phase 2 của RT-DETR-L, phân tích kết quả và các hướng phát triển bám theo hạn chế quan sát được.":
            "Chương 4 trình bày môi trường, cấu hình huấn luyện, benchmark bốn mô hình, Phase 2 của RT-DETR-L, kết quả A0R/A1-DINO, phân tích đối chứng và các hướng phát triển bám theo hạn chế quan sát được.",
    }
    for old, new in replacements.items():
        replace_exact(doc, old, new)


def update_chapter3(doc: Document) -> None:
    replace_exact(
        doc,
        "Kiến trúc hệ thống gồm bốn khối liên kết: dữ liệu và tiền xử lý; các mô hình benchmark; huấn luyện–lựa chọn mô hình; Phase 2 và đánh giá cuối. Dữ liệu được chuẩn hóa về cùng sáu lớp trước khi đi qua Stage 1 và Stage 2. XWOD validation được dùng để lựa chọn RT-DETR-L; cấu hình này tiếp tục fine-tuning đa miền và chỉ được đưa sang frozen test sau khi đã cố định.",
        "Kiến trúc hệ thống gồm các khối liên kết: dữ liệu và tiền xử lý; các mô hình benchmark; huấn luyện–lựa chọn mô hình; Phase 2; lựa chọn dữ liệu sau Phase 2; và đánh giá frozen. Dữ liệu được chuẩn hóa về cùng sáu lớp trước Stage 1 và Stage 2. XWOD validation được dùng để lựa chọn RT-DETR-L; checkpoint Phase 2 sau đó là điểm khởi tạo chung của A0R và A1-DINO.",
    )
    replace_exact(
        doc,
        "Khối dữ liệu đảm nhiệm đọc ảnh, chuyển đổi nhãn và tổ chức train, validation, test. Khối mô hình gồm YOLOv8n, YOLO11n, Faster R-CNN và RT-DETR-L. Khối huấn luyện điều phối COCO pretrained → BDD100K → XWOD và tách rõ validation khỏi test. Khối cuối sử dụng RT-DETR-L Phase 2 để đánh giá frozen trên XWOD, BDD, DAWN và ACDC.",
        "Khối dữ liệu đảm nhiệm đọc ảnh, chuyển đổi nhãn và tổ chức train, validation, test. Khối mô hình gồm YOLOv8n, YOLO11n, Faster R-CNN và RT-DETR-L. Khối huấn luyện điều phối COCO pretrained → BDD100K → XWOD → Phase 2. Sau Phase 2, hai nhánh A0R và A1-DINO sử dụng cùng checkpoint, cùng ngân sách 5.000 ảnh BDD100K train và cùng training budget trước khi Phase 2/A0R/A1-DINO được đánh giá frozen trên XWOD, ACDC, DAWN và BDD.",
    )

    old_protocol = find_exact(doc, "3.6. Giao thức thực nghiệm")
    insert_heading(old_protocol, "3.6. Thí nghiệm lựa chọn dữ liệu sau Phase 2", "Heading 2")
    insert_body(
        old_protocol,
        "Sau khi Phase 2 hoàn tất, checkpoint tốt nhất của RT-DETR-L Phase 2 được dùng làm điểm khởi tạo chung cho hai ablation lựa chọn dữ liệu. Cả hai thí nghiệm chỉ khai thác phần BDD100K train có nhãn chưa được sử dụng làm candidate pool và không sử dụng BDD validation/test để chọn dữ liệu.",
    )
    insert_heading(old_protocol, "3.6.1. A0R – Random Rare-Class Sampling", "Heading 3")
    insert_body(
        old_protocol,
        "A0R là control experiment nhằm kiểm tra liệu lợi ích có chủ yếu đến từ việc bổ sung thêm dữ liệu liên quan đến lớp hiếm hay không. Candidate pool được lọc từ phần BDD100K train chưa sử dụng để giữ các ảnh liên quan đến rare classes; từ đó 5.000 ảnh được chọn ngẫu nhiên. Mô hình tiếp tục huấn luyện từ Phase 2 best checkpoint với epochs = 20, patience = 8, lr0 = 1×10⁻⁵, batch = 16 và seed = 42.",
    )
    insert_heading(old_protocol, "3.6.2. A1-DINO – Failure-Driven DINOv2 Retrieval", "Heading 3")
    insert_body(
        old_protocol,
        "A1-DINO kiểm tra liệu lựa chọn dữ liệu dựa trên failure/hard examples của Phase 2 có tạo lợi ích vượt quá random sampling hay không. Các mẫu khó liên quan đến bicycle, motorcycle và bus từ XWOD train và ACDC train được dùng làm query. Image embedding được trích xuất bằng DINOv2 [16], chuẩn hóa và so sánh bằng cosine similarity với cùng unused labeled BDD100K train candidate pool; global Top-5.000 ảnh được chọn để tiếp tục huấn luyện từ checkpoint Phase 2.",
    )
    insert_body(
        old_protocol,
        "Candidate pool đã có nhãn, vì vậy A1-DINO được mô tả là failure-driven similarity retrieval, không phải active learning theo nghĩa nghiêm ngặt. Validation và test không tham gia candidate selection; các tập test chỉ được sử dụng sau khi cấu hình đánh giá đã được cố định.",
    )
    insert_heading(old_protocol, "3.6.3. Giao thức đối chứng", "Heading 3")
    insert_body(
        old_protocol,
        "A0R và A1-DINO sử dụng cùng Phase 2 best checkpoint, cùng candidate pool BDD100K train, cùng 5.000 ảnh bổ sung, cùng cấu hình epochs = 20, patience = 8, lr0 = 1×10⁻⁵, batch = 16, seed = 42 và cùng giao thức frozen evaluation. Thiết kế này dùng A0R làm đối chứng để tách ảnh hưởng của lượng dữ liệu bổ sung khỏi ảnh hưởng riêng của chiến lược truy hồi DINOv2.",
    )
    insert_caption(old_protocol, "Bảng 3.6. Thiết kế đối chứng A0R và A1-DINO")
    make_table(
        doc,
        old_protocol,
        [
            ["Thuộc tính", "A0R", "A1-DINO"],
            ["Khởi tạo", "RT-DETR-L Phase 2 best checkpoint", "RT-DETR-L Phase 2 best checkpoint"],
            ["Candidate pool", "Unused labeled BDD100K train; rare-class candidates", "Cùng candidate pool với A0R"],
            ["Lựa chọn", "Random 5.000 ảnh", "DINOv2 cosine similarity; global Top-5.000"],
            ["Training budget", "20 epoch; patience 8; lr0 1×10⁻⁵; batch 16; seed 42", "Giống A0R"],
            ["Đánh giá", "XWOD/ACDC/DAWN/BDD frozen test", "Cùng giao thức với A0R"],
        ],
        widths_cm=[3.0, 6.1, 6.1],
        font_size=9.5,
    )

    old_protocol.text = "3.7. Giao thức thực nghiệm"
    format_heading(old_protocol)
    find_exact(doc, "3.6.1. Phân tách train, validation và test").text = "3.7.1. Phân tách train, validation và test"
    find_exact(doc, "3.6.2. Độ đo và nguyên tắc diễn giải").text = "3.7.2. Độ đo và nguyên tắc diễn giải"
    find_exact(doc, "3.6.3. Khả năng tái lập và tính công bằng").text = "3.7.3. Khả năng tái lập và tính công bằng"
    for title in (
        "3.7.1. Phân tách train, validation và test",
        "3.7.2. Độ đo và nguyên tắc diễn giải",
        "3.7.3. Khả năng tái lập và tính công bằng",
    ):
        format_heading(find_exact(doc, title))
    find_exact(doc, "Bảng 3.6. Ma trận các giai đoạn thực nghiệm").text = "Bảng 3.7. Ma trận các giai đoạn và ablation thực nghiệm"
    find_exact(doc, "Bảng 3.7. Giao thức đánh giá theo từng tập dữ liệu").text = "Bảng 3.8. Giao thức đánh giá theo từng tập dữ liệu"
    replace_exact(
        doc,
        "Train được dùng để tối ưu tham số; validation phục vụ early stopping, chọn best checkpoint và lựa chọn RT-DETR-L; test chỉ được dùng cho đánh giá frozen sau khi cấu hình đã cố định. Bảng 3.6 tóm tắt vai trò của ba giai đoạn thực nghiệm chính.",
        "Train được dùng để tối ưu tham số; validation phục vụ early stopping và lựa chọn checkpoint ở các giai đoạn tương ứng; test chỉ được dùng cho đánh giá frozen sau khi cấu hình đã cố định. Bảng 3.7 tóm tắt vai trò của các giai đoạn chính và hai ablation sau Phase 2.",
    )
    replace_exact(
        doc,
        "XWOD val được dùng chọn checkpoint, còn XWOD test đo hiệu năng trên cùng miền dữ liệu thời tiết bất lợi. DAWN test đo khả năng chuyển sang một nguồn dữ liệu không tham gia huấn luyện. ACDC test được báo cáo ở hai bối cảnh: trước Phase 2 là đánh giá chéo dữ liệu; sau Phase 2 là đánh giá trên split được giữ lại của cùng bộ dữ liệu. Các vai trò này được ghi ngay trong bảng kết quả để người đọc không đánh đồng ba loại phép đo.",
        "XWOD val được dùng chọn checkpoint ở quy trình chính, còn XWOD test đo hiệu năng trên cùng miền dữ liệu thời tiết bất lợi. DAWN test đo khả năng chuyển sang một nguồn dữ liệu không tham gia huấn luyện. ACDC test trước Phase 2 là đánh giá chéo dữ liệu; sau Phase 2 là held-out same-source evaluation. Bảng 3.8 ghi rõ vai trò từng split; không tập test nào tham gia model selection hoặc candidate selection.",
    )
    replace_exact(
        doc,
        "Các chỉ số tổng thể được tính trên cùng sáu lớp, nhưng kết quả theo lớp vẫn phải đi kèm số lượng đối tượng. DAWN chỉ có sáu instance bicycle nên một giá trị AP cao hoặc thấp của lớp này chưa đủ để kết luận mô hình ổn định. Tương tự, cải thiện mAP trung bình không được xem là thành công nếu recall giảm rõ hoặc các lớp hiếm tiếp tục suy giảm. Kết quả âm được giữ lại trong báo cáo vì chúng xác định giới hạn của phương án đã thử và tránh tạo cảm giác mọi thay đổi đều dẫn đến cải thiện.",
        "Các chỉ số tổng thể được tính trên cùng sáu lớp, nhưng kết quả theo lớp vẫn phải được xem xét cùng phân bố đối tượng. DAWN chỉ có sáu instance bicycle nên một giá trị AP cao hoặc thấp của lớp này chưa đủ để kết luận mô hình ổn định. Khi diễn giải kết quả, khóa luận xem xét đồng thời mAP, Recall và hiệu năng theo lớp nhằm tránh kết luận quá mức chỉ từ chỉ số trung bình.",
    )
    replace_exact(
        doc,
        "Từ thiết kế trên, Chương 4 trình bày môi trường và cấu hình, benchmark bốn mô hình, Phase 2 của RT-DETR-L, frozen evaluation, phân tích theo lớp và thời tiết, cùng các hạn chế và hướng phát triển.",
        "Từ thiết kế trên, Chương 4 trình bày môi trường và cấu hình, benchmark bốn mô hình, Phase 2 của RT-DETR-L, frozen evaluation của Phase 2/A0R/A1-DINO, phân tích theo lớp và thời tiết, cùng các hạn chế và hướng phát triển.",
    )

    matrix = doc.tables[10] if doc.tables[9].rows[0].cells[0].text != "Giai đoạn" else doc.tables[9]
    append_table_row(matrix, ["A0R", "Phase 2 data + 5.000 ảnh BDD train hiếm (random)", "Cấu hình đối chứng cố định", "XWOD/ACDC/DAWN/BDD frozen test"])
    append_table_row(matrix, ["A1-DINO", "Phase 2 data + 5.000 ảnh BDD train truy hồi", "Cấu hình đối chứng cố định", "XWOD/ACDC/DAWN/BDD frozen test"])


def update_phase2_tables_and_text(doc: Document) -> None:
    table42 = next(t for t in doc.tables if t.rows[0].cells[0].text == "Mô hình" and t.rows[0].cells[1].text == "Giai đoạn" and len(t.columns) == 5)
    append_table_row(table42, ["RT-DETR-L", "A0R", "20", "patience=8", "Phase 2 ckpt; random 5.000; lr0=10⁻⁵; batch=16"])
    append_table_row(table42, ["RT-DETR-L", "A1-DINO", "20", "patience=8", "Phase 2 ckpt; DINOv2 Top-5.000; lr0=10⁻⁵; batch=16"])

    phase2_table = next(t for t in doc.tables if t.rows[0].cells[0].text == "Tập" and len(t.columns) == 7 and t.rows[1].cells[0].text == "XWOD test")
    official = [
        ["XWOD test", "0,8063", "0,7448", "0,7786", "0,5096"],
        ["BDD test", "0,6975", "0,5582", "0,5938", "0,3535"],
        ["DAWN test", "0,8351", "0,7659", "0,7954", "0,5147"],
        ["ACDC test", "0,6005", "0,3840", "0,4068", "0,2470"],
    ]
    for i, row in enumerate(official, 1):
        for j, value in enumerate(row):
            format_cell(phase2_table.cell(i, j), value, bold=False, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)

    replace_exact(
        doc,
        "Mô hình cuối cùng đạt P/R/mAP50/mAP50-95 lần lượt 0,806/0,745/0,779/0,510 trên XWOD test; 0,698/0,558/0,594/0,354 trên BDD test; 0,835/0,766/0,795/0,515 trên DAWN test; và 0,600/0,384/0,407/0,247 trên ACDC test.",
        "RT-DETR-L Phase 2 đạt P/R/mAP50/mAP50-95 lần lượt 0,8063/0,7448/0,7786/0,5096 trên XWOD test; 0,6975/0,5582/0,5938/0,3535 trên BDD test; 0,8351/0,7659/0,7954/0,5147 trên DAWN test; và 0,6005/0,3840/0,4068/0,2470 trên ACDC test.",
    )


def insert_chapter4_ablation(doc: Document) -> None:
    old_analysis = find_exact(doc, "4.4. Phân tích kết quả")
    insert_heading(old_analysis, "4.4. Đánh giá chiến lược lựa chọn dữ liệu sau Phase 2", "Heading 2")
    insert_heading(old_analysis, "4.4.1. Thiết lập thí nghiệm", "Heading 3")
    insert_body(
        old_analysis,
        "Phase 2 là baseline multi-domain và là checkpoint khởi tạo chung cho A0R và A1-DINO. Hai ablation sử dụng cùng candidate pool từ phần BDD100K train chưa dùng, cùng ngân sách 5.000 ảnh, cùng cấu hình 20 epoch, patience 8, lr0 = 1×10⁻⁵, batch 16, seed 42 và cùng 12 frozen evaluations trên XWOD, ACDC, DAWN và BDD. Sự khác biệt chính nằm ở chiến lược chọn ảnh: random rare-class sampling đối với A0R và global image-level DINOv2 similarity retrieval đối với A1-DINO.",
    )
    insert_heading(old_analysis, "4.4.2. Kết quả tổng thể trên bốn tập kiểm thử", "Heading 3")
    insert_caption(old_analysis, "Bảng 4.7. So sánh mAP50-95 của Phase 2, A0R và A1-DINO trên các tập kiểm thử")
    make_table(
        doc,
        old_analysis,
        [
            ["Mô hình", "XWOD", "ACDC", "DAWN", "BDD", "Trung bình"],
            ["Phase 2", "0,5096", "0,2470", "0,5147", "0,3535", "0,4062"],
            ["A0R", "0,5096", "0,2542", "0,5160", "0,3596", "0,4099"],
            ["A1-DINO", "0,5004", "0,2424", "0,5253", "0,3542", "0,4056"],
        ],
        widths_cm=[3.0, 2.3, 2.3, 2.3, 2.3, 2.6],
        font_size=10,
        bold_cells={(1, 1), (2, 1), (2, 2), (3, 3), (2, 4), (2, 5)},
    )
    insert_body(
        old_analysis,
        "A0R đạt mAP50-95 trung bình cao nhất là 0,4099, tăng 0,0037 so với 0,4062 của Phase 2; A1-DINO đạt 0,4056, giảm 0,0006. So với Phase 2, A0R giữ nguyên XWOD ở 0,5096, tăng ACDC 0,0072, tăng DAWN 0,0013 và tăng BDD 0,0061. A1-DINO tăng DAWN 0,0106 nhưng giảm XWOD 0,0092 và ACDC 0,0046; BDD gần như giữ nguyên với mức tăng 0,0007.",
    )
    insert_caption(old_analysis, "Bảng 4.8. Precision, Recall, mAP50 và mAP50-95 của các biến thể sau Phase 2", page_break_before=True)
    make_table(
        doc,
        old_analysis,
        [
            ["Mô hình", "Tập", "Precision", "Recall", "mAP50", "mAP50-95"],
            ["Phase 2", "XWOD", "0,8063", "0,7448", "0,7786", "0,5096"],
            ["Phase 2", "ACDC", "0,6005", "0,3840", "0,4068", "0,2470"],
            ["Phase 2", "DAWN", "0,8351", "0,7659", "0,7954", "0,5147"],
            ["Phase 2", "BDD", "0,6975", "0,5582", "0,5938", "0,3535"],
            ["A0R", "XWOD", "0,8253", "0,7387", "0,7815", "0,5096"],
            ["A0R", "ACDC", "0,5723", "0,4012", "0,4188", "0,2542"],
            ["A0R", "DAWN", "0,8239", "0,7699", "0,7932", "0,5160"],
            ["A0R", "BDD", "0,6976", "0,5651", "0,5990", "0,3596"],
            ["A1-DINO", "XWOD", "0,8041", "0,7331", "0,7683", "0,5004"],
            ["A1-DINO", "ACDC", "0,6021", "0,3871", "0,4045", "0,2424"],
            ["A1-DINO", "DAWN", "0,8290", "0,7589", "0,8173", "0,5253"],
            ["A1-DINO", "BDD", "0,7007", "0,5626", "0,5935", "0,3542"],
        ],
        widths_cm=[2.5, 2.4, 2.5, 2.5, 2.4, 2.6],
        font_size=9.5,
    )
    insert_body(
        old_analysis,
        "A0R cải thiện Recall trên ACDC từ 0,3840 lên 0,4012 và trên BDD từ 0,5582 lên 0,5651. Tuy nhiên, Precision ACDC giảm từ 0,6005 xuống 0,5723, nên kết quả không cho thấy A0R cải thiện mọi metric. Với A1-DINO, DAWN mAP50 tăng từ 0,7954 lên 0,8173 và mAP50-95 tăng từ 0,5147 lên 0,5253, trong khi các thay đổi trên những miền còn lại không tạo thành một xu hướng cải thiện nhất quán.",
    )
    insert_heading(old_analysis, "4.4.3. Phân tích rare classes trên XWOD", "Heading 3")
    insert_caption(old_analysis, "Bảng 4.9. mAP50-95 theo lớp của Phase 2, A0R và A1-DINO trên XWOD test")
    make_table(
        doc,
        old_analysis,
        [
            ["Lớp", "Phase 2", "A0R", "A1-DINO"],
            ["person", "0,4560", "0,4532", "0,4484"],
            ["bicycle", "0,6520", "0,6582", "0,6377"],
            ["car", "0,5745", "0,5762", "0,5694"],
            ["motorcycle", "0,3837", "0,3845", "0,3679"],
            ["bus", "0,3959", "0,3917", "0,3883"],
            ["truck", "0,5955", "0,5936", "0,5904"],
        ],
        widths_cm=[4.0, 3.7, 3.7, 3.7],
        font_size=10,
    )
    insert_body(
        old_analysis,
        "Trên XWOD test, A0R tăng bicycle 0,0062, car 0,0017 và motorcycle 0,0008 so với Phase 2, nhưng giảm person 0,0028, bus 0,0042 và truck 0,0019. A1-DINO giảm mAP50-95 ở cả sáu lớp; bicycle giảm 0,0143 và motorcycle giảm 0,0158. Vì vậy, thiết kế global image-level retrieval hiện tại chưa cung cấp bằng chứng rằng DINOv2 similarity retrieval cải thiện các rare classes trên miền đích XWOD.",
    )
    insert_heading(old_analysis, "4.4.4. Thảo luận", "Heading 3")
    insert_body(
        old_analysis,
        "Do A0R và A1-DINO có cùng checkpoint khởi tạo, ngân sách dữ liệu, candidate pool, training budget và evaluation protocol, A0R là control trực tiếp cho chiến lược retrieval. Kết quả cho thấy trong thiết lập này, việc bổ sung dữ liệu liên quan đến lớp hiếm mang lại lợi ích tổng thể rõ hơn so với global image-level cosine similarity retrieval. Kết luận này chỉ áp dụng trong phạm vi thiết kế của khóa luận và không hàm ý random sampling luôn vượt learned retrieval.",
    )
    insert_body(
        old_analysis,
        "A1-DINO có dấu hiệu hỗ trợ transfer sang DAWN nhưng không cải thiện ổn định trên XWOD, ACDC và BDD. Một giả thuyết cần kiểm chứng thêm là embedding toàn ảnh có thể ưu tiên tương đồng về cảnh hoặc ngữ cảnh, chưa trực tiếp phản ánh độ khó định vị ở mức đối tượng và chưa kiểm soát đầy đủ tính đa dạng của Top-K. Đây là giả thuyết định hướng future work, không phải quan hệ nhân quả đã được chứng minh.",
    )


def renumber_existing_chapter4(doc: Document) -> None:
    heading_map = {
        "4.4. Phân tích kết quả": "4.5. Phân tích kết quả",
        "4.4.1. So sánh RT-DETR-L qua các giai đoạn": "4.5.1. So sánh RT-DETR-L qua các giai đoạn",
        "4.4.2. Phân tích theo lớp": "4.5.2. Phân tích theo lớp",
        "4.4.3. Phân tích theo điều kiện thời tiết": "4.5.3. Phân tích theo điều kiện thời tiết",
        "4.5. Các phân tích bổ trợ": "4.6. Các phân tích bổ trợ",
        "4.5.1. Quên thảm họa và khả năng phục hồi": "4.6.1. Quên thảm họa và khả năng phục hồi",
        "4.5.2. Đánh đổi giữa độ chính xác và tốc độ": "4.6.2. Đánh đổi giữa độ chính xác và tốc độ",
        "4.6. Thảo luận": "4.7. Thảo luận",
        "4.6.1. Các phát hiện chính": "4.7.1. Các phát hiện chính",
        "4.6.2. Hạn chế": "4.7.2. Hạn chế",
        "4.6.3. Hướng phát triển": "4.7.3. Hướng phát triển",
    }
    for old, new in heading_map.items():
        paragraph = find_exact(doc, old)
        paragraph.text = new
        format_heading(paragraph)
    caption_map = {
        "Bảng 4.7. Biến động mAP50-95 của RT-DETR-L qua các giai đoạn": "Bảng 4.10. Biến động mAP50-95 của RT-DETR-L qua các giai đoạn",
        "Bảng 4.8. mAP50-95 theo lớp của RT-DETR-L Phase 2": "Bảng 4.11. mAP50-95 theo lớp của RT-DETR-L Phase 2",
        "Bảng 4.9. mAP50-95 theo điều kiện thời tiết của RT-DETR-L Phase 2": "Bảng 4.12. mAP50-95 theo điều kiện thời tiết của RT-DETR-L Phase 2",
    }
    for old, new in caption_map.items():
        paragraph = find_exact(doc, old)
        paragraph.text = new
        paragraph.style = doc.styles["Table Caption"]
        paragraph.paragraph_format.keep_with_next = True


def update_discussion_future_conclusion(doc: Document) -> None:
    replace_exact(
        doc,
        "RT-DETR-L đạt mAP50-95 cao nhất trên XWOD validation và được lựa chọn cho Phase 2. YOLOv8n và YOLO11n có tốc độ cao hơn nhưng mAP50-95 thấp hơn. Stage 2 cải thiện thích nghi thời tiết bất lợi nhưng gây suy giảm BDD; Phase 2 phục hồi phần lớn mức giảm này, cải thiện ACDC, giữ XWOD ổn định và giảm nhẹ trên DAWN. RT-DETR-L Phase 2 là mô hình cuối cùng của đề tài.",
        "RT-DETR-L đạt mAP50-95 cao nhất trên XWOD validation và được lựa chọn cho Phase 2. Phase 2 phục hồi phần lớn suy giảm BDD sau Stage 2, cải thiện ACDC và duy trì XWOD, đồng thời làm baseline đa miền cho các ablation sau Phase 2. Trong ba cấu hình được frozen-evaluate, A0R đạt macro mAP50-95 cao nhất 0,4099; Phase 2 đạt 0,4062 và A1-DINO đạt 0,4056.",
    )
    replace_exact(
        doc,
        "Kết quả cho thấy fine-tuning đa miền tạo ra sự cân bằng tốt hơn giữa miền giao thông và miền thời tiết bất lợi, nhưng chưa giải quyết đồng đều mọi điều kiện. Kết luận của đề tài vì vậy tập trung vào hiệu quả của cấu hình RT-DETR-L Phase 2 thay vì quy nguyên nhân cho một nguồn dữ liệu đơn lẻ.",
        "A0R giữ nguyên XWOD và cải thiện ACDC, DAWN, BDD ở mAP50-95, nhưng không cải thiện mọi metric hoặc mọi lớp. A1-DINO đạt kết quả cao nhất trên DAWN nhưng giảm XWOD và ACDC. Do đó, bằng chứng hiện tại ủng hộ lợi ích của việc bổ sung dữ liệu lớp hiếm trong thiết lập này, nhưng chưa hỗ trợ giả thuyết rằng global image-level DINOv2 retrieval tốt hơn random rare-class sampling.",
    )
    replace_exact(
        doc,
        "Mô hình cuối vẫn có recall thấp trên ACDC, đặc biệt trong điều kiện ban đêm; bicycle, motorcycle và bus còn nhạy với số lượng mẫu, kích thước nhỏ, che khuất và chênh lệch miền. DAWN giảm nhẹ sau Phase 2, cho thấy cân bằng nhiều miền vẫn là một bài toán khó.",
        "Các cấu hình vẫn có hạn chế trên ACDC, đặc biệt trong điều kiện ban đêm; bicycle, motorcycle và bus còn nhạy với số lượng mẫu, kích thước nhỏ, che khuất và chênh lệch miền. A1-DINO cải thiện DAWN nhưng giảm cả sáu lớp trên XWOD, cho thấy chiến lược truy hồi toàn ảnh hiện tại chưa giải quyết ổn định mục tiêu rare-class robustness.",
    )
    future_replacements = {
        "Cải thiện lớp hiếm. Một hướng tiếp theo là khai thác phần BDD100K train chưa sử dụng để bổ sung có kiểm soát các ảnh chứa bicycle, motorcycle và bus. A0R có thể được xây dựng như một đối chứng random rare-class data-selection: xác lập candidate pool, chọn ngẫu nhiên một ngân sách ảnh cố định chứa lớp hiếm và fine-tune lại mô hình. Đối chứng này giúp tách ảnh hưởng của việc tăng dữ liệu khỏi ảnh hưởng của chiến lược lựa chọn mẫu. Class-aware sampling và object-level augmentation cũng có thể được khảo sát trong cùng giao thức.":
            "Truy hồi ở mức đối tượng. Kết quả A1-DINO gợi ý cần khảo sát object-level hoặc crop-level DINOv2 embeddings thay cho embedding toàn ảnh, nhằm liên kết trực tiếp hơn giữa mẫu truy vấn và đặc trưng của bicycle, motorcycle, bus.",
        "Lựa chọn dữ liệu dựa trên failure case. Thay vì chọn ngẫu nhiên, A1-DINO có thể sử dụng RT-DETR-L Phase 2 để xác định các mẫu khó, tập trung vào lớp hiếm hoặc bị bỏ sót, rồi tìm ảnh tương đồng trong candidate pool BDD100K đã có nhãn. Đây là failure-driven similarity retrieval, không phải active learning theo nghĩa nghiêm ngặt.":
            "Truy hồi có điều kiện theo lớp. Candidate retrieval có thể được thực hiện riêng cho bicycle, motorcycle và bus với ngân sách cân bằng, thay vì sử dụng một global Top-K chung cho mọi query và mọi lớp.",
        "DINOv2 có thể đóng vai trò cơ chế biểu diễn và truy hồi hình ảnh cho A1-DINO. Embedding tổng quát của các mẫu khó và candidate pool được so sánh bằng cosine similarity để xếp hạng ảnh nguồn; sau đó một ngân sách cố định được chọn để fine-tune lại RT-DETR-L. Cách phân vai này cho phép RT-DETR xác định mô hình thất bại ở đâu, còn DINOv2 xác định ảnh nguồn nào có đặc trưng gần với failure case. So sánh A0R với A1-DINO sẽ cho biết lựa chọn có chủ đích có tạo lợi ích vượt quá việc chỉ tăng dữ liệu lớp hiếm hay không.":
            "Lựa chọn có kiểm soát độ đa dạng. Một hướng tiếp theo là kết hợp failure score, semantic similarity và diversity; có thể phân cụm candidate trước khi chọn Top-K hoặc bổ sung uncertainty để hạn chế việc lấy nhiều ảnh quá giống nhau.",
        "Cải thiện điều kiện ban đêm. Có thể bổ sung dữ liệu night thật, áp dụng photometric nighttime augmentation, tạo dữ liệu ban đêm tổng hợp hoặc fine-tune theo miền ánh sáng. Các phương án cần được đánh giá riêng để phân biệt lợi ích của lượng dữ liệu với lợi ích của kiểu biến đổi ảnh.":
            "Cải thiện điều kiện ban đêm. Có thể bổ sung dữ liệu night thật, áp dụng photometric nighttime augmentation, tạo dữ liệu ban đêm tổng hợp hoặc fine-tune theo miền ánh sáng. Các phương án cần được đánh giá riêng để phân biệt lợi ích của lượng dữ liệu với lợi ích của kiểu biến đổi ảnh.",
        "Giảm quên miền và tối ưu triển khai. Các hướng phù hợp gồm replay có kiểm soát, balanced multi-domain sampling, regularization và continual/domain-adaptive fine-tuning. Đối với triển khai, có thể nghiên cứu distillation, quantization, TensorRT, biến thể RT-DETR nhẹ hơn và đánh giá trên thiết bị biên.":
            "Giảm quên miền và tối ưu triển khai. Các hướng phù hợp gồm balanced replay, balanced multi-domain sampling và regularization. Đối với triển khai, có thể nghiên cứu distillation, quantization, TensorRT, biến thể RT-DETR nhẹ hơn và đánh giá trên thiết bị biên.",
    }
    for old, new in future_replacements.items():
        replace_exact(doc, old, new)

    conclusion_updates = {
        "Khóa luận đã xây dựng và đánh giá quy trình phát hiện người và phương tiện trong điều kiện thời tiết bất lợi thông qua benchmark bốn kiến trúc YOLOv8n, YOLO11n, Faster R-CNN và RT-DETR-L. Trên cơ sở kết quả XWOD validation, RT-DETR-L được lựa chọn để tiếp tục huấn luyện ở Phase 2.":
            "Khóa luận đã xây dựng quy trình phát hiện người và phương tiện trong điều kiện thời tiết bất lợi, benchmark bốn kiến trúc YOLOv8n, YOLO11n, Faster R-CNN và RT-DETR-L, đồng thời tách rõ validation khỏi frozen test. RT-DETR-L đạt kết quả cao nhất trên XWOD validation và được lựa chọn để tiếp tục huấn luyện đa miền ở Phase 2.",
        "Phase 2 thực hiện fine-tuning đa miền với dữ liệu XWOD, ACDC và phần dữ liệu BDD100K dùng để phát lại. Kết quả tổng hợp cho thấy RT-DETR-L Phase 2 phục hồi đáng kể hiệu năng trên BDD, cải thiện trên ACDC và duy trì tương đối ổn định trên XWOD, dù kết quả trên DAWN giảm nhẹ.":
            "Phase 2 thực hiện fine-tuning đa miền với XWOD, ACDC và BDD100K, qua đó phục hồi đáng kể BDD sau quên miền ở Stage 2, tăng ACDC và duy trì hiệu năng mạnh trên XWOD. Phase 2 được sử dụng làm baseline và checkpoint khởi tạo chung cho hai post-Phase2 data-selection ablations A0R và A1-DINO.",
        "RT-DETR-L Phase 2 vì vậy được xác định là mô hình cuối cùng trong phạm vi khóa luận. Tuy nhiên, hiệu năng trong điều kiện ban đêm và trên một số lớp khó vẫn còn hạn chế, cho thấy mô hình chưa thích nghi đồng đều với mọi miền dữ liệu và điều kiện quan sát.":
            "Thí nghiệm sau Phase 2 cho thấy A0R đạt mAP50-95 trung bình cao nhất 0,4099 trên bốn tập kiểm thử, so với 0,4062 của Phase 2 và 0,4056 của A1-DINO. A0R giữ XWOD ở 0,5096, tăng ACDC lên 0,2542, đạt 0,5160 trên DAWN và tăng BDD lên 0,3596. A1-DINO đạt DAWN cao nhất 0,5253 nhưng giảm XWOD xuống 0,5004 và ACDC xuống 0,2424.",
        "Các hướng tiếp tục cải thiện dữ liệu lớp hiếm, lựa chọn mẫu theo failure case, tăng cường điều kiện ban đêm, giảm quên miền và tối ưu triển khai đã được trình bày tại mục 4.6.3. Đây là cơ sở để mở rộng nghiên cứu trong các vòng thực nghiệm tiếp theo.":
            "Trong thiết lập và phạm vi thí nghiệm của khóa luận, việc bổ sung dữ liệu liên quan đến lớp hiếm mang lại lợi ích tổng thể rõ hơn so với truy hồi dựa trên độ tương đồng embedding toàn ảnh bằng DINOv2. Kết quả chưa cho phép kết luận random sampling luôn vượt learned retrieval; thay vào đó, chúng định hướng các nghiên cứu tiếp theo về object-level embedding, class-conditioned retrieval, diversity-aware selection, robustness ban đêm, balanced replay và tối ưu triển khai tại mục 4.7.3.",
    }
    for old, new in conclusion_updates.items():
        replace_exact(doc, old, new)


def add_reference(doc: Document) -> None:
    appendix = find_exact(doc, "PHỤ LỤC")
    reference = appendix.insert_paragraph_before(
        '[16] M. Oquab và cộng sự, "DINOv2: Learning Robust Visual Features without Supervision," Transactions on Machine Learning Research, 2024. [Trực tuyến]. Địa chỉ: https://openreview.net/forum?id=a68SUt6zFt',
        style="Normal",
    )
    reference.alignment = WD_ALIGN_PARAGRAPH.LEFT
    reference.paragraph_format.first_line_indent = Cm(-0.63)
    reference.paragraph_format.left_indent = Cm(0.63)
    reference.paragraph_format.line_spacing = 1.5
    reference.paragraph_format.space_after = Pt(0)
    for run in reference.runs:
        set_run_font(run, 13)


def enable_update_fields(doc: Document) -> None:
    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def final_format_safeguards(doc: Document) -> None:
    for paragraph in doc.paragraphs:
        if paragraph.style.name.startswith("Heading"):
            format_heading(paragraph)
        if paragraph.style.name in ("Table Caption", "Figure Caption"):
            paragraph.paragraph_format.keep_together = True
        if paragraph.style.name == "Table Caption":
            paragraph.paragraph_format.keep_with_next = True
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for row in table.rows:
            set_no_split(row)
        if table.rows:
            set_repeat_header(table.rows[0])
    if "Heading 33" in [s.name for s in doc.styles]:
        stale = doc.styles["Heading 33"]._element
        stale.getparent().remove(stale)


def revise(source: Path, output: Path) -> None:
    doc = Document(source)
    update_intro(doc)
    update_chapter3(doc)
    update_phase2_tables_and_text(doc)
    insert_chapter4_ablation(doc)
    renumber_existing_chapter4(doc)
    replace_exact(doc, "4.3. Phase 2 và mô hình cuối cùng", "4.3. Phase 2 và mô hình nền đa miền")
    replace_exact(
        doc,
        "Chương này trình bày môi trường triển khai, cấu hình huấn luyện và kết quả đánh giá của quy trình nhiều giai đoạn. RT-DETR-L được lựa chọn bằng XWOD validation, tiếp tục ở Phase 2 và được xem là mô hình cuối cùng trong phạm vi thực nghiệm của đề tài.",
        "Chương này trình bày môi trường triển khai, benchmark bốn kiến trúc, Phase 2 và các post-Phase2 data-selection ablations. RT-DETR-L được lựa chọn bằng XWOD validation; Phase 2 là baseline đa miền, còn A0R và A1-DINO là hai nhánh tiếp tục huấn luyện từ cùng checkpoint Phase 2.",
    )
    replace_exact(
        doc,
        "Pipeline thực nghiệm gồm COCO pretrained → Stage 1 trên BDD100K → Stage 2 trên XWOD → XWOD validation → Phase 2 multi-domain fine-tuning. DAWN không tham gia huấn luyện; ACDC test sau Phase 2 là held-out same-source evaluation; cả bốn tập test chỉ được sử dụng sau khi cấu hình đã được cố định.",
        "Pipeline thực nghiệm gồm COCO pretrained → Stage 1 trên BDD100K → Stage 2 trên XWOD → XWOD validation → RT-DETR-L Phase 2 trên XWOD + ACDC + BDD → A0R/A1-DINO → frozen evaluation trên XWOD, ACDC, DAWN và BDD. DAWN không tham gia huấn luyện; BDD validation/test không tham gia candidate selection; frozen test không được dùng để chọn mô hình.",
    )
    replace_exact(
        doc,
        "Bảng dưới đây tổng hợp cấu hình huấn luyện của RT-DETR-L Phase 2, là mô hình cuối cùng của đề tài.",
        "Bảng dưới đây tổng hợp cấu hình huấn luyện của RT-DETR-L Phase 2, là baseline đa miền và checkpoint khởi tạo chung cho hai ablation A0R và A1-DINO.",
    )
    update_discussion_future_conclusion(doc)
    add_reference(doc)
    enable_update_fields(doc)
    final_format_safeguards(doc)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    revise(args.source, args.output)


if __name__ == "__main__":
    main()
