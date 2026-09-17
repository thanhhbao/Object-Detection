#!/usr/bin/env python3
"""Normalize thesis formatting against the supplied university template.

The script never overwrites its input. It preserves document content and media,
while normalizing page setup, sections, styles, headings, captions and tables.
"""

from __future__ import annotations

import argparse
import copy
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


MAJOR_FRONT_TITLES = {
    "LỜI MỞ ĐẦU",
    "LỜI CẢM ƠN",
    "NHẬN XÉT CỦA GIẢNG VIÊN HƯỚNG DẪN",
    "NHẬN XÉT CỦA GIẢNG VIÊN PHẢN BIỆN",
    "MỤC LỤC",
    "DANH MỤC CHỮ VIẾT TẮT",
    "DANH MỤC HÌNH",
    "DANH MỤC BẢNG",
    "TÀI LIỆU THAM KHẢO",
}


def set_east_asia_font(style_or_run, name: str) -> None:
    style_or_run.font.name = name
    rpr = style_or_run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), name)


def remove_child(parent, tag: str) -> None:
    child = parent.find(qn(tag))
    if child is not None:
        parent.remove(child)


def set_keep(paragraph, *, next_: bool | None = None, together: bool | None = None) -> None:
    if next_ is not None:
        paragraph.paragraph_format.keep_with_next = next_
    if together is not None:
        paragraph.paragraph_format.keep_together = together


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char_sep = OxmlElement("w:fldChar")
    fld_char_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    for el in (fld_char_begin, instr, fld_char_sep, text, fld_char_end):
        run._r.append(el)


def clear_container(container) -> None:
    for paragraph in container.paragraphs:
        for run in paragraph.runs:
            paragraph._p.remove(run._r)
    while len(container.paragraphs) > 1:
        container._element.remove(container.paragraphs[-1]._p)


def add_section_break_after(paragraph, source_sect_pr) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    old = ppr.find(qn("w:sectPr"))
    if old is not None:
        ppr.remove(old)
    ppr.append(copy.deepcopy(source_sect_pr))


def remove_explicit_page_breaks(paragraph) -> None:
    for br in paragraph._p.xpath('.//w:br[@w:type="page"]'):
        br.getparent().remove(br)


def set_page_numbering(section, fmt: str | None, start: int | None) -> None:
    sect_pr = section._sectPr
    pg = sect_pr.find(qn("w:pgNumType"))
    if fmt is None and start is None:
        if pg is not None:
            sect_pr.remove(pg)
        return
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        sect_pr.append(pg)
    if fmt:
        pg.set(qn("w:fmt"), fmt)
    if start is not None:
        pg.set(qn("w:start"), str(start))


def normalize_sections(doc: Document) -> None:
    # Existing section break is immediately after the second cover. Add one
    # after the first cover and one before Chapter 1, producing four sections.
    paragraphs = doc.paragraphs
    final_sect_pr = doc._element.body.sectPr
    existing_break = next(p._p.pPr.sectPr for p in paragraphs if p._p.pPr is not None and p._p.pPr.sectPr is not None)

    first_cover_end = next(p for p in paragraphs if p.text.strip().startswith("Tp HCM") or p.text.strip().startswith("Tp.HCM"))
    add_section_break_after(first_cover_end, existing_break)
    first_cover_idx = paragraphs.index(first_cover_end)
    if first_cover_idx + 1 < len(paragraphs):
        remove_explicit_page_breaks(paragraphs[first_cover_idx + 1])

    chapter_one = next(p for p in paragraphs if p.text.strip().startswith("CHƯƠNG 1"))
    chapter_idx = paragraphs.index(chapter_one)
    previous = paragraphs[chapter_idx - 1]
    add_section_break_after(previous, final_sect_pr)

    # Refresh the section proxies after modifying the XML.
    sections = list(doc.sections)
    if len(sections) != 4:
        raise RuntimeError(f"Expected 4 sections after normalization, found {len(sections)}")

    for section in sections:
        section.start_type = WD_SECTION.NEW_PAGE
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3.5)
        section.right_margin = Cm(2.0)
        section.header_distance = Cm(1.27)
        section.footer_distance = Cm(1.27)
        # Preserve the original blank-cover and PAGE footer relationships.
        # The copied section properties deliberately reuse those valid parts:
        # sections 1-2 use the blank footer, sections 3-4 use the PAGE footer.
    set_page_numbering(sections[0], None, None)
    set_page_numbering(sections[1], None, None)
    set_page_numbering(sections[2], "lowerRoman", 1)
    set_page_numbering(sections[3], "decimal", 1)


def normalize_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    set_east_asia_font(normal, "Times New Roman")
    normal.font.size = Pt(13)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.widow_control = True

    try:
        normal_web = doc.styles["Normal (Web)"]
        normal_web.base_style = normal
        set_east_asia_font(normal_web, "Times New Roman")
        normal_web.font.size = Pt(13)
        normal_web.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        normal_web.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        normal_web.paragraph_format.space_before = Pt(0)
        normal_web.paragraph_format.space_after = Pt(0)
    except KeyError:
        pass

    h1 = doc.styles["Heading 1"]
    set_east_asia_font(h1, "Times New Roman")
    h1.font.size = Pt(16)
    h1.font.bold = True
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(12)
    h1.paragraph_format.keep_with_next = True
    h1.paragraph_format.keep_together = True
    h1.paragraph_format.page_break_before = True

    h2 = doc.styles["Heading 2"]
    set_east_asia_font(h2, "Times New Roman")
    h2.font.size = Pt(13)
    h2.font.bold = True
    h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h2.paragraph_format.left_indent = Cm(0)
    h2.paragraph_format.first_line_indent = Cm(0)
    h2.paragraph_format.space_before = Pt(2)
    h2.paragraph_format.space_after = Pt(0)
    h2.paragraph_format.keep_with_next = True
    h2.paragraph_format.keep_together = True

    try:
        h3 = doc.styles["Heading 3"]
    except KeyError:
        h3 = doc.styles.add_style("Heading 3", WD_STYLE_TYPE.PARAGRAPH)
    h3.base_style = normal
    set_east_asia_font(h3, "Times New Roman")
    h3.font.size = Pt(13)
    h3.font.bold = True
    h3.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h3.paragraph_format.left_indent = Cm(0)
    h3.paragraph_format.first_line_indent = Cm(0)
    h3.paragraph_format.space_before = Pt(2)
    h3.paragraph_format.space_after = Pt(0)
    h3.paragraph_format.keep_with_next = True
    h3.paragraph_format.keep_together = True
    h3.element.get_or_add_pPr().get_or_add_outlineLvl().val = 2

    try:
        caption = doc.styles["Caption"]
    except KeyError:
        caption = doc.styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)
    caption.base_style = normal
    set_east_asia_font(caption, "Times New Roman")
    caption.font.size = Pt(12)
    caption.font.italic = True
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.left_indent = Cm(0)
    caption.paragraph_format.first_line_indent = Cm(0)
    caption.paragraph_format.line_spacing = 1.0
    caption.paragraph_format.space_before = Pt(6)
    caption.paragraph_format.space_after = Pt(16)
    caption.paragraph_format.keep_together = True

    try:
        source = doc.styles["Figure Source"]
    except KeyError:
        source = doc.styles.add_style("Figure Source", WD_STYLE_TYPE.PARAGRAPH)
    source.base_style = normal
    set_east_asia_font(source, "Times New Roman")
    source.font.size = Pt(11)
    source.font.italic = True
    source.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    source.paragraph_format.left_indent = Cm(0)
    source.paragraph_format.first_line_indent = Cm(0)
    source.paragraph_format.line_spacing = 1.0
    source.paragraph_format.space_before = Pt(0)
    source.paragraph_format.space_after = Pt(12)

    for toc_name, left_cm, bold in (("TOC 1", 0.0, True), ("TOC 2", 0.39, False), ("TOC 3", 0.78, False)):
        try:
            style = doc.styles[toc_name]
        except KeyError:
            continue
        set_east_asia_font(style, "Times New Roman")
        style.font.size = Pt(13)
        style.font.bold = bold
        style.paragraph_format.left_indent = Cm(left_cm)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.line_spacing = 1.0


def strip_direct_heading_format(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    for tag in ("w:ind", "w:jc", "w:spacing", "w:numPr", "w:keepNext", "w:keepLines", "w:pageBreakBefore"):
        remove_child(ppr, tag)


def normalize_paragraphs(doc: Document) -> None:
    paragraphs = doc.paragraphs

    # The copied template contains thick blue paragraph borders on both covers.
    # Remove borders only through the end of the second cover.
    second_cover_end = [p for p in paragraphs if p.text.strip().startswith("Tp HCM") or p.text.strip().startswith("Tp.HCM")][1]
    cover_end_idx = paragraphs.index(second_cover_end)
    for paragraph in paragraphs[: cover_end_idx + 1]:
        ppr = paragraph._p.pPr
        if ppr is not None:
            remove_child(ppr, "w:pBdr")

    # Convert the six Chapter 1 auto-numbered headings into stable text numbers,
    # matching the explicit numbering used by the remaining chapters/template.
    chapter = 0
    chapter_one_counter = 0
    for paragraph in paragraphs:
        text = paragraph.text.strip()
        if text.startswith("CHƯƠNG 1"):
            chapter = 1
        elif text.startswith("CHƯƠNG 2"):
            chapter = 2
        elif text.startswith("CHƯƠNG 3"):
            chapter = 3
            paragraph.text = text.replace("CHƯƠNG 3.", "CHƯƠNG 3:", 1)
        elif text.startswith("CHƯƠNG 4"):
            chapter = 4
            paragraph.text = text.replace("CHƯƠNG 4.", "CHƯƠNG 4:", 1)

        if paragraph.style.name == "Heading 32":
            paragraph.style = doc.styles["Heading 3"]

        if paragraph.style.name.startswith("Heading"):
            if chapter == 1 and paragraph.style.name == "Heading 2" and not text[:1].isdigit():
                chapter_one_counter += 1
                paragraph.text = f"1.{chapter_one_counter}. {text}"
            strip_direct_heading_format(paragraph)
            if paragraph.style.name == "Heading 1":
                paragraph.paragraph_format.page_break_before = True
            else:
                set_keep(paragraph, next_=True, together=True)

    # Major preliminary parts must not start halfway down the preceding part.
    for paragraph in paragraphs:
        text = paragraph.text.strip()
        if text in MAJOR_FRONT_TITLES:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.left_indent = Cm(0)
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.page_break_before = True
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.keep_together = True
            for run in paragraph.runs:
                set_east_asia_font(run, "Times New Roman")
                run.font.size = Pt(16)
                run.font.bold = True

    in_manual_lists = False
    for idx, paragraph in enumerate(paragraphs):
        text = paragraph.text.strip()
        if text in {"DANH MỤC HÌNH", "DANH MỤC BẢNG"}:
            in_manual_lists = True
            continue
        if text.startswith("CHƯƠNG 1"):
            in_manual_lists = False

        if paragraph.style.name in {"Caption2", "Caption3", "Caption1"}:
            if text.startswith("Nguồn:"):
                paragraph.style = doc.styles["Figure Source"]
            else:
                paragraph.style = doc.styles["Caption"]
                paragraph.paragraph_format.left_indent = Cm(0)
                paragraph.paragraph_format.first_line_indent = Cm(0)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if text.startswith("Bảng"):
                    set_keep(paragraph, next_=True, together=True)
                elif text.startswith("Hình") and idx > 0:
                    set_keep(paragraphs[idx - 1], next_=True, together=True)
            continue

        if in_manual_lists and (text.startswith("Hình ") or text.startswith("Bảng ")):
            paragraph.paragraph_format.left_indent = Cm(0)
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.space_after = Pt(5)
            continue

        if paragraph.style.name in {"Normal", "Normal (Web)"} and text and idx > cover_end_idx:
            # Preserve centered signatures, title-like lines, image/equation anchors,
            # and code blocks; normalize ordinary prose only.
            is_title_like = text in MAJOR_FRONT_TITLES or (text.isupper() and len(text) < 90)
            is_signature = any(key in text for key in ("Ký tên", "Sinh viên thực hiện", "Giáo viên hướng dẫn", "Giáo viên phản biện"))
            is_code = any(run.font.name == "Consolas" for run in paragraph.runs)
            if not is_title_like and not is_signature and not is_code:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                paragraph.paragraph_format.left_indent = Cm(0)
                paragraph.paragraph_format.first_line_indent = Cm(1.0)
                paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.widow_control = True


def set_cell_shading(cell, fill: str | None) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if fill is None:
        if shd is not None:
            tc_pr.remove(shd)
        return
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_row_property(row, tag: str) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn(tag)) is None:
        tr_pr.append(OxmlElement(tag))


def normalize_tables(doc: Document) -> None:
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        try:
            table.style = "Table Grid"
        except KeyError:
            pass
        for row_idx, row in enumerate(table.rows):
            set_row_property(row, "w:cantSplit")
            if row_idx == 0:
                set_row_property(row, "w:tblHeader")
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_shading(cell, None)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.left_indent = Cm(0)
                    paragraph.paragraph_format.first_line_indent = Cm(0)
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.0
                    if row_idx == 0:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        set_east_asia_font(run, "Times New Roman")
                        run.font.size = Pt(12)
                        if row_idx == 0:
                            run.font.bold = True


def mark_fields_dirty(doc: Document) -> None:
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")
    for fld in doc._element.xpath("//w:fldChar[@w:fldCharType='begin']"):
        fld.set(qn("w:dirty"), "true")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--backup", type=Path)
    args = parser.parse_args()

    args.output.parent.mkdir(parents=True, exist_ok=True)
    if args.backup:
        args.backup.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(args.source, args.backup)

    doc = Document(args.source)
    normalize_sections(doc)
    normalize_styles(doc)
    normalize_paragraphs(doc)
    normalize_tables(doc)
    mark_fields_dirty(doc)
    doc.core_properties.title = "Khóa luận tốt nghiệp – bản đã chuẩn hóa định dạng"
    doc.save(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
