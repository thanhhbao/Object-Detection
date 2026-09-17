#!/usr/bin/env python3
"""Rewrite Chapter 3 of the thesis while preserving the rest of the DOCX."""

from __future__ import annotations

from pathlib import Path
import sys

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
FIGURES = ROOT / "docs" / "thesis" / "figures"


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, *, bold: bool = False) -> None:
    cell.text = text
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        for run in paragraph.runs:
            run.bold = bold
            run.font.size = Pt(10)


class ChapterBuilder:
    def __init__(self, document: Document, anchor):
        self.document = document
        self.anchor = anchor
        self.blocks = []

    def paragraph(self, text: str = "", style: str = "Normal", *, align=None):
        p = self.document.add_paragraph(style=style)
        p.add_run(text)
        if align is not None:
            p.alignment = align
        self.blocks.append(p._p)
        return p

    def heading2(self, text: str):
        return self.paragraph(text, "Heading 2")

    def heading3(self, text: str):
        return self.paragraph(text, "Heading 32")

    def caption(self, text: str):
        return self.paragraph(text, "Caption2", align=WD_ALIGN_PARAGRAPH.CENTER)

    def picture(self, path: Path, width: float = 6.15):
        p = self.document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(width))
        self.blocks.append(p._p)
        return p

    def table(self, headers: list[str], rows: list[list[str]], widths=None):
        table = self.document.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        for index, header in enumerate(headers):
            set_cell_text(table.rows[0].cells[index], header, bold=True)
            set_cell_shading(table.rows[0].cells[index], "D9EAF7")
        set_repeat_table_header(table.rows[0])
        for values in rows:
            cells = table.add_row().cells
            for index, value in enumerate(values):
                set_cell_text(cells[index], value)
        if widths:
            for row in table.rows:
                for index, width in enumerate(widths):
                    row.cells[index].width = Inches(width)
        self.blocks.append(table._tbl)
        return table

    def commit(self):
        for block in self.blocks:
            self.anchor.addprevious(block)


def remove_chapter3_body(document: Document):
    body = document.element.body
    ch3 = None
    ch4 = None
    for child in body.iterchildren():
        if not child.tag.endswith("}p"):
            continue
        texts = child.xpath(".//w:t/text()")
        text = "".join(texts).strip()
        if text.startswith("CHƯƠNG 3") and ch3 is None:
            p_style = child.find("./w:pPr/w:pStyle", namespaces=child.nsmap)
            if p_style is not None and p_style.get(qn("w:val")) == "Heading1":
                ch3 = child
        elif ch3 is not None and text.startswith("CHƯƠNG 4"):
            ch4 = child
            break
    if ch3 is None or ch4 is None:
        raise RuntimeError("Không tìm thấy ranh giới Chương 3/Chương 4.")
    current = ch3.getnext()
    while current is not None and current is not ch4:
        next_element = current.getnext()
        body.remove(current)
        current = next_element
    return ch3, ch4


def add_chapter3(builder: ChapterBuilder) -> None:
    builder.paragraph(
        "Chương 2 đã trình bày nền tảng lý thuyết của bài toán phát hiện đối tượng, các họ mô hình tiêu biểu, "
        "học chuyển giao, cơ chế chú ý và các bộ dữ liệu liên quan. Trên cơ sở đó, Chương 3 chuyển các mục tiêu "
        "nghiên cứu thành một thiết kế hệ thống và kế hoạch thực nghiệm có thể kiểm chứng. Nội dung không chỉ mô "
        "tả kiến trúc phần mềm mà còn xác định rõ vai trò của từng nguồn dữ liệu, không gian nhãn thống nhất, luồng "
        "huấn luyện, biến kiểm soát, tập đánh giá và nguyên tắc truy vết kết quả."
    )
    builder.paragraph(
        "Thiết kế được xây dựng theo hai yêu cầu xuyên suốt. Thứ nhất, mô hình phải hoạt động trên sáu lớp giao thông "
        "theo đúng thứ tự person, bicycle, car, motorcycle, bus và truck. Thứ hai, mọi kết luận phải tách bạch giữa "
        "độ chính xác tuyệt đối và tính phù hợp triển khai. Vì vậy, mô hình có độ chính xác cao nhất trong benchmark "
        "không mặc nhiên là mô hình thực dụng cuối cùng; quyết định còn xét tốc độ, kích thước mô hình, khả năng triển "
        "khai và mức ổn định khi chuyển miền. Kết quả định lượng của các thiết kế trong chương này được trình bày ở Chương 4."
    )

    builder.heading2("3.1. Phân tích bài toán và yêu cầu hệ thống")
    builder.paragraph(
        "Bài toán đầu vào là ảnh hoặc chuỗi khung hình giao thông có thể bị suy giảm bởi mưa, sương mù, tuyết, ngập "
        "hoặc bão cát. Đầu ra của mô hình là tập các hộp giới hạn, nhãn lớp và độ tin cậy tương ứng. Khó khăn chính "
        "không chỉ nằm ở nhiễu hình ảnh mà còn ở chênh lệch phân phối giữa các bộ dữ liệu, mất cân bằng lớp và số "
        "lượng nhỏ của các đối tượng hiếm. Do đó, hệ thống phải giải quyết đồng thời ba lớp vấn đề: chuẩn hóa dữ liệu, "
        "học biểu diễn có khả năng chuyển miền và đánh giá không thiên lệch."
    )

    builder.heading3("3.1.1. Mục tiêu và phạm vi thiết kế")
    builder.paragraph(
        "Mục tiêu thứ nhất là xây dựng một pipeline tái sử dụng được cho bốn kiến trúc đại diện: YOLOv8n, YOLO11n, "
        "Faster R-CNN và RT-DETR. Pipeline phải cho phép huấn luyện theo các giai đoạn COCO pretrained → BDD100K → "
        "XWOD, đánh giá trên XWOD test và DAWN val, sau đó mở rộng mô hình thực dụng được chọn bằng dữ liệu Phase 2."
    )
    builder.paragraph(
        "Mục tiêu thứ hai là nghiên cứu cách cải thiện mô hình trong điều kiện dữ liệu đích nhỏ và mất cân bằng. Các "
        "hướng khảo sát gồm cơ chế chú ý SE, kết hợp XWOD với ACDC và BDD replay, lấy mẫu tăng cường cho các lớp hiếm, "
        "tăng kích thước ảnh đầu vào, tăng năng lực mô hình và copy-paste. Mỗi hướng được coi là một giả thuyết thực "
        "nghiệm; chỉ những thay đổi tạo cải thiện nhất quán mới được giữ trong cấu hình cuối. CBAM không thuộc nhóm "
        "thực nghiệm đã triển khai và chỉ được xem là hướng nghiên cứu tiếp theo."
    )
    builder.paragraph(
        "Mục tiêu thứ ba là cung cấp một mô hình thực dụng có thể phục vụ demo ảnh/video. Phạm vi khóa luận tập trung "
        "vào phát hiện sáu lớp, không mở rộng sang theo vết đối tượng, ước lượng khoảng cách hay điều khiển phương tiện. "
        "Tốc độ được báo cáo trên môi trường phần cứng xác định; vì vậy không diễn giải FPS như một hằng số độc lập với "
        "GPU, phiên bản thư viện, batch size và kích thước ảnh."
    )

    builder.heading3("3.1.2. Yêu cầu chức năng và tiêu chí chấp nhận")
    builder.paragraph(
        "Các yêu cầu chức năng bao phủ toàn bộ vòng đời từ dữ liệu đến suy luận. Bên cạnh mô-đun demo, pipeline nghiên "
        "cứu phải lưu được cấu hình và tạo kết quả đánh giá theo từng lớp để có thể đối chiếu giữa các run. Bảng 3.1 "
        "trình bày yêu cầu cùng tiêu chí chấp nhận, giúp phân biệt một chức năng đã được triển khai với một mô tả mang tính ý tưởng."
    )
    builder.caption("Bảng 3.1. Yêu cầu chức năng và tiêu chí chấp nhận của hệ thống")
    builder.table(
        ["Mã", "Chức năng", "Mô tả", "Tiêu chí chấp nhận"],
        [
            ["FR1", "Tiếp nhận dữ liệu", "Nhận ảnh hoặc video từ tệp đầu vào.", "Đọc được định dạng hợp lệ và báo lỗi rõ ràng với tệp hỏng."],
            ["FR2", "Phát hiện sáu lớp", "Sinh hộp giới hạn, nhãn và độ tin cậy.", "Chỉ số lớp tuân thủ thứ tự canonical 0–5."],
            ["FR3", "Điều chỉnh hậu xử lý", "Cho phép cấu hình ngưỡng confidence và IoU/NMS khi áp dụng.", "Thay đổi ngưỡng làm thay đổi kết quả mà không phải huấn luyện lại."],
            ["FR4", "Trực quan hóa", "Vẽ hộp, nhãn, độ tin cậy và thống kê theo lớp.", "Kết quả hiển thị khớp với đầu ra số của mô hình."],
            ["FR5", "Xuất kết quả", "Lưu ảnh/video đã chú thích và bảng thống kê.", "Tệp xuất có thể mở lại và giữ đúng số lượng phát hiện."],
            ["FR6", "Huấn luyện nhiều giai đoạn", "Hỗ trợ Stage 1, Stage 2 và Phase 2 theo cấu hình.", "Mỗi run có checkpoint, cấu hình và nhật ký riêng."],
            ["FR7", "Đánh giá đa miền", "Tính P, R, mAP50 và mAP50-95 trên các split quy định.", "Kết quả tổng thể và theo lớp được lưu cùng tên dataset/split."],
            ["FR8", "Truy vết thực nghiệm", "Liên kết checkpoint với args, results và log đánh giá.", "Có thể xác định nguồn của mỗi số liệu được dùng trong báo cáo."],
        ],
        widths=[0.55, 1.25, 2.15, 2.35],
    )

    builder.heading3("3.1.3. Yêu cầu phi chức năng và ràng buộc nghiên cứu")
    builder.paragraph(
        "Các yêu cầu phi chức năng được diễn giải theo khả năng đo lường. Độ chính xác phải được xem đồng thời với recall "
        "và kết quả theo lớp, vì mAP trung bình có thể che khuất sự suy giảm của bicycle, motorcycle hoặc bus. Khả năng "
        "tổng quát hóa phải gắn với đúng quan hệ train/evaluation: DAWN là đánh giá ngoài miền và không tham gia Stage 2; "
        "ACDC val là zero-shot trước Phase 2 nhưng trở thành tập held-out cùng nguồn sau khi ACDC train đã được sử dụng."
    )
    builder.caption("Bảng 3.2. Yêu cầu phi chức năng và cách kiểm chứng")
    builder.table(
        ["Mã", "Yêu cầu", "Cách kiểm chứng"],
        [
            ["NFR1", "Hiệu quả suy luận", "Đo ms/ảnh hoặc FPS trong cùng phần cứng, batch size và kích thước đầu vào."],
            ["NFR2", "Độ chính xác", "Báo cáo P, R, mAP50, mAP50-95 tổng thể và mAP50-95 theo sáu lớp."],
            ["NFR3", "Khả năng chuyển miền", "Đánh giá riêng trên DAWN val; diễn giải ACDC val theo thời điểm trước/sau Phase 2."],
            ["NFR4", "Tính gọn nhẹ", "So sánh tham số, chi phí tính toán, độ phân giải và độ chính xác, không chỉ một chỉ số."],
            ["NFR5", "Khả năng tái lập", "Lưu seed, args, checkpoint, results.csv và log evaluation; ghi rõ các gap provenance."],
            ["NFR6", "Tính đúng dữ liệu", "Kiểm tra ảnh/nhãn, class ID, split và số mẫu hợp lệ trước khi đánh giá."],
            ["NFR7", "Khả năng mở rộng", "Tách mô-đun dữ liệu, huấn luyện, đánh giá và ứng dụng để thay thế độc lập."],
        ],
        widths=[0.65, 1.6, 4.1],
    )

    builder.heading2("3.2. Kiến trúc tổng thể của hệ thống")
    builder.paragraph(
        "Hệ thống được tổ chức thành bốn tầng: dữ liệu, mô hình, huấn luyện–đánh giá và ứng dụng. Việc phân tầng giúp "
        "một thay đổi ở bộ dữ liệu hoặc mô hình không buộc phải viết lại toàn bộ pipeline. Đồng thời, kết quả nghiên cứu "
        "được tách khỏi giao diện demo: giao diện chỉ tiêu thụ checkpoint đã chọn, không quyết định cách huấn luyện hay đánh giá."
    )

    builder.heading3("3.2.1. Các tầng chức năng")
    builder.paragraph(
        "Tầng dữ liệu quản lý BDD100K, XWOD, DAWN và ACDC, bao gồm ánh xạ lớp, chuyển đổi nhãn, kiểm tra ảnh hỏng, "
        "tổ chức split và tạo tệp YAML. Tầng mô hình cung cấp adapter cho các kiến trúc Ultralytics và torchvision, "
        "để đầu ra cuối cùng cùng quy về hộp giới hạn và sáu lớp. Tầng huấn luyện–đánh giá thực hiện fine-tune, lưu "
        "checkpoint, tổng hợp metrics và chạy đánh giá theo từng split. Tầng ứng dụng thực hiện suy luận, hậu xử lý, "
        "trực quan hóa và xuất kết quả."
    )
    builder.picture(FIGURES / "hinh_3_1_kien_truc.png")
    builder.caption("Hình 3.1. Kiến trúc phân tầng của hệ thống")
    builder.paragraph(
        "Luồng phụ thuộc được thiết kế một chiều: dữ liệu chuẩn hóa được cung cấp cho huấn luyện; huấn luyện tạo checkpoint; "
        "checkpoint được khóa trước khi đánh giá; kết quả đánh giá mới được dùng cho lựa chọn mô hình. Cách tổ chức này hạn "
        "chế việc điều chỉnh mô hình trực tiếp theo tập test và giảm nguy cơ rò rỉ thông tin giữa các bước."
    )

    builder.heading3("3.2.2. Luồng dữ liệu và kiểm soát chất lượng")
    builder.paragraph(
        "Mỗi nguồn dữ liệu đi qua chuỗi kiểm tra gồm: xác nhận tệp ảnh có thể giải mã, đọc nhãn gốc, ánh xạ về sáu lớp, "
        "loại nhãn ngoài phạm vi, chuẩn hóa tọa độ, kiểm tra hộp nằm trong biên ảnh và thống kê phân bố lớp. Ảnh không "
        "có đối tượng mục tiêu vẫn có thể được giữ như mẫu âm nếu cấu hình dataset cho phép; ảnh hỏng hoặc sai định dạng "
        "phải bị loại và được ghi nhận, không âm thầm làm thay đổi mẫu số đánh giá."
    )
    builder.paragraph(
        "Một trường hợp cụ thể là XWOD test có 2.321 tệp nguồn nhưng một tệp mang phần mở rộng JPG thực chất được nhận "
        "diện là GIF không hợp lệ đối với pipeline hiện tại. Ultralytics bỏ qua tệp này, nên số ảnh hợp lệ thực tế trong "
        "evaluation là 2.320. Báo cáo vì vậy phân biệt rõ số tệp nguồn và số ảnh được đánh giá; chênh lệch này không được "
        "xem là một split chưa giải thích."
    )

    builder.heading3("3.2.3. Luồng huấn luyện, đánh giá và lựa chọn")
    builder.paragraph(
        "Luồng nghiên cứu bắt đầu từ checkpoint pretrained, tiếp tục qua benchmark nhiều giai đoạn và chỉ sau khi benchmark "
        "mới chọn mô hình thực dụng để mở rộng. Mọi tập đánh giá được gọi bằng tên dataset và split cụ thể nhằm tránh nhầm "
        "XWOD val với XWOD test. DAWN val chỉ được dùng để đo khả năng chuyển miền; ACDC train chỉ xuất hiện ở Phase 2, "
        "trong khi ACDC val luôn được giữ lại để đánh giá."
    )
    builder.picture(FIGURES / "hinh_3_2_pipeline.png")
    builder.caption("Hình 3.2. Pipeline huấn luyện nhiều giai đoạn và đánh giá đa miền")

    builder.heading2("3.3. Thiết kế và chuẩn hóa dữ liệu")
    builder.paragraph(
        "Dữ liệu là biến thiết kế quan trọng nhất của đề tài vì bốn nguồn khác nhau về miền ảnh, hệ nhãn và mục đích sử "
        "dụng. Thiết kế dữ liệu không gộp tùy ý tất cả các tập, mà gán vai trò rõ ràng cho từng nguồn và duy trì ranh giới "
        "train/val/test. Điều này cho phép phân biệt cải thiện do thích nghi miền với cải thiện do vô tình nhìn thấy dữ liệu đánh giá."
    )

    builder.heading3("3.3.1. Vai trò và phân chia các bộ dữ liệu")
    builder.caption("Bảng 3.3. Vai trò của các bộ dữ liệu trong pipeline")
    builder.table(
        ["Bộ dữ liệu", "Quy mô được dùng", "Vai trò", "Nguyên tắc tách biệt"],
        [
            ["BDD100K", "BDD30K train ≈ 30.000; val ≈ 10.000", "Stage 1 và nguồn replay ở Phase 2", "Replay thực tế là 2.000 ảnh, không diễn giải là 30% BDD30K."],
            ["XWOD", "train 5.945; val 1.744; test 2.321 tệp nguồn/2.320 ảnh hợp lệ", "Dữ liệu thời tiết bất lợi chính cho Stage 2 và Phase 2", "Val phục vụ chọn checkpoint; test chỉ dùng báo cáo cuối."],
            ["DAWN", "val 206 ảnh, 1.806 instances", "Đánh giá ngoài miền", "Không dùng trong Stage 2 hoặc Phase 2 training."],
            ["ACDC", "train 1.572; val 398", "Mở rộng Phase 2 và đánh giá held-out", "Train được gộp ở Phase 2; val không tham gia huấn luyện."],
        ],
        widths=[1.0, 1.6, 1.8, 2.0],
    )
    builder.paragraph(
        "Trước Phase 2, ACDC val có thể được xem là đánh giá cross-dataset/zero-shot vì mô hình chưa học từ ACDC train. "
        "Sau Phase 2, cách gọi đúng là held-out evaluation trên cùng nguồn dữ liệu ACDC. Sự thay đổi ngữ nghĩa này phải "
        "được giữ nhất quán khi so sánh các mốc mô hình trong Chương 4."
    )

    builder.heading3("3.3.2. Không gian nhãn canonical")
    builder.paragraph(
        "Không gian nhãn thống nhất gồm sáu lớp và là hợp đồng dữ liệu giữa mọi mô-đun. Thứ tự bắt buộc là 0 person, "
        "1 bicycle, 2 car, 3 motorcycle, 4 bus, 5 truck. Mọi tệp YAML, nhãn YOLO, tensor đầu ra, bảng kết quả và màu "
        "trực quan hóa phải tuân thủ cùng thứ tự; chỉ cần hoán đổi car và motorcycle cũng khiến số liệu theo lớp bị gán sai nghĩa."
    )
    builder.caption("Bảng 3.4. Ánh xạ nhãn về sáu lớp mục tiêu")
    builder.table(
        ["ID", "Lớp mục tiêu", "BDD100K", "XWOD", "ACDC/Cityscapes"],
        [
            ["0", "person", "person, rider", "person", "person (24), rider (25)"],
            ["1", "bicycle", "bike", "bicycle", "bicycle (33)"],
            ["2", "car", "car", "car", "car (26)"],
            ["3", "motorcycle", "motor", "motorcycle", "motorcycle (32)"],
            ["4", "bus", "bus", "bus", "bus (28)"],
            ["5", "truck", "truck", "truck", "truck (27)"],
            ["—", "Loại bỏ", "traffic light, traffic sign, train", "lớp ngoài sáu lớp", "train (31), lớp nền"],
        ],
        widths=[0.45, 1.1, 1.55, 1.2, 2.0],
    )

    builder.heading3("3.3.3. Chuyển đổi nhãn và tiền xử lý ảnh")
    builder.paragraph(
        "Nhãn hộp giới hạn được chuyển về định dạng YOLO gồm class_id, tâm hộp (cx, cy), chiều rộng w và chiều cao h, "
        "trong đó bốn tọa độ được chuẩn hóa theo kích thước ảnh. Với ACDC, nhãn phân đoạn instance/panoptic được giải mã "
        "theo ID lớp; mỗi vùng đối tượng hợp lệ được bao bởi hình chữ nhật nhỏ nhất và ánh xạ sang lớp canonical. Các vùng "
        "nền, lớp ngoài phạm vi và hộp suy biến được loại bỏ trước khi ghi nhãn."
    )
    builder.paragraph(
        "Ảnh được letterbox về kích thước đầu vào thay vì kéo giãn độc lập theo hai chiều. Tỉ lệ co giãn được giữ nguyên, "
        "phần thiếu được đệm và tọa độ hộp được biến đổi tương ứng. Cấu hình benchmark và các run cơ sở sử dụng imgsz 640; "
        "các thí nghiệm độ phân giải cao sử dụng imgsz 960. Kích thước ảnh vì vậy là một biến thực nghiệm, không phải một "
        "giá trị cố định chung cho mọi run."
    )

    builder.heading3("3.3.4. Thiết kế dữ liệu Phase 2 và xử lý mất cân bằng")
    builder.paragraph(
        "Tập cơ sở của Phase 2 v2/v3 gồm 5.945 ảnh XWOD train, 1.572 ảnh ACDC train và 2.000 ảnh BDD replay, tổng cộng "
        "9.517 ảnh. BDD replay có vai trò giữ lại một phần tri thức miền giao thông thông thường và hạn chế quên thảm họa, "
        "nhưng chỉ chiếm khoảng 6,7% so với quy mô BDD30K ban đầu. Field metadata cũ bdd_replay_ratio = 0.3, nếu xuất hiện, "
        "không đại diện cho composition đã build và không được dùng để mô tả 2.000 ảnh replay."
    )
    builder.caption("Bảng 3.5. Thành phần tập huấn luyện Phase 2 v3")
    builder.table(
        ["Thành phần", "Số ảnh", "Vai trò"],
        [
            ["XWOD train", "5.945", "Nguồn thời tiết bất lợi chính"],
            ["ACDC train", "1.572", "Bổ sung miền adverse conditions và kiểu suy giảm khác"],
            ["BDD replay", "2.000", "Giữ một phần tri thức miền giao thông thông thường"],
            ["Tập cơ sở", "9.517", "Composition trước oversampling"],
            ["Phase 2 v3 sau oversampling", "13.430", "Tăng xuất hiện ảnh chứa lớp hiếm"],
            ["Validation", "1.744", "XWOD val, không tham gia training"],
        ],
        widths=[2.1, 1.0, 3.2],
    )
    builder.paragraph(
        "Phase 2 v3 áp dụng oversampling theo sự hiện diện của lớp: bicycle ×2, motorcycle ×3 và bus ×3. Đây là lấy mẫu "
        "lại ở mức ảnh, không sinh hộp giả và không thay đổi nhãn gốc. Một ảnh có nhiều lớp hiếm có thể được chọn theo quy "
        "tắc xây dựng dataset, vì vậy số ảnh sau lấy mẫu phải được xác nhận từ manifest thực tế thay vì suy ra bằng phép "
        "nhân độc lập. Mục tiêu của chiến lược là tăng tần suất cập nhật gradient cho các lớp hiếm mà vẫn giữ dữ liệu gốc."
    )

    builder.heading2("3.4. Lựa chọn kiến trúc phát hiện đối tượng")
    builder.paragraph(
        "Benchmark ban đầu bao phủ ba họ kiến trúc để tránh kết luận chỉ dựa trên một dòng mô hình. YOLOv8n và YOLO11n "
        "đại diện cho detector một giai đoạn gọn nhẹ; Faster R-CNN đại diện cho detector hai giai đoạn; RT-DETR đại diện "
        "cho hướng Transformer thời gian thực. Việc lựa chọn biến thể cụ thể cân bằng giữa tính đại diện học thuật và khả "
        "năng huấn luyện trên tài nguyên sẵn có."
    )

    builder.heading3("3.4.1. Đặc điểm các mô hình khảo sát")
    builder.paragraph(
        "YOLOv8n sử dụng backbone/neck dựa trên các khối Conv, C2f và SPPF, head tách nhánh phân loại–hồi quy và cơ chế "
        "dự đoán anchor-free. YOLO11n tiếp tục hướng một giai đoạn với các khối mới nhằm cải thiện hiệu quả tham số. Hai "
        "mô hình nano phù hợp cho yêu cầu triển khai, đồng thời tạo cơ sở so sánh giữa hai thế hệ Ultralytics."
    )
    builder.paragraph(
        "Faster R-CNN với ResNet-50–FPN tạo proposal bằng RPN rồi phân loại và tinh chỉnh hộp ở giai đoạn thứ hai. Kiến "
        "trúc này thường có chi phí lớn hơn nhưng là đối chứng quan trọng cho giả thuyết rằng xử lý hai giai đoạn có thể "
        "hữu ích với đối tượng nhỏ. RT-DETR sử dụng hybrid encoder và decoder truy vấn để dự đoán tập đối tượng, cung cấp "
        "một đối chứng Transformer với độ chính xác cao nhưng chi phí triển khai lớn hơn các biến thể nano."
    )
    builder.caption("Bảng 3.6. Vai trò so sánh của bốn kiến trúc benchmark")
    builder.table(
        ["Mô hình", "Họ kiến trúc", "Ưu tiên khảo sát", "Hạn chế cần kiểm soát"],
        [
            ["YOLOv8n", "Một giai đoạn, anchor-free", "Cân bằng tốc độ, độ chính xác và deployability", "Năng lực mô hình nhỏ; nhạy với vật thể nhỏ/hiếm"],
            ["YOLO11n", "Một giai đoạn, anchor-free", "Hiệu quả tham số của thế hệ mới", "Khác biệt triển khai có thể ảnh hưởng so sánh"],
            ["Faster R-CNN", "Hai giai đoạn", "Đối chứng proposal-based", "Nặng, chậm; pipeline đánh giá khác Ultralytics"],
            ["RT-DETR", "Transformer detector", "Tham chiếu độ chính xác và recall", "Chi phí tính toán và triển khai cao hơn"],
        ],
        widths=[1.0, 1.45, 2.0, 2.0],
    )

    builder.heading3("3.4.2. Tiêu chí lựa chọn mô hình thực dụng")
    builder.paragraph(
        "Quyết định lựa chọn sử dụng bốn nhóm bằng chứng: mAP50-95 và mAP50, recall, thời gian suy luận trong cùng điều "
        "kiện, cùng độ phức tạp triển khai. Mô hình đạt accuracy tuyệt đối cao nhất được ghi nhận như tham chiếu accuracy; "
        "mô hình tiếp tục Phase 2 là mô hình đạt trade-off tốt nhất cho mục tiêu thực dụng. Do đó, việc RT-DETR đạt độ chính "
        "xác cao nhất ở benchmark ban đầu không mâu thuẫn với việc YOLOv8n được chọn làm practical model."
    )
    builder.paragraph(
        "Tốc độ chỉ được so sánh khi các phép đo dùng cùng phần cứng, batch size, chế độ precision và kích thước ảnh. Nếu "
        "điều kiện khác nhau hoặc log không đầy đủ, kết luận tốc độ được giới hạn ở mức định tính. Precision của Faster R-CNN "
        "không được suy diễn khi nguồn log không cung cấp; ô dữ liệu tương ứng phải để không khả dụng thay vì nội suy."
    )

    builder.heading2("3.5. Thiết kế quy trình huấn luyện nhiều giai đoạn")
    builder.paragraph(
        "Quy trình chính gồm một checkpoint pretrained và hai stage benchmark, sau đó là Phase 2 cho mô hình thực dụng. "
        "Mục đích của phân kỳ là thu hẹp khoảng cách miền theo thứ tự từ ảnh tổng quát, ảnh giao thông thông thường đến ảnh "
        "thời tiết bất lợi. Các run phải lưu riêng cấu hình, checkpoint tốt nhất và kết quả validation để tránh nhầm lẫn "
        "giữa stage hoặc dùng tên thư mục như bằng chứng thay cho metadata."
    )

    builder.heading3("3.5.1. Khởi tạo từ COCO và Stage 1 trên BDD100K")
    builder.paragraph(
        "Trọng số COCO cung cấp các đặc trưng thị giác tổng quát và giúp giảm chi phí huấn luyện từ đầu. Ở Stage 1, mô "
        "hình được fine-tune trên BDD30K train sau khi nhãn đã được ánh xạ về sáu lớp. Stage này đóng vai trò domain "
        "adaptation sang cảnh giao thông: cấu trúc đường, người tham gia giao thông, tỉ lệ đối tượng và nền đô thị. BDD val "
        "được dùng theo dõi hội tụ và lựa chọn checkpoint của Stage 1."
    )

    builder.heading3("3.5.2. Stage 2 trên XWOD và benchmark đa kiến trúc")
    builder.paragraph(
        "Stage 2 tiếp tục thích nghi mô hình với XWOD train, nơi ảnh chứa các kiểu suy giảm thời tiết bất lợi. XWOD val "
        "được dùng chọn checkpoint, trong khi XWOD test được giữ cho báo cáo hiệu năng cuối stage. Sau đó, cùng checkpoint "
        "được đánh giá trên DAWN val để đo khả năng chuyển miền. DAWN không được dùng để huấn luyện Stage 2, chọn epoch "
        "hay điều chỉnh trực tiếp siêu tham số."
    )
    builder.paragraph(
        "Bốn kiến trúc đi qua cùng logic dữ liệu và cùng bộ độ đo. Tuy nhiên, do Ultralytics và torchvision có pipeline "
        "huấn luyện khác nhau, tính công bằng được đảm bảo ở cấp dataset, split, class mapping và metric; những khác biệt "
        "không thể loại bỏ hoàn toàn ở optimizer hoặc implementation phải được công bố thay vì giả định là đồng nhất tuyệt đối."
    )

    builder.heading3("3.5.3. Phase 2 với dữ liệu gộp và BDD replay")
    builder.paragraph(
        "Sau benchmark, Phase 2 chỉ tiếp tục trên practical model. Dữ liệu gộp kết hợp XWOD train, ACDC train và 2.000 "
        "ảnh BDD replay; phiên bản v3 bổ sung oversampling lớp hiếm. Vai trò của XWOD là giữ miền đích chính, ACDC mở rộng "
        "đa dạng điều kiện bất lợi, còn BDD replay giảm nguy cơ quên hoàn toàn cảnh giao thông thông thường. Validation của "
        "Phase 2 vẫn là XWOD val để giữ tiêu chí chọn checkpoint nhất quán."
    )
    builder.paragraph(
        "Hai run final_yolov8n_phase2_v3_rare và final_yolov8n_phase2_v3_960 cùng sử dụng chiến lược/dataset Phase 2 v3 "
        "và khác chắc chắn ở imgsz 640 so với 960. Tên run không chứng minh run 960 được khởi tạo trực tiếp từ best.pt của "
        "run 640; quan hệ checkpoint chỉ được khẳng định khi args, log hoặc metadata ghi rõ. Vì bằng chứng lineage hiện chưa "
        "đủ, thiết kế và phần kết quả không mô tả hai run như một chuỗi fine-tune trực tiếp."
    )

    builder.heading3("3.5.4. Kiểm soát huấn luyện và chọn checkpoint")
    builder.paragraph(
        "Trong mỗi nhóm so sánh, dữ liệu train/val, class mapping, metric chính và quy tắc chọn checkpoint được giữ cố định. "
        "Seed, optimizer, learning rate schedule, batch size, AMP, patience và phiên bản thư viện phải được lấy từ cấu hình "
        "run, không phục dựng từ trí nhớ. Checkpoint best được lựa chọn theo kết quả validation; tập test không tham gia "
        "early stopping. Khi tăng imgsz hoặc đổi kích thước mô hình, batch size có thể phải thay đổi vì bộ nhớ, và khác biệt "
        "này phải được ghi trong metadata thực nghiệm."
    )

    builder.heading2("3.6. Thiết kế khảo sát cơ chế chú ý SE")
    builder.paragraph(
        "SE được triển khai như một ablation kiến trúc, không phải thành phần mặc định của mô hình cuối. Mục tiêu là kiểm "
        "tra liệu tái cân bằng kênh đặc trưng có giúp mô hình giảm ảnh hưởng nhiễu thời tiết hay không. Thiết kế này tách biệt "
        "với CBAM: CBAM chỉ được trình bày ở phần lý thuyết/hướng phát triển và không được mô tả như thí nghiệm đã thực hiện."
    )

    builder.heading3("3.6.1. Vị trí tích hợp và nguyên lý hoạt động")
    builder.paragraph(
        "Một khối SE được chèn sau SPPF của YOLOv8n, nơi feature map đã chứa thông tin ngữ nghĩa mức cao nhưng chưa đi qua "
        "neck đa tỉ lệ. Nhánh squeeze dùng global average pooling để biến mỗi kênh thành một thống kê; nhánh excitation "
        "dùng hai phép biến đổi có bottleneck và sigmoid để tạo trọng số kênh. Feature map đầu vào sau đó được nhân theo "
        "kênh với vector trọng số. Cách chèn một khối duy nhất giới hạn số biến thay đổi và giữ chi phí bổ sung nhỏ."
    )
    builder.picture(FIGURES / "hinh_3_3_se.png")
    builder.caption("Hình 3.3. Vị trí tích hợp mô-đun SE sau khối SPPF của YOLOv8n")

    builder.heading3("3.6.2. Thiết kế ablation và quy tắc ra quyết định")
    builder.paragraph(
        "Ablation so sánh YOLOv8n chuẩn với YOLOv8n+SE trên cùng dữ liệu và giao thức đánh giá. Biến độc lập là sự hiện "
        "diện của SE; các biến kiểm soát gồm split, class mapping, kích thước ảnh và quy tắc chọn checkpoint. Kết quả được "
        "đánh giá trên cả XWOD và DAWN để phân biệt cải thiện cùng miền với cải thiện chuyển miền."
    )
    builder.paragraph(
        "Quy tắc lựa chọn không dựa trên một chênh lệch mAP duy nhất. SE chỉ được giữ nếu cải thiện đủ nhất quán giữa các "
        "dataset và không gây suy giảm đáng kể recall. Nếu tác động trung tính hoặc trái chiều giữa XWOD và DAWN, kết luận "
        "phù hợp là SE chưa tạo cải thiện nhất quán, không phải SE hoàn toàn vô dụng. Kiến trúc checkpoint final phải được "
        "kiểm tra trực tiếp để xác nhận có hay không có SE, thay vì suy luận từ tên checkpoint tạm."
    )

    builder.heading2("3.7. Thiết kế thực nghiệm và giao thức đánh giá")
    builder.paragraph(
        "Các thực nghiệm được tổ chức theo thứ tự từ lựa chọn họ mô hình đến tối ưu dữ liệu, độ phân giải và năng lực. "
        "Mỗi thí nghiệm trả lời một câu hỏi riêng; kết quả âm được giữ lại vì giúp loại bỏ cấu hình không phù hợp. Bảng 3.7 "
        "tóm tắt biến thay đổi và tiêu chí của từng nhóm."
    )
    builder.caption("Bảng 3.7. Ma trận thiết kế các nhóm thực nghiệm")
    builder.table(
        ["Nhóm", "Cấu hình so sánh", "Biến chính", "Tập đánh giá", "Mục đích"],
        [
            ["Benchmark", "YOLOv8n, YOLO11n, Faster R-CNN, RT-DETR", "Kiến trúc", "XWOD test, DAWN val", "Chọn practical model và ghi nhận accuracy reference"],
            ["SE ablation", "YOLOv8n vs YOLOv8n+SE", "Chú ý kênh", "XWOD test, DAWN val", "Kiểm tra tính nhất quán giữa miền"],
            ["Phase 2", "Dữ liệu gộp và v3 oversampling", "Composition/lấy mẫu", "XWOD, DAWN, ACDC", "Cải thiện lớp hiếm và đa miền"],
            ["Độ phân giải", "v3 @640 vs v3 @960", "imgsz", "XWOD, DAWN, ACDC", "Đo lợi ích chi tiết không gian"],
            ["Accuracy ceiling", "YOLOv8n vs YOLOv8s @960", "Năng lực mô hình", "XWOD, DAWN, ACDC", "Ước lượng trần độ chính xác theo capacity"],
            ["Copy-paste", "v3 @960 vs copy-paste @960", "Tăng cường dữ liệu", "XWOD, DAWN, ACDC", "Kiểm tra lợi ích và nguy cơ artifact"],
        ],
        widths=[0.85, 1.55, 1.05, 1.35, 1.65],
    )

    builder.heading3("3.7.1. Benchmark bốn mô hình")
    builder.paragraph(
        "Benchmark sử dụng pipeline COCO pretrained → Stage 1 BDD100K → Stage 2 XWOD. Sau Stage 2, checkpoint được đánh "
        "giá trên XWOD test và DAWN val. Các chỉ số chính là P, R, mAP50 và mAP50-95; tốc độ và độ phức tạp được dùng như "
        "tiêu chí triển khai. Mục tiêu không phải tìm mô hình thắng trên mọi chỉ số mà tìm hai mốc tham chiếu: kiến trúc có "
        "accuracy tuyệt đối cao và kiến trúc có trade-off thực dụng để tiếp tục Phase 2."
    )

    builder.heading3("3.7.2. Đánh giá chiến lược Phase 2 và lớp hiếm")
    builder.paragraph(
        "Các cấu hình Phase 2 được so sánh với mốc trước mở rộng để xác định đóng góp của dữ liệu ACDC, BDD replay và "
        "oversampling. Ngoài metric trung bình, phân tích bắt buộc xem mAP50-95 của bicycle, motorcycle và bus, vì mục tiêu "
        "của v3 là tăng tín hiệu cho các lớp này. Đồng thời cần kiểm tra person, car và truck để phát hiện đánh đổi do thay "
        "đổi phân bố sampling."
    )
    builder.paragraph(
        "Đánh giá ACDC được đặt trong đúng bối cảnh: mốc trước Phase 2 là cross-dataset, còn mốc sau Phase 2 là held-out "
        "evaluation. Vì hai trạng thái không hoàn toàn cùng ý nghĩa, mức cải thiện phải được diễn giải như hiệu quả thích nghi "
        "sau khi dùng ACDC train, không phải bằng chứng zero-shot sau huấn luyện."
    )

    builder.heading3("3.7.3. Thực nghiệm độ phân giải đầu vào")
    builder.paragraph(
        "Hai run v3_rare @640 và v3_960 @960 dùng cùng Phase 2 v3 và cùng rare-class oversampling. Biến quan sát chính là "
        "kích thước ảnh. Ảnh 960 có khả năng bảo toàn nhiều chi tiết của đối tượng nhỏ nhưng làm tăng bộ nhớ, FLOPs và thời "
        "gian suy luận. Vì checkpoint initialization giữa hai run chưa được chứng minh, phân tích chỉ xem chúng là hai run "
        "cùng chiến lược dữ liệu với khác biệt chắc chắn về imgsz, không xem run 960 là phần tiếp nối trực tiếp của run 640."
    )

    builder.heading3("3.7.4. Thực nghiệm năng lực mô hình")
    builder.paragraph(
        "YOLOv8s Phase 2 v3 @960 được dùng như accuracy ceiling/reference, trong khi YOLOv8n v3 @960 là practical model. "
        "So sánh này đánh giá mức lợi ích khi tăng capacity trong cùng chiến lược dữ liệu và độ phân giải. Kết luận phải xem "
        "đồng thời mAP và recall theo dataset; nếu mô hình lớn tăng mAP nhưng recall ACDC gần như không đổi, cách diễn giải "
        "đúng là capacity cải thiện định vị/độ chính xác nhưng chưa giải quyết hoàn toàn nút thắt đối tượng nhỏ hoặc hiếm."
    )

    builder.heading3("3.7.5. Thực nghiệm copy-paste")
    builder.paragraph(
        "Copy-paste được thiết kế để tăng sự xuất hiện của đối tượng hiếm bằng cách ghép vùng đối tượng vào ảnh đích. Run "
        "được giữ cùng imgsz 960 và đánh giá trên ba tập để kiểm tra liệu lợi ích có tổng quát hay chỉ xuất hiện trên miền "
        "huấn luyện. Nếu hiệu năng giảm, kết quả âm vẫn được báo cáo. Các nguyên nhân như lệch ngữ cảnh, ánh sáng, tỉ lệ hoặc "
        "biên cắt chỉ được nêu dưới dạng giả thuyết, trừ khi có phân tích định lượng trực tiếp chứng minh."
    )

    builder.heading3("3.7.6. Độ đo, ngưỡng và nguyên tắc so sánh")
    builder.paragraph(
        "Precision đo tỉ lệ dự đoán đúng trong số dự đoán dương, recall đo tỉ lệ đối tượng thật được phát hiện, mAP50 dùng "
        "ngưỡng IoU 0,5 và mAP50-95 trung bình trên các ngưỡng 0,50:0,05:0,95. mAP50-95 được ưu tiên khi so sánh độ chính "
        "xác hộp, nhưng recall được theo dõi riêng vì bỏ sót người/phương tiện là rủi ro quan trọng trong giám sát giao thông."
    )
    builder.paragraph(
        "Mọi bảng phải ghi rõ dataset, split, số ảnh hợp lệ, số instances và imgsz. So sánh per-class chỉ có ý nghĩa khi "
        "thứ tự lớp khớp canonical và số mẫu đủ lớn. Ví dụ, DAWN val chỉ có sáu instances bicycle; metric của lớp này có "
        "độ bất định cao và không được diễn giải như bằng chứng ổn định chỉ từ một giá trị AP."
    )
    builder.caption("Bảng 3.8. Giao thức đánh giá và các lưu ý diễn giải")
    builder.table(
        ["Tập đánh giá", "Vai trò", "Thông tin bắt buộc", "Lưu ý"],
        [
            ["XWOD val", "Chọn checkpoint", "1.744 ảnh; metric validation", "Không thay thế cho XWOD test."],
            ["XWOD test", "Đánh giá cùng miền", "2.321 tệp nguồn; 2.320 ảnh hợp lệ", "Một tệp GIF gắn đuôi JPG bị loại."],
            ["DAWN val", "Đánh giá ngoài miền", "206 ảnh; 1.806 instances", "Không tham gia training; bicycle chỉ 6 instances."],
            ["ACDC val trước Phase 2", "Cross-dataset/zero-shot", "398 ảnh; 2.830 instances", "Mô hình chưa học ACDC train."],
            ["ACDC val sau Phase 2", "Held-out cùng nguồn", "398 ảnh; 2.830 instances", "Không gọi là pure zero-shot."],
        ],
        widths=[1.4, 1.5, 1.9, 1.7],
    )

    builder.heading2("3.8. Thiết kế tái lập và truy vết provenance")
    builder.paragraph(
        "Mỗi run được định danh bằng tên duy nhất và phải lưu tối thiểu args.yaml hoặc cấu hình tương đương, results.csv, "
        "checkpoint best/last và log đánh giá. Kết quả dùng trong báo cáo được ưu tiên theo thứ tự: log/checkpoint/dataset "
        "đã kiểm chứng, memory tổng hợp, rồi mới đến prose cũ trong báo cáo. Nếu ba nguồn mâu thuẫn, không chỉnh số chỉ để "
        "làm bảng trông nhất quán; phải ghi nhận conflict và dùng nguồn có provenance mạnh hơn."
    )

    builder.heading3("3.8.1. Quy tắc định danh cấu hình final")
    builder.paragraph(
        "Cấu hình practical final được định danh bởi kiến trúc, dataset strategy và imgsz: YOLOv8n chuẩn, Phase 2 v3 rare "
        "oversampling, imgsz 960. Việc xác nhận kiến trúc dựa trên kiểm tra trực tiếp checkpoint, không dựa vào tên tệp tạm; "
        "checkpoint final không chứa SE. YOLOv8s Phase 2 v3 @960 được định danh riêng là accuracy ceiling, không thay thế "
        "mô hình thực dụng."
    )

    builder.heading3("3.8.2. Các giới hạn provenance và cách báo cáo")
    builder.paragraph(
        "Hai giới hạn còn mở không làm thay đổi thiết kế: log text đầy đủ của benchmark bốn mô hình chưa đồng đều như log "
        "tái đánh giá final, và lineage chính xác giữa v3_rare @640 với v3_960 @960 chưa được chứng minh. Vì vậy, báo cáo "
        "giữ lại các số benchmark đã xuất hiện nhất quán trong nguồn tổng hợp nhưng không nâng mức khẳng định vượt quá bằng "
        "chứng; precision Faster R-CNN để ở trạng thái không khả dụng nếu không có log gốc."
    )
    builder.paragraph(
        "Thiết kế ở Chương 3 tạo khung để Chương 4 trình bày kết quả theo đúng câu hỏi nghiên cứu: mô hình nào có accuracy "
        "cao nhất, mô hình nào phù hợp triển khai, SE có cải thiện nhất quán hay không, Phase 2 v3 và imgsz 960 tác động ra "
        "sao, và liệu tăng capacity hoặc copy-paste có giải quyết được nút thắt còn lại. Cách tổ chức này bảo đảm kết luận "
        "cuối cùng xuất phát từ bằng chứng thực nghiệm thay vì tên run hoặc giả định thiết kế."
    )


def update_manual_figure_table_lists(document: Document) -> None:
    """Refresh the existing manual figure/table lists without touching covers."""
    figure_titles = [
        p.text.strip() for p in document.paragraphs
        if p.style.name == "Caption2" and p.text.strip().startswith("Hình ")
    ]
    table_titles = [
        p.text.strip() for p in document.paragraphs
        if p.style.name == "Caption2" and p.text.strip().startswith("Bảng ")
    ]
    paragraphs = document.paragraphs
    figure_heading = next(p for p in paragraphs if p.text.strip() == "DANH MỤC HÌNH")
    table_heading = next(p for p in paragraphs if p.text.strip() == "DANH MỤC BẢNG")
    chapter1 = next(p for p in paragraphs if p.style.name == "Heading 1" and p.text.startswith("CHƯƠNG 1"))

    def clear_between(start, end):
        current = start._p.getnext()
        while current is not None and current is not end._p:
            following = current.getnext()
            current.getparent().remove(current)
            current = following

    clear_between(figure_heading, table_heading)
    clear_between(table_heading, chapter1)
    for title in figure_titles:
        p = document.add_paragraph(title, style="Normal")
        table_heading._p.addprevious(p._p)
    for title in table_titles:
        p = document.add_paragraph(title, style="Normal")
        chapter1._p.addprevious(p._p)


def mark_fields_dirty(document: Document) -> None:
    settings = document.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")
    for fld_char in document.element.xpath(".//w:fldChar"):
        if fld_char.get(qn("w:fldCharType")) == "begin":
            fld_char.set(qn("w:dirty"), "true")


def main(path_str: str) -> None:
    path = Path(path_str)
    document = Document(path)
    _chapter3_heading, chapter4_heading = remove_chapter3_body(document)
    builder = ChapterBuilder(document, chapter4_heading)
    add_chapter3(builder)
    builder.commit()
    update_manual_figure_table_lists(document)
    mark_fields_dirty(document)
    document.save(path)
    print(f"Đã ghi lại Chương 3: {path}")


if __name__ == "__main__":
    main(sys.argv[1])
