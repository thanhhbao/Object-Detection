#!/usr/bin/env python3
"""Produce a more cohesive, natural academic revision of thesis Chapter 3."""

from pathlib import Path
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from rewrite_thesis_chapter3 import (
    ChapterBuilder,
    FIGURES,
    mark_fields_dirty,
    remove_chapter3_body,
    update_manual_figure_table_lists,
)


def add_chapter3(builder: ChapterBuilder) -> None:
    p = builder.paragraph
    h2 = builder.heading2
    h3 = builder.heading3
    cap = builder.caption
    table = builder.table

    p(
        "Chương 2 đã trình bày cơ sở lý thuyết của bài toán phát hiện đối tượng, đặc điểm của các họ mô hình và những "
        "khó khăn nảy sinh khi chất lượng ảnh suy giảm do thời tiết. Từ nền tảng đó, chương này xây dựng phương án giải "
        "quyết cụ thể cho đề tài. Trọng tâm của chương không nằm ở việc mô tả lại từng thuật toán, mà ở cách tổ chức dữ "
        "liệu, lựa chọn mô hình và thiết kế thực nghiệm sao cho các kết quả ở Chương 4 có thể được kiểm chứng và diễn giải "
        "một cách nhất quán."
    )
    p(
        "Quá trình thiết kế được đặt trong hai ràng buộc có quan hệ trực tiếp với nhau. Mô hình phải đủ nhạy để phát hiện "
        "người và phương tiện trong ảnh bị giảm độ tương phản, che khuất hoặc nhiễu; đồng thời, mô hình cuối phải đủ gọn để "
        "có khả năng triển khai trong một hệ thống giám sát thực tế. Vì vậy, độ chính xác cao nhất không phải là tiêu chí "
        "duy nhất. Kích thước mô hình, tốc độ suy luận, recall trên dữ liệu ngoài miền và mức độ ổn định giữa các tập đánh "
        "giá đều được đưa vào quá trình lựa chọn."
    )

    h2("3.1. Phân tích bài toán và tiêu chí thiết kế")
    h3("3.1.1. Đặc tả bài toán và phạm vi nghiên cứu")
    p(
        "Đầu vào của hệ thống là ảnh tĩnh hoặc khung hình trích từ video giao thông. Ảnh có thể được ghi nhận trong điều "
        "kiện bình thường hoặc dưới các hiện tượng bất lợi như mưa, sương mù, tuyết, ngập và bão cát. Đầu ra là tập các "
        "hộp giới hạn, trong đó mỗi hộp gồm vị trí đối tượng, nhãn lớp và độ tin cậy. Không gian nhãn được giới hạn ở sáu "
        "lớp thường gặp trong giao thông đường bộ: person, bicycle, car, motorcycle, bus và truck."
    )
    p(
        "So với ảnh giao thông thông thường, dữ liệu thời tiết bất lợi gây khó khăn ở cả mức điểm ảnh lẫn mức phân phối. "
        "Sương mù và mưa làm suy giảm tương phản; tuyết hoặc ánh sáng phản xạ làm thay đổi màu sắc và biên vật thể; trong "
        "khi ngập nước hoặc bão cát có thể che khuất một phần đối tượng. Các lớp bicycle, motorcycle và bus còn xuất hiện "
        "ít hơn đáng kể so với car và person. Nếu chỉ tối ưu chỉ số trung bình, mô hình có thể đạt kết quả tương đối tốt "
        "nhưng vẫn bỏ sót nhiều đối tượng hiếm. Thiết kế thực nghiệm vì thế phải theo dõi cả kết quả tổng thể lẫn kết quả "
        "theo từng lớp."
    )
    p(
        "Phạm vi đề tài dừng ở bài toán phát hiện đối tượng. Các chức năng theo vết, nhận dạng biển số, ước lượng khoảng "
        "cách và dự báo quỹ đạo không được đưa vào mô hình. Việc giới hạn phạm vi giúp tập trung tài nguyên vào hai câu hỏi "
        "chính: kiến trúc nào tạo được sự cân bằng hợp lý giữa độ chính xác và khả năng triển khai; chiến lược dữ liệu nào "
        "giúp mô hình thích nghi tốt hơn với điều kiện thời tiết bất lợi."
    )

    h3("3.1.2. Yêu cầu đối với hệ thống")
    p(
        "Yêu cầu của hệ thống được xác định cho cả phần nghiên cứu và phần ứng dụng. Ở phần nghiên cứu, quy trình phải quản "
        "lý được nhiều nguồn dữ liệu, nhiều kiến trúc và nhiều giai đoạn huấn luyện mà không làm thay đổi ý nghĩa của nhãn. "
        "Ở phần ứng dụng, bộ trọng số đã chọn phải nhận được ảnh hoặc video, thực hiện suy luận và trả về kết quả có thể "
        "quan sát hoặc lưu trữ. Bảng 3.1 tổng hợp các yêu cầu chính và cách kiểm tra tương ứng."
    )
    cap("Bảng 3.1. Yêu cầu và tiêu chí kiểm chứng của hệ thống")
    table(
        ["Mã", "Yêu cầu", "Nội dung", "Tiêu chí kiểm chứng"],
        [
            ["YC1", "Chuẩn hóa dữ liệu", "Đưa nhãn từ các bộ dữ liệu về cùng sáu lớp mục tiêu.", "Class ID trong dữ liệu, mô hình và báo cáo có cùng ý nghĩa."],
            ["YC2", "Huấn luyện nhiều giai đoạn", "Thực hiện Stage 1, Stage 2 và Phase 2 theo cấu hình độc lập.", "Mỗi lần huấn luyện có cấu hình, trọng số và nhật ký riêng."],
            ["YC3", "Đánh giá đa miền", "Đánh giá trên XWOD, DAWN và ACDC theo đúng vai trò của từng phần dữ liệu.", "Bảng kết quả ghi rõ bộ dữ liệu, split, số ảnh và kích thước đầu vào."],
            ["YC4", "Phát hiện và trực quan hóa", "Sinh hộp giới hạn, nhãn, độ tin cậy và thống kê theo lớp.", "Kết quả trực quan khớp với đầu ra số của mô hình."],
            ["YC5", "Hiệu quả triển khai", "Theo dõi tốc độ, kích thước và chi phí tính toán cùng với độ chính xác.", "Các phép đo tốc độ được thực hiện trong cùng điều kiện phần cứng."],
            ["YC6", "Khả năng tái lập", "Lưu tham số huấn luyện, trọng số và log đánh giá.", "Mỗi số liệu trong báo cáo có thể truy về nguồn thực nghiệm."],
        ],
        widths=[0.55, 1.4, 2.2, 2.3],
    )
    p(
        "Độ chính xác được đánh giá bằng Precision, Recall, mAP50 và mAP50-95. Trong đó, mAP50-95 được sử dụng để phản ánh "
        "chất lượng định vị hộp ở nhiều ngưỡng IoU, còn Recall cho biết mức độ bỏ sót đối tượng. Đối với bài toán giám sát, "
        "recall có ý nghĩa riêng vì một mô hình ít dự đoán sai nhưng bỏ sót nhiều người hoặc phương tiện vẫn chưa đáp ứng tốt "
        "yêu cầu thực tế. Tốc độ suy luận được báo cáo kèm môi trường đo; không sử dụng một ngưỡng FPS tuyệt đối cho mọi loại "
        "GPU và mọi kích thước ảnh."
    )

    h2("3.2. Kiến trúc hệ thống và luồng xử lý")
    h3("3.2.1. Kiến trúc tổng thể")
    p(
        "Kiến trúc hệ thống được chia thành bốn tầng: dữ liệu, mô hình, huấn luyện–đánh giá và ứng dụng. Cách phân chia này "
        "phản ánh đúng luồng phụ thuộc của bài toán. Dữ liệu sau chuẩn hóa là đầu vào của quá trình huấn luyện; huấn luyện tạo "
        "ra bộ trọng số; bộ trọng số đã khóa mới được chuyển sang đánh giá và triển khai. Nhờ đó, giao diện ứng dụng không can "
        "thiệp vào quy trình nghiên cứu, còn việc thay một mô hình không đòi hỏi xây dựng lại toàn bộ phần xử lý dữ liệu."
    )
    p(
        "Tầng dữ liệu đảm nhiệm việc đọc ảnh, chuyển đổi nhãn, kiểm tra tính hợp lệ và tổ chức các tập train, val, test. "
        "Tầng mô hình đóng gói bốn kiến trúc so sánh và biến thể SE để chúng có thể sử dụng cùng không gian sáu lớp. Tầng "
        "huấn luyện–đánh giá điều phối các giai đoạn tinh chỉnh, lưu bộ trọng số tốt nhất và tạo kết quả tổng thể cũng như theo "
        "lớp. Tầng ứng dụng tiếp nhận mô hình cuối cùng phục vụ triển khai để suy luận trên ảnh hoặc video, sau đó vẽ hộp và xuất thống kê."
    )
    builder.picture(FIGURES / "hinh_3_1_kien_truc.png")
    cap("Hình 3.1. Kiến trúc tổng thể của hệ thống")
    p(
        "Giữa các tầng tồn tại một nguyên tắc kiểm soát quan trọng: tập test không được dùng để chọn epoch hoặc điều chỉnh "
        "trực tiếp siêu tham số. Checkpoint được chọn trên tập validation của giai đoạn tương ứng, sau đó mới được đánh giá "
        "trên XWOD test và các tập ngoài miền. Nguyên tắc này làm giảm nguy cơ mô hình bị tối ưu gián tiếp theo tập kiểm thử."
    )

    h3("3.2.2. Luồng xử lý từ dữ liệu đến kết quả")
    p(
        "Mỗi bộ dữ liệu trước hết được kiểm tra khả năng giải mã ảnh và sự tồn tại của nhãn. Nhãn gốc được ánh xạ về sáu "
        "lớp, các lớp ngoài phạm vi bị loại bỏ, sau đó tọa độ hộp được chuẩn hóa theo định dạng của mô hình. Các kiểm tra biên "
        "được thực hiện để loại hộp có diện tích bằng không, tọa độ đảo hoặc nằm hoàn toàn ngoài ảnh. Sau bước này, thống kê "
        "số ảnh và số đối tượng theo lớp được dùng để phát hiện sai lệch trong quá trình chuyển đổi."
    )
    p(
        "Dữ liệu hợp lệ được đưa vào quy trình huấn luyện theo từng giai đoạn. Mỗi lần huấn luyện tạo một thư mục riêng chứa cấu hình, "
        "kết quả theo epoch và bộ trọng số. Khi huấn luyện kết thúc, bộ trọng số tốt nhất được đánh giá bằng đúng kích thước "
        "ảnh và split đã quy định. Kết quả đánh giá được lưu cùng tên bộ trọng số thay vì chỉ sao chép thủ công vào bảng tổng "
        "hợp. Cách tổ chức này đặc biệt cần thiết khi nhiều lần huấn luyện có tên gần giống nhau nhưng khác về kích thước ảnh hoặc chiến "
        "lược dữ liệu."
    )
    builder.picture(FIGURES / "hinh_3_2_pipeline.png")
    cap("Hình 3.2. Luồng huấn luyện và đánh giá của đề tài")
    p(
        "Ở giai đoạn triển khai, ảnh đầu vào được letterbox về kích thước mà bộ trọng số yêu cầu, sau đó đi qua mô hình và "
        "bước hậu xử lý. Các dự đoán dưới ngưỡng tin cậy bị loại; những hộp còn lại được gán màu theo lớp và chuyển về tọa "
        "độ của ảnh gốc. Đối với video, quy trình được lặp trên từng khung hình. Phần triển khai không làm thay đổi trọng số "
        "và do đó không ảnh hưởng đến kết quả đánh giá đã công bố."
    )

    h2("3.3. Thiết kế và chuẩn hóa dữ liệu")
    h3("3.3.1. Vai trò của từng bộ dữ liệu")
    p(
        "Bốn bộ dữ liệu không được sử dụng như bốn nguồn tương đương. BDD100K có quy mô lớn và phản ánh cảnh giao thông đa "
        "dạng, nhưng phần lớn ảnh không tập trung vào thời tiết bất lợi; vì vậy bộ này phù hợp cho bước thích nghi miền giao "
        "thông. XWOD là nguồn dữ liệu thời tiết bất lợi chính của Stage 2. DAWN có quy mô nhỏ hơn và được giữ bên ngoài quá "
        "trình huấn luyện để đo khả năng chuyển miền. ACDC được đưa vào Phase 2 nhằm mở rộng kiểu suy giảm ảnh và điều kiện "
        "môi trường mà mô hình được quan sát."
    )
    cap("Bảng 3.2. Vai trò và quy mô dữ liệu được sử dụng")
    table(
        ["Bộ dữ liệu", "Quy mô", "Vai trò", "Tập được giữ lại để đánh giá"],
        [
            ["BDD100K", "BDD30K train khoảng 30.000 ảnh; val khoảng 10.000 ảnh", "Stage 1 và nguồn BDD replay", "BDD val"],
            ["XWOD", "Train 5.945; val 1.744; test 2.321 tệp nguồn", "Dữ liệu thời tiết bất lợi chính", "XWOD val và XWOD test"],
            ["DAWN", "206 ảnh; 1.806 đối tượng", "Đánh giá chuyển miền", "Toàn bộ DAWN val; không dùng để huấn luyện"],
            ["ACDC", "Train 1.572; val 398", "Mở rộng dữ liệu trong Phase 2", "ACDC val"],
        ],
        widths=[1.05, 1.9, 1.9, 1.65],
    )
    p(
        "Cách diễn giải ACDC val phụ thuộc vào thời điểm đánh giá. Trước Phase 2, mô hình chưa học từ ACDC train nên ACDC "
        "val là phép đánh giá chéo dữ liệu, có thể xem là zero-shot theo nguồn dữ liệu. Sau khi ACDC train được đưa vào Phase "
        "2, ACDC val trở thành tập giữ lại cùng nguồn. Kết quả sau Phase 2 vì vậy cho thấy hiệu quả thích nghi với ACDC, "
        "không còn là bằng chứng zero-shot thuần túy. DAWN không thay đổi vai trò: bộ này không tham gia bất kỳ giai đoạn "
        "huấn luyện nào và tiếp tục là phép đo ngoài miền."
    )
    p(
        "XWOD test có 2.321 tệp ảnh ở nguồn. Trong quá trình đánh giá, tệp flooding_test_00217.jpg bị thư viện loại vì nội "
        "dung thực tế được nhận diện là GIF89a thay vì ảnh JPEG hợp lệ. Do đó, quá trình đánh giá xử lý 2.320 ảnh. Việc ghi đồng "
        "thời số tệp nguồn và số ảnh hợp lệ giúp giải thích chênh lệch mà không làm thay đổi dữ liệu chỉ để khớp số lượng."
    )

    h3("3.3.2. Chuẩn hóa nhãn và ảnh đầu vào")
    p(
        "Không gian nhãn chung được xem như một hợp đồng giữa dữ liệu, mô hình và báo cáo. Thứ tự bắt buộc là 0 person, "
        "1 bicycle, 2 car, 3 motorcycle, 4 bus và 5 truck. Thứ tự này được khai báo trong tệp cấu hình dữ liệu, dùng khi "
        "ghi nhãn YOLO và giữ nguyên khi tổng hợp kết quả theo lớp. Đặc biệt, car đứng trước motorcycle; đảo hai lớp này sẽ "
        "không làm chương trình báo lỗi nhưng khiến toàn bộ bảng phân tích theo lớp sai ý nghĩa."
    )
    cap("Bảng 3.3. Ánh xạ nhãn về không gian sáu lớp")
    table(
        ["ID", "Lớp đích", "BDD100K", "XWOD", "ACDC/Cityscapes"],
        [
            ["0", "person", "person, rider", "person", "person (24), rider (25)"],
            ["1", "bicycle", "bike", "bicycle", "bicycle (33)"],
            ["2", "car", "car", "car", "car (26)"],
            ["3", "motorcycle", "motor", "motorcycle", "motorcycle (32)"],
            ["4", "bus", "bus", "bus", "bus (28)"],
            ["5", "truck", "truck", "truck", "truck (27)"],
            ["—", "Ngoài phạm vi", "traffic light, traffic sign, train", "Các lớp khác", "train (31), lớp nền"],
        ],
        widths=[0.45, 1.05, 1.65, 1.2, 2.0],
    )
    p(
        "BDD100K và XWOD cung cấp hộp giới hạn nên có thể chuyển trực tiếp về biểu diễn YOLO gồm class_id, tọa độ tâm và "
        "kích thước hộp đã chuẩn hóa. ACDC cung cấp nhãn phân đoạn; mỗi instance hợp lệ được tách theo mã lớp rồi bao bởi "
        "hình chữ nhật nhỏ nhất. Các vùng nền, lớp ngoài phạm vi và hộp suy biến bị loại trước khi ghi tệp nhãn. Với person "
        "và rider, việc gộp nhãn được thực hiện nhất quán để phù hợp với mục tiêu phát hiện người tham gia giao thông."
    )
    p(
        "Ảnh được thay đổi kích thước bằng letterbox. Phép biến đổi giữ nguyên tỉ lệ khung hình, thu phóng ảnh theo cạnh giới "
        "hạn rồi đệm phần còn thiếu. Nhãn hộp được biến đổi theo cùng hệ số và độ lệch đệm. Cách làm này tránh hiện tượng "
        "kéo giãn người hoặc phương tiện, vốn có thể làm sai đặc trưng hình dạng. Kích thước 640 được dùng ở benchmark và "
        "các run cơ sở; kích thước 960 là biến khảo sát ở giai đoạn sau, không phải cấu hình chung cho mọi thực nghiệm."
    )

    h3("3.3.3. Thành phần dữ liệu Phase 2 và xử lý lớp hiếm")
    p(
        "Phase 2 được xây dựng sau khi đã chọn YOLOv8n làm mô hình thực dụng. Tập cơ sở của phiên bản v2 và v3 gồm 5.945 "
        "ảnh XWOD train, 1.572 ảnh ACDC train và 2.000 ảnh BDD replay, tổng cộng 9.517 ảnh. XWOD giữ vai trò miền đích chính; "
        "ACDC bổ sung các kiểu suy giảm và bối cảnh khác; phần BDD replay giúp duy trì một lượng tri thức về cảnh giao thông "
        "thông thường khi mô hình tiếp tục học trên dữ liệu bất lợi."
    )
    cap("Bảng 3.4. Thành phần dữ liệu Phase 2 v3")
    table(
        ["Thành phần", "Số ảnh", "Mục đích sử dụng"],
        [
            ["XWOD train", "5.945", "Duy trì miền thời tiết bất lợi chính"],
            ["ACDC train", "1.572", "Mở rộng điều kiện và dạng suy giảm ảnh"],
            ["BDD replay", "2.000", "Hạn chế quên miền giao thông thông thường"],
            ["Tập cơ sở", "9.517", "Dữ liệu trước khi lấy mẫu tăng cường"],
            ["Phase 2 v3", "13.430", "Tập train sau khi tăng mẫu cho lớp hiếm"],
            ["Validation", "1.744", "XWOD val, không tham gia cập nhật trọng số"],
        ],
        widths=[1.8, 1.0, 3.4],
    )
    p(
        "Số 2.000 là số ảnh BDD replay thực tế trong tập dữ liệu đã xây dựng. Con số này tương đương khoảng 6,7% của BDD30K, "
        "không phải 30%. Trường bdd_replay_ratio = 0.3 còn xuất hiện trong một số metadata cũ được xem là giá trị mặc định "
        "không phản ánh composition cuối cùng. Báo cáo sử dụng số tuyệt đối 2.000 ảnh để tránh biến một trường cấu hình lỗi "
        "thời thành đặc điểm của dữ liệu thực nghiệm."
    )
    p(
        "Phiên bản v3 tăng tần suất xuất hiện của ảnh chứa bicycle lên hai lần, motorcycle lên ba lần và bus lên ba lần. "
        "Đây là oversampling ở mức ảnh; ảnh và nhãn gốc không bị biến đổi, cũng không có hộp giả được tạo thêm. Vì một ảnh "
        "có thể chứa nhiều lớp, quy mô 13.430 ảnh được lấy từ manifest sau khi build thay vì tính bằng cách cộng riêng số "
        "mẫu của từng lớp. Cách lấy mẫu này ưu tiên thêm cơ hội cập nhật cho lớp hiếm nhưng vẫn giữ được bối cảnh thật của ảnh."
    )

    h2("3.4. Lựa chọn mô hình và thiết kế quy trình huấn luyện")
    h3("3.4.1. Các kiến trúc được đưa vào benchmark")
    p(
        "Bốn mô hình được lựa chọn để đại diện cho ba hướng phát hiện đối tượng. YOLOv8n và YOLO11n là mô hình một giai "
        "đoạn, có kích thước nhỏ và phù hợp với yêu cầu suy luận nhanh. Faster R-CNN đại diện cho kiến trúc hai giai đoạn, "
        "trong đó mạng đề xuất vùng tạo ứng viên trước khi phân loại và hiệu chỉnh hộp. RT-DETR đại diện cho detector dựa "
        "trên Transformer, sử dụng truy vấn đối tượng và cơ chế dự đoán theo tập."
    )
    p(
        "Việc so sánh không nhằm chứng minh một họ kiến trúc luôn vượt trội, mà nhằm quan sát cách các đặc tính khác nhau biểu "
        "hiện trên dữ liệu thời tiết bất lợi. Detector hai giai đoạn có thể xử lý vùng ứng viên kỹ hơn nhưng tốn nhiều thời "
        "gian. RT-DETR có năng lực biểu diễn mạnh nhưng chi phí tính toán cao. Hai biến thể YOLO nano có lợi thế rõ về kích "
        "thước và hệ sinh thái triển khai, song có thể gặp giới hạn với vật thể nhỏ hoặc bị che khuất."
    )
    cap("Bảng 3.5. Đặc điểm và vai trò của các mô hình benchmark")
    table(
        ["Mô hình", "Hướng tiếp cận", "Vai trò trong so sánh", "Đánh đổi chính"],
        [
            ["YOLOv8n", "Một giai đoạn, anchor-free", "Ứng viên mô hình thực dụng", "Gọn và nhanh; năng lực biểu diễn giới hạn hơn mô hình lớn"],
            ["YOLO11n", "Một giai đoạn, anchor-free", "Đối chiếu thế hệ YOLO mới", "Hiệu quả tham số cao; kết quả vẫn phụ thuộc miền dữ liệu"],
            ["Faster R-CNN", "Hai giai đoạn", "Đối chứng proposal-based", "Chi phí suy luận và bộ nhớ lớn"],
            ["RT-DETR", "Transformer detector", "Mốc tham chiếu độ chính xác", "Độ chính xác cao nhưng nặng hơn các mô hình nano"],
        ],
        widths=[1.0, 1.5, 1.8, 2.1],
    )
    p(
        "Tiêu chí lựa chọn được tách thành hai khái niệm. Mô hình có độ chính xác cao nhất trong benchmark được dùng làm mốc "
        "tham chiếu về độ chính xác. Mô hình thực dụng được chọn theo sự cân bằng giữa mAP, recall, tốc độ, kích thước và khả năng "
        "tích hợp vào ứng dụng. Do đó, RT-DETR có thể đứng đầu về độ chính xác trong benchmark ban đầu, trong khi YOLOv8n "
        "được chọn để phát triển tiếp ở Phase 2. Hai kết luận này trả lời hai tiêu chí khác nhau và không mâu thuẫn."
    )

    h3("3.4.2. Quy trình học chuyển giao nhiều giai đoạn")
    p(
        "Các mô hình không được huấn luyện từ đầu. Bộ trọng số đã huấn luyện trước trên COCO cung cấp đặc trưng tổng quát về hình dạng, "
        "biên và cấu trúc đối tượng. Stage 1 tiếp tục tinh chỉnh trên BDD30K để đưa mô hình về miền giao thông đường phố. "
        "BDD có quy mô đủ lớn để mô hình học lại phân bố sáu lớp mục tiêu mà chưa phải đối diện ngay với dữ liệu thời tiết "
        "bất lợi có quy mô nhỏ và mất cân bằng."
    )
    p(
        "Stage 2 sử dụng XWOD train để thích nghi với miền thời tiết bất lợi. XWOD val phục vụ theo dõi hội tụ và chọn "
        "bộ trọng số, còn XWOD test chỉ được dùng sau khi cấu hình đã được xác định. Bộ trọng số Stage 2 sau đó được đánh giá "
        "trên DAWN val. Vì DAWN không tham gia huấn luyện hoặc chọn epoch, kết quả trên bộ này phản ánh khả năng chuyển miền "
        "tốt hơn so với kết quả trên XWOD."
    )
    p(
        "Sau benchmark, YOLOv8n được đưa vào Phase 2 với dữ liệu gộp. Hai lần huấn luyện final_yolov8n_phase2_v3_rare và "
        "final_yolov8n_phase2_v3_960 cùng sử dụng dữ liệu và chiến lược lấy mẫu của Phase 2 v3; khác biệt chắc chắn được "
        "xác nhận là imgsz 640 và 960. Tuy nhiên, các bằng chứng hiện có chưa chỉ ra rõ bộ trọng số khởi tạo của lần huấn luyện 960. "
        "Vì vậy, luận văn không mô tả lần huấn luyện 960 là bước tinh chỉnh trực tiếp từ best.pt của lần huấn luyện 640."
    )

    h3("3.4.3. Kiểm soát quá trình huấn luyện")
    p(
        "Một phép so sánh chỉ có giá trị khi biến cần khảo sát được tách khỏi các yếu tố còn lại. Trong benchmark, tập dữ liệu, "
        "split, cách ánh xạ lớp và bộ độ đo được giữ thống nhất. Các mô hình sử dụng chung logic huấn luyện nhiều giai đoạn, "
        "nhưng không giả định mọi chi tiết triển khai của Ultralytics và torchvision hoàn toàn giống nhau. Những khác biệt "
        "về bộ tối ưu, cách tính hàm mất mát hoặc hậu xử lý được xem là một phần của cách cài đặt và phải được công bố khi diễn giải."
    )
    p(
        "Trong từng run, checkpoint tốt nhất được chọn theo validation thay vì epoch cuối. Seed, batch size, learning rate, "
        "optimizer, patience, AMP và phiên bản thư viện được lấy từ tệp cấu hình thực tế. Khi tăng kích thước ảnh từ 640 lên "
        "960 hoặc chuyển từ YOLOv8n sang YOLOv8s, batch size có thể phải thay đổi do giới hạn bộ nhớ. Những thay đổi như vậy "
        "không bị che giấu dưới nhận định 'cùng thiết lập', mà được ghi nhận như điều kiện của thí nghiệm."
    )
    p(
        "Tốc độ suy luận cũng được kiểm soát theo cùng nguyên tắc. Giá trị ms/ảnh hoặc FPS chỉ được đặt cạnh nhau khi được đo "
        "trên cùng phần cứng, cùng chế độ precision, batch size và kích thước đầu vào. Nếu nguồn log không thỏa điều kiện đó, "
        "tốc độ chỉ được dùng để mô tả xu hướng. Tương tự, Precision của Faster R-CNN không được nội suy từ các độ đo khác khi "
        "log gốc không cung cấp giá trị này."
    )

    h2("3.5. Thiết kế các hướng cải tiến")
    h3("3.5.1. Khảo sát cơ chế chú ý SE")
    p(
        "SE được khảo sát như một phép loại trừ thành phần kiến trúc. Một khối SE được đặt sau SPPF của YOLOv8n, tại vị trí "
        "feature map đã tích lũy đặc trưng ngữ nghĩa mức cao trước khi chuyển sang phần tổng hợp đa tỉ lệ. Nhánh squeeze dùng "
        "global average pooling để tóm tắt đáp ứng của từng kênh. Nhánh excitation tạo trọng số qua hai phép biến đổi và hàm "
        "sigmoid; các trọng số này được nhân trở lại feature map để tái cân bằng kênh."
    )
    builder.picture(FIGURES / "hinh_3_3_se.png")
    cap("Hình 3.3. Vị trí tích hợp mô-đun SE trong YOLOv8n")
    p(
        "Việc chỉ chèn một khối giúp giới hạn số biến thay đổi và làm rõ tác động của SE. Mô hình chuẩn và mô hình có SE được "
        "đánh giá trên cả XWOD lẫn DAWN. Nếu SE chỉ tăng nhẹ chỉ số cùng miền nhưng làm giảm recall hoặc kết quả ngoài miền, "
        "cải thiện đó không đủ để đưa mô-đun vào cấu hình cuối. Ngược lại, kết quả âm cũng không được diễn giải thành việc "
        "SE không có giá trị trong mọi trường hợp; kết luận chỉ giới hạn ở cấu hình và dữ liệu của đề tài."
    )
    p(
        "Bộ trọng số cuối cùng được kiểm tra trực tiếp ở cấp kiến trúc. Các lớp của mô hình gồm Conv, C2f, Bottleneck, SPPF, "
        "Upsample, Concat, Detect và DFL; không có SEBlock hoặc lớp chú ý tương đương. Vì vậy, SE được ghi nhận là phép loại trừ thành phần "
        "đã thực hiện nhưng không được chọn. CBAM không được triển khai trong chuỗi thực nghiệm và chỉ có thể xuất hiện như "
        "một hướng nghiên cứu tiếp theo."
    )

    h3("3.5.2. Cải tiến dữ liệu, độ phân giải và năng lực mô hình")
    p(
        "Sau ablation kiến trúc, các hướng cải tiến tập trung vào ba nút thắt: thiếu mẫu lớp hiếm, thiếu chi tiết không gian và "
        "giới hạn năng lực biểu diễn. Phase 2 v3 giải quyết nút thắt thứ nhất bằng cách kết hợp thêm ACDC, BDD replay và "
        "oversampling. Run v3_rare ở kích thước 640 được dùng để quan sát tác động của chiến lược dữ liệu trước khi xem xét "
        "độ phân giải cao hơn."
    )
    p(
        "Run v3_960 dùng cùng chiến lược Phase 2 v3 nhưng tăng ảnh đầu vào lên 960. Kích thước lớn hơn giữ lại nhiều chi tiết "
        "của người và phương tiện nhỏ, đồng thời làm tăng chi phí tính toán. Thực nghiệm này vì thế phải xem xét cả mAP và "
        "tốc độ. Quan hệ checkpoint giữa run 640 và 960 chưa được xác minh nên khác biệt kết quả được mô tả như so sánh giữa "
        "hai run cùng dữ liệu, không phải mức tăng chắc chắn do một bước fine-tune nối tiếp."
    )
    p(
        "YOLOv8s Phase 2 v3 ở kích thước 960 được dùng để ước lượng trần độ chính xác khi tăng năng lực mô hình. Cấu hình này "
        "không thay thế YOLOv8n trong vai trò mô hình thực dụng. Nếu YOLOv8s tăng mAP nhưng recall trên ACDC ít thay đổi, có "
        "thể kết luận rằng năng lực mô hình giúp cải thiện biểu diễn và định vị, song chưa giải quyết hoàn toàn khó khăn của đối tượng "
        "nhỏ hoặc hiếm. Không nên mở rộng kết luận thành 'năng lực mô hình không phải nút thắt'."
    )
    p(
        "Copy-paste được thử nghiệm để tăng số lần xuất hiện của đối tượng hiếm bằng cách ghép vùng đối tượng vào ảnh khác. "
        "Do phép ghép có thể làm thay đổi quan hệ giữa vật thể và bối cảnh, kết quả được đánh giá trên cả XWOD, DAWN và ACDC. "
        "Nếu hiệu năng giảm, các khả năng như biên cắt không tự nhiên, khác biệt chiếu sáng, tỉ lệ hoặc ngữ cảnh chỉ được nêu "
        "như giả thuyết. Không có phân tích riêng chứng minh một trong các yếu tố đó là nguyên nhân duy nhất."
    )

    h2("3.6. Thiết kế thực nghiệm và bảo đảm khả năng tái lập")
    h3("3.6.1. Ma trận thực nghiệm")
    p(
        "Các run được tổ chức thành những nhóm trả lời câu hỏi riêng. Benchmark xác định đặc tính của từng họ mô hình. SE "
        "ablation kiểm tra thay đổi kiến trúc. Phase 2 và oversampling kiểm tra thay đổi dữ liệu. Thực nghiệm 960 kiểm tra "
        "độ phân giải; YOLOv8s kiểm tra năng lực mô hình; copy-paste kiểm tra một hình thức tăng cường dữ liệu. Việc tách câu "
        "hỏi giúp tránh quy toàn bộ chênh lệch cho một nguyên nhân khi nhiều cấu hình đã thay đổi cùng lúc."
    )
    cap("Bảng 3.6. Ma trận các nhóm thực nghiệm")
    table(
        ["Nhóm", "Cấu hình", "Biến khảo sát", "Tập đánh giá", "Câu hỏi cần trả lời"],
        [
            ["Benchmark", "YOLOv8n, YOLO11n, Faster R-CNN, RT-DETR", "Kiến trúc", "XWOD test, DAWN val", "Mô hình nào chính xác nhất và mô hình nào phù hợp triển khai?"],
            ["SE", "YOLOv8n chuẩn và YOLOv8n+SE", "Chú ý kênh", "XWOD test, DAWN val", "SE có cải thiện nhất quán giữa các miền không?"],
            ["Phase 2 v3", "Dữ liệu gộp và oversampling", "Composition dữ liệu", "XWOD, DAWN, ACDC", "Tăng mẫu lớp hiếm có cải thiện mà không làm suy giảm lớp phổ biến không?"],
            ["Độ phân giải", "v3_rare @640 và v3_960 @960", "Kích thước ảnh", "XWOD, DAWN, ACDC", "Chi tiết không gian bổ sung có bù được chi phí suy luận không?"],
            ["Năng lực", "YOLOv8n và YOLOv8s v3 @960", "Kích thước mô hình", "XWOD, DAWN, ACDC", "Tăng năng lực mô hình còn đem lại mức cải thiện nào?"],
            ["Copy-paste", "v3 @960 và copy-paste @960", "Tăng cường dữ liệu", "XWOD, DAWN, ACDC", "Ảnh ghép có giúp các lớp hiếm tổng quát tốt hơn không?"],
        ],
        widths=[0.8, 1.45, 1.0, 1.25, 1.85],
    )

    h3("3.6.2. Giao thức đánh giá và diễn giải kết quả")
    p(
        "XWOD val được dùng chọn checkpoint, còn XWOD test đo hiệu năng trên cùng miền dữ liệu thời tiết bất lợi. DAWN val "
        "đo khả năng chuyển sang một nguồn dữ liệu không tham gia huấn luyện. ACDC val được báo cáo ở hai bối cảnh: trước "
        "Phase 2 là đánh giá chéo dữ liệu; sau Phase 2 là đánh giá trên split được giữ lại của cùng bộ dữ liệu. Các vai trò "
        "này được ghi ngay trong bảng kết quả để người đọc không đánh đồng ba loại phép đo."
    )
    cap("Bảng 3.7. Giao thức đánh giá theo từng tập dữ liệu")
    table(
        ["Tập", "Quy mô đánh giá", "Vai trò", "Lưu ý khi diễn giải"],
        [
            ["XWOD val", "1.744 ảnh", "Chọn checkpoint", "Không dùng thay cho XWOD test."],
            ["XWOD test", "2.320 ảnh hợp lệ/2.321 tệp nguồn", "Đánh giá cùng miền", "Một tệp GIF gắn đuôi JPG bị Ultralytics loại."],
            ["DAWN val", "206 ảnh; 1.806 đối tượng", "Đánh giá ngoài miền", "Bicycle chỉ có 6 đối tượng nên AP có độ bất định cao."],
            ["ACDC val trước Phase 2", "398 ảnh; 2.830 đối tượng", "Đánh giá chéo dữ liệu", "Mô hình chưa học từ ACDC train."],
            ["ACDC val sau Phase 2", "398 ảnh; 2.830 đối tượng", "Đánh giá held-out", "Không còn là zero-shot thuần túy."],
        ],
        widths=[1.55, 1.55, 1.35, 2.0],
    )
    p(
        "Các chỉ số tổng thể được tính trên cùng sáu lớp, nhưng kết quả theo lớp vẫn phải đi kèm số lượng đối tượng. DAWN "
        "chỉ có sáu instance bicycle nên một giá trị AP cao hoặc thấp của lớp này chưa đủ để kết luận mô hình ổn định. Tương "
        "tự, cải thiện mAP trung bình không được xem là thành công nếu recall giảm rõ hoặc các lớp hiếm tiếp tục suy giảm. "
        "Kết quả âm được giữ lại trong báo cáo vì chúng xác định giới hạn của phương án đã thử và tránh tạo cảm giác mọi thay "
        "đổi đều dẫn đến cải thiện."
    )

    h3("3.6.3. Quản lý kết quả và giới hạn bằng chứng")
    p(
        "Mỗi kết quả được gắn với tên run, checkpoint, cấu hình dữ liệu, kích thước ảnh và log evaluation. Thứ tự ưu tiên khi "
        "xử lý mâu thuẫn là bằng chứng trực tiếp từ checkpoint, dataset và log; tiếp theo là bản tổng hợp provenance; cuối "
        "cùng mới là câu chữ cũ trong báo cáo. Số liệu không được điều chỉnh chỉ để các bảng có vẻ đồng nhất. Khi thiếu log, "
        "ô dữ liệu được để ở trạng thái không khả dụng hoặc mức xác nhận được ghi rõ."
    )
    p(
        "Đối với cấu hình cuối, bộ trọng số YOLOv8n Phase 2 v3 ở kích thước 960 đã được nạp lại và đánh giá trực tiếp trên "
        "ba tập. Kiến trúc kiểm tra được là YOLOv8n chuẩn, không chứa SE. YOLOv8s Phase 2 v3 ở kích thước 960 được giữ như "
        "mốc trần độ chính xác. Hai giới hạn provenance còn lại là log text benchmark bốn mô hình chưa đầy đủ đồng đều và "
        "quan hệ kế thừa trọng số giữa v3_rare với v3_960 chưa được chứng minh. Các giới hạn này không đòi hỏi huấn luyện lại mô hình cuối, "
        "nhưng giới hạn mức độ khẳng định trong phần thảo luận."
    )
    p(
        "Từ thiết kế trên, Chương 4 lần lượt trình bày môi trường triển khai, kết quả benchmark, ablation SE, quá trình mở "
        "rộng Phase 2 và các thí nghiệm bổ trợ. Cách sắp xếp này cho phép đối chiếu trực tiếp mỗi kết quả với câu hỏi thực "
        "nghiệm đã nêu, đồng thời phân biệt rõ mô hình có độ chính xác cao nhất với mô hình được chọn cho mục tiêu triển khai."
    )


def reset_toc_cache(document: Document) -> None:
    """Replace stale cached TOC entries with a clean field for Word to rebuild."""
    body = document.element.body
    start = None
    stop = None
    children = list(body.iterchildren())
    for index, child in enumerate(children):
        if not child.tag.endswith("}p"):
            continue
        instructions = [node.text or "" for node in child.xpath(".//w:instrText")]
        if any(text.strip().startswith("TOC ") for text in instructions):
            start = child
            for following in children[index + 1 :]:
                if not following.tag.endswith("}p"):
                    stop = following
                    break
                style = following.find("./w:pPr/w:pStyle", namespaces=following.nsmap)
                style_id = style.get(qn("w:val")) if style is not None else ""
                if not style_id.startswith("TOC"):
                    stop = following
                    break
            break
    if start is None or stop is None:
        raise RuntimeError("Không xác định được vùng mục lục để làm mới.")

    current = start
    while current is not stop:
        following = current.getnext()
        body.remove(current)
        current = following

    paragraph = document.add_paragraph(style="toc 1")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    paragraph.add_run()._r.append(begin)
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = ' TOC \\o "1-3" \\h \\z '
    paragraph.add_run()._r.append(instruction)
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    paragraph.add_run()._r.append(separate)
    paragraph.add_run("Đang cập nhật mục lục...")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    paragraph.add_run()._r.append(end)
    stop.addprevious(paragraph._p)


def main(path_str: str) -> None:
    path = Path(path_str)
    document = Document(path)
    _, chapter4_heading = remove_chapter3_body(document)
    builder = ChapterBuilder(document, chapter4_heading)
    add_chapter3(builder)
    builder.commit()
    update_manual_figure_table_lists(document)
    reset_toc_cache(document)
    mark_fields_dirty(document)
    document.save(path)
    print(f"Đã ghi bản Chương 3 học thuật: {path}")


if __name__ == "__main__":
    main(sys.argv[1])
