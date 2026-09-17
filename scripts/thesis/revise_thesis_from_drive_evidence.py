#!/usr/bin/env python3
"""Revise the thesis using values verified from the 2026-09-01 Drive audit."""

from __future__ import annotations

import argparse
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


def set_paragraph(paragraph, text: str) -> None:
    """Replace paragraph text while retaining its paragraph style."""
    paragraph.text = text


def replace_prefix(doc: Document, prefix: str, replacement: str, expected: int = 1) -> None:
    matches = [p for p in doc.paragraphs if p.text.strip().startswith(prefix)]
    if len(matches) != expected:
        raise RuntimeError(f"Expected {expected} paragraph(s) starting {prefix!r}, found {len(matches)}")
    for paragraph in matches:
        set_paragraph(paragraph, replacement)


def replace_contains(doc: Document, needle: str, replacement: str, expected: int = 1) -> None:
    matches = [p for p in doc.paragraphs if needle in p.text]
    if len(matches) != expected:
        raise RuntimeError(f"Expected {expected} paragraph(s) containing {needle!r}, found {len(matches)}")
    for paragraph in matches:
        set_paragraph(paragraph, replacement)


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def set_cell(cell, text: str, bold: bool = False) -> None:
    cell.text = text
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        for run in paragraph.runs:
            run.bold = bold
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def fill_table(table, rows: list[list[str]]) -> None:
    if len(table.rows) != len(rows) or len(table.columns) != len(rows[0]):
        raise RuntimeError(
            f"Table shape mismatch: document={len(table.rows)}x{len(table.columns)}, "
            f"data={len(rows)}x{len(rows[0])}"
        )
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            set_cell(table.cell(row_index, column_index), value, bold=row_index == 0)


def replace_table(doc: Document, old_table, rows: list[list[str]]):
    new_table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    if old_table.style:
        new_table.style = old_table.style
    new_table.autofit = True
    fill_table(new_table, rows)
    old_element = old_table._element
    new_element = new_table._element
    old_element.addprevious(new_element)
    old_element.getparent().remove(old_element)
    return new_table


def set_table_font(table, size: float) -> None:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(size)
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def replace_caption_image(doc: Document, caption_prefix: str, image_path: Path) -> None:
    """Replace the image immediately before a body caption, preserving its size/layout."""
    captions = [
        (index, paragraph)
        for index, paragraph in enumerate(doc.paragraphs)
        if index > 400 and paragraph.text.strip().startswith(caption_prefix)
    ]
    if len(captions) != 1:
        raise RuntimeError(f"Expected one body caption {caption_prefix!r}, found {len(captions)}")
    caption_index, _ = captions[0]
    for paragraph in reversed(doc.paragraphs[max(0, caption_index - 3) : caption_index]):
        blips = paragraph._p.xpath(".//a:blip")
        if blips:
            rid = blips[0].get(qn("r:embed"))
            doc.part.related_parts[rid]._blob = image_path.read_bytes()
            return
    raise RuntimeError(f"No image found before caption {caption_prefix!r}")


def remove_empty_paragraphs_before_headings(doc: Document) -> None:
    """Remove accumulated empty paragraphs that can create an unintended blank page."""
    for heading in list(doc.paragraphs):
        if not heading.style.name.startswith("Heading 1"):
            continue
        while True:
            previous = heading._p.getprevious()
            if previous is None or previous.tag != qn("w:p"):
                break
            protected = previous.xpath(".//w:br[@w:type='page'] | ./w:pPr/w:sectPr | .//w:drawing")
            text = "".join(previous.itertext()).strip()
            if text or protected:
                break
            previous.getparent().remove(previous)


def shade_header(table) -> None:
    for cell in table.rows[0].cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = tc_pr.find(qn("w:shd"))
        if shd is None:
            shd = OxmlElement("w:shd")
            tc_pr.append(shd)
        shd.set(qn("w:fill"), "D9EAF7")


def replace_drive_tree(doc: Document) -> None:
    start = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "adverse_weather_project/")
    end = next(i for i, p in enumerate(doc.paragraphs) if i > start and p.text.startswith("Phụ lục B."))
    old = list(doc.paragraphs[start:end])
    lines = [
        "object_detection_in_adverse_weather/",
        "├── evaluations/",
        "│   ├── stage1/                 # Raw metrics, logs và biểu đồ Stage 1",
        "│   ├── stage2/                 # XWOD val và frozen-test Stage 2",
        "│   ├── phase2/                 # Frozen-test RT-DETR-L Phase 2",
        "│   └── weather_comparison/     # Đánh giá theo điều kiện thời tiết",
        "├── runs/",
        "│   ├── stage1/                 # Run archive và reproducibility metadata",
        "│   ├── stage2/                 # Run archive và training log",
        "│   └── phase2/                 # args.yaml, results.csv và checkpoint",
        "├── official_backup_from_old_machine/",
        "│   └── protocol/               # Dataset manifest, môi trường và script protocol",
        "└── dataset_protocol/           # Bản sửa định dạng ảnh XWOD có checksum",
    ]
    for index, line in enumerate(lines):
        set_paragraph(old[index], line)
    for paragraph in old[len(lines) :]:
        remove_paragraph(paragraph)


def revise(input_path: Path, output_path: Path, assets_dir: Path | None = None) -> None:
    doc = Document(input_path)

    # Abstract/introduction and unsafe Chapter-1 claims.
    replace_contains(
        doc,
        "đồng thời khảo sát tác động của cơ chế chú ý Squeeze-and-Excitation (SE)",
        "Xuất phát từ thực tiễn đó, khóa luận thực hiện đề tài \"Thiết kế và phát triển mô hình học sâu phát hiện người và phương tiện trong điều kiện thời tiết bất lợi phục vụ giám sát giao thông thông minh\". Đề tài xây dựng quy trình huấn luyện nhiều giai đoạn, so sánh ba họ kiến trúc và thiết kế khảo sát cơ chế chú ý Squeeze-and-Excitation (SE). Do kho bằng chứng Drive chưa lưu artifact huấn luyện hoặc đánh giá SE, báo cáo không đưa ra kết luận định lượng cho nhánh này.",
        expected=1,
    )
    replace_contains(
        doc,
        "và khảo sát tác động của cơ chế chú ý Squeeze-and-Excitation (SE)",
        "Khóa luận thực hiện đề tài \"Thiết kế và phát triển mô hình học sâu phát hiện người và phương tiện trong điều kiện thời tiết bất lợi phục vụ giám sát giao thông thông minh\". Mục tiêu là xây dựng quy trình huấn luyện nhiều giai đoạn, so sánh ba họ kiến trúc và thiết kế khảo sát SE. Do kho bằng chứng Drive chưa lưu artifact SE, báo cáo không đưa ra kết luận định lượng cho nhánh này.",
        expected=1,
    )
    replace_prefix(
        doc,
        "Hệ thống giám sát giao thông phải vận hành liên tục",
        "Hệ thống giám sát giao thông phải vận hành trong nhiều điều kiện quan sát khác nhau. Sương mù làm suy giảm tương phản, mưa tạo vệt nhiễu và phản xạ, cát bụi gây ám màu, còn tuyết hoặc thiếu sáng có thể che khuất chi tiết của đối tượng. Những biến đổi này làm tăng chênh lệch miền giữa dữ liệu huấn luyện và dữ liệu vận hành, từ đó gây khó khăn cho cả phân loại lẫn định vị hộp giới hạn.",
    )
    replace_prefix(
        doc,
        "Để thấy rõ tác động thực tế, hãy hình dung",
        "Mức suy giảm mAP không thể quy đổi trực tiếp thành tỉ lệ phương tiện bị bỏ sót hoặc nhận dạng sai, vì mAP tổng hợp đường cong Precision–Recall trên nhiều lớp và ngưỡng IoU. Do đó, khóa luận đánh giá tác động của thời tiết bằng Precision, Recall, mAP50 và mAP50-95 trên các split được giữ cố định, thay vì diễn giải mAP như độ chính xác trên từng phương tiện.",
    )
    replace_prefix(
        doc,
        "Xây dựng quy trình học chuyển giao gồm pretrained",
        "Xây dựng quy trình học chuyển giao pretrained → BDD100K → XWOD; tiếp tục RT-DETR-L bằng tập dữ liệu gộp đăng ký dưới tên phase2_merged_yolo; sử dụng XWOD validation để chọn checkpoint và chỉ dùng các tập test cho đánh giá cuối sau khi cấu hình đã cố định.",
    )
    replace_prefix(
        doc,
        "Khảo sát tác động của cơ chế chú ý Squeeze-and-Excitation",
        "Thiết kế giao thức khảo sát cơ chế chú ý Squeeze-and-Excitation (SE) trên YOLOv8n; chỉ báo cáo kết quả định lượng khi có run, config và evaluation artifact truy vết được.",
    )
    replace_prefix(
        doc,
        "Về dữ liệu, BDD100K được dùng cho Stage 1",
        "Về dữ liệu, BDD100K được dùng cho Stage 1; XWOD là dữ liệu thời tiết bất lợi chính cho Stage 2; DAWN chỉ dùng để đánh giá ngoài miền; ACDC được đánh giá chéo trước Phase 2 và được dùng ở nhánh thích nghi đa miền. Cấu hình Phase 2 trỏ tới phase2_merged_yolo, nhưng Drive không lưu manifest của tập gộp nên quy mô sau lấy mẫu được ghi là chưa xác minh.",
    )
    replace_prefix(
        doc,
        "Về mô hình, đề tài khảo sát các hướng tiếp cận tiêu biểu",
        "Về mô hình, đề tài khảo sát YOLOv8n, YOLO11n, Faster R-CNN và RT-DETR-L, đại diện cho các hướng một giai đoạn, hai giai đoạn và Transformer detector. SE được trình bày như thiết kế ablation trên YOLOv8n; không có kết quả định lượng SE trong kho Drive, còn CBAM chỉ là hướng phát triển.",
    )
    replace_prefix(
        doc,
        "Về huấn luyện, các mô hình bắt đầu từ trọng số pretrained",
        "Về huấn luyện, các mô hình bắt đầu từ trọng số pretrained, tinh chỉnh trên BDD100K ở Stage 1 và tiếp tục trên XWOD ở Stage 2. XWOD validation được dùng để theo dõi hội tụ, chọn best checkpoint và làm căn cứ lựa chọn RT-DETR-L cho Phase 2; XWOD, BDD, DAWN và ACDC test chỉ được dùng cho đánh giá cuối sau khi cấu hình đã cố định.",
    )
    replace_prefix(
        doc,
        "Để khắc phục hạn chế về quy mô và mất cân bằng",
        "Đề tài áp dụng học chuyển giao theo chuỗi pretrained → BDD100K → XWOD. RT-DETR-L được tiếp tục bằng dataset đăng ký là phase2_merged_yolo nhằm fine-tune đa miền. Các source manifest xác nhận tổng 37.188 ảnh train cơ sở; merged manifest không có trong Drive nên quy mô và hệ số lấy mẫu cuối không được xem là số liệu đã xác minh.",
    )
    replace_prefix(
        doc,
        "Chương 3 trình bày phần phân tích và thiết kế hệ thống",
        "Chương 3 trình bày yêu cầu bài toán, kiến trúc hệ thống, chuẩn hóa dữ liệu, quy trình huấn luyện, tiêu chí lựa chọn mô hình và thiết kế khảo sát SE.",
    )
    replace_prefix(
        doc,
        "Chương 4 trình bày quá trình triển khai, thực nghiệm và đánh giá",
        "Chương 4 trình bày môi trường, cấu hình huấn luyện, kết quả benchmark, metadata Phase 2, phân tích theo lớp và thời tiết, cùng trạng thái bằng chứng của SE và các ablation sau Phase 2.",
    )
    replace_prefix(
        doc,
        "Phần tài liệu tham khảo liệt kê các công trình",
        "Phần tài liệu tham khảo liệt kê các công trình và tài liệu được sử dụng. Phần phụ lục cung cấp cấu trúc kho bằng chứng, cấu hình Phase 2 và các hình minh họa kết quả phát hiện đối tượng.",
    )

    # Methodology and dataset protocol.
    replace_prefix(
        doc,
        "Nhãn của XWOD được ánh xạ",
        "Nhãn của XWOD được ánh xạ về sáu lớp person, bicycle, car, motorcycle, bus và truck. XWOD train được dùng để tối ưu Stage 2, XWOD validation dùng cho early stopping và model selection, còn XWOD test chỉ dùng để báo cáo hiệu năng cùng miền sau khi checkpoint đã được cố định.",
    )
    replace_prefix(
        doc,
        "Quy trình huấn luyện được xây dựng theo hướng lũy tiến",
        "Quy trình huấn luyện được xây dựng theo hướng lũy tiến: pretrained → BDD100K → XWOD. RT-DETR-L sau Stage 2 được tiếp tục bằng dataset đăng ký là phase2_merged_yolo. DAWN không tham gia huấn luyện hoặc model selection và được giữ làm phép đánh giá ngoài miền.",
    )
    replace_prefix(
        doc,
        "XWOD test có 2.321 tệp ảnh ở nguồn",
        "Protocol chính thức ghi nhận XWOD gồm 6.006 ảnh train, 1.001 ảnh validation và 3.003 ảnh test. Artifact định dạng cho flooding_test_00217 cho thấy tệp ảnh gắn sai định dạng đã được sửa và kiểm tra bằng SHA-256 trước khi tạo manifest; vì vậy official evaluation sử dụng đầy đủ split test mới.",
    )
    replace_prefix(
        doc,
        "Phase 2 được thực hiện với RT-DETR-L",
        "Run Phase 2 sử dụng RT-DETR-L và trỏ tới dataset phase2_merged_yolo. Các manifest nguồn xác nhận XWOD train có 6.006 ảnh, ACDC train 1.182 ảnh và BDD100K train 30.000 ảnh, tổng cơ sở là 37.188 ảnh. Tuy nhiên, Drive không chứa manifest của tập gộp nên quy mô cuối sau lấy mẫu và danh sách ảnh phát lại chưa thể kiểm tra độc lập.",
    )
    replace_prefix(
        doc,
        "Các quy mô 6k, 1,2k và 30k",
        "Báo cáo sử dụng số đếm chính xác từ các source manifest thay cho ký hiệu làm tròn. Tổng 37.188 chỉ là tổng của ba split nguồn; không được diễn giải là kích thước train loader cuối khi merged manifest chưa có trong Drive.",
    )
    replace_prefix(
        doc,
        "Chiến lược lấy mẫu lớp hiếm tăng tần suất ảnh chứa bicycle",
        "Thiết kế hiện tại mô tả lấy mẫu tăng cường cho bicycle, motorcycle và bus. Do Drive không lưu merged manifest hoặc provenance của bước lấy mẫu, hệ số nhân, số bản sao và quy mô train cuối không được đưa vào kết quả đã xác minh.",
    )
    replace_prefix(
        doc,
        "Kết quả Stage 2 cho thấy RT-DETR-L đạt mAP50-95 cao nhất",
        "Re-evaluation trên XWOD validation cho thấy RT-DETR-L đạt mAP50-95 0,559, cao hơn YOLO11n 0,506, YOLOv8n 0,503 và Faster R-CNN 0,421. Đây là căn cứ validation để lựa chọn RT-DETR-L cho Phase 2; các kết quả XWOD, DAWN, BDD và ACDC test được dùng cho đánh giá frozen sau đó.",
    )
    replace_prefix(
        doc,
        "Sau thực nghiệm so sánh, RT-DETR-L được đưa vào Phase 2",
        "Sau khi được lựa chọn bằng XWOD validation, RT-DETR-L được đưa vào Phase 2 từ best checkpoint Stage 2. Run dùng ảnh 640×640, batch 16, AdamW, lr0 = 5×10⁻⁵, patience 20 và seed 42; dataset được đăng ký dưới tên phase2_merged_yolo.",
    )
    replace_prefix(
        doc,
        "Trong mỗi lần huấn luyện, bộ trọng số tốt nhất được chọn",
        "Trong mỗi lần huấn luyện, best checkpoint được chọn theo validation. Phase 2 chạy tối đa 50 epoch, đạt mAP50-95 validation cao nhất 0,562 tại epoch 13 và dừng sau epoch 33 theo patience 20. Thời gian tích lũy trong results.csv là 31.689,8 giây (8,803 giờ); best.pt có kích thước 66.156.359 byte.",
    )
    replace_prefix(
        doc,
        "Khảo sát SE thuộc nhánh YOLOv8n riêng",
        "SE được giữ ở mức thiết kế ablation trên YOLOv8n và không thuộc RT-DETR-L Phase 2. Inventory Drive không có config, checkpoint, training log hoặc evaluation result của nhánh SE, vì vậy báo cáo không tuyên bố thí nghiệm đã hoàn tất và không gán số liệu. CBAM chỉ là hướng nghiên cứu tiếp theo.",
    )
    replace_prefix(
        doc,
        "Trong đề tài này, SE được khảo sát như một biến thể ablation",
        "Trong đề tài này, SE được thiết kế như một ablation trên YOLOv8n theo nguyên tắc thay đổi một biến duy nhất. Đây là phương án thực nghiệm dự kiến; do Drive chưa có run và evaluation artifact, nội dung chỉ mô tả kiến trúc và giao thức, không khẳng định thí nghiệm đã hoàn tất.",
    )
    replace_prefix(
        doc,
        "SE được khảo sát như một phép loại trừ thành phần kiến trúc",
        "Thiết kế ablation dự kiến chèn một khối SE sau SPPF của YOLOv8n, tại feature map đã tích lũy đặc trưng ngữ nghĩa mức cao. Nhánh squeeze dùng global average pooling; nhánh excitation tạo trọng số qua hai phép biến đổi và sigmoid, sau đó tái cân bằng các kênh đặc trưng.",
    )
    replace_prefix(
        doc,
        "Việc chỉ chèn một khối giúp giới hạn số biến thay đổi",
        "Thiết kế chỉ chèn một khối nhằm giới hạn số biến thay đổi. Một phép so sánh hợp lệ cần giữ cố định dataset, split, seed và cấu hình huấn luyện, sau đó đánh giá trên XWOD validation và frozen test. Vì các artifact này chưa có trong Drive, báo cáo không kết luận SE có lợi hay bất lợi.",
    )
    replace_prefix(
        doc,
        "Hướng cải tiến dữ liệu tập trung vào hai vấn đề",
        "Thiết kế Phase 2 hướng tới joint multi-domain fine-tuning để giảm chênh lệch miền và mất cân bằng lớp. Cấu hình run xác nhận dataset phase2_merged_yolo, nhưng merged manifest không có trong Drive; vì vậy thành phần chi tiết và hệ số lấy mẫu chỉ được xem là thông tin thiết kế chưa đủ provenance.",
    )
    replace_prefix(
        doc,
        "Các nhóm thực nghiệm trả lời những câu hỏi riêng",
        "Các nhóm thực nghiệm trả lời những câu hỏi riêng: Stage 1 đo thích nghi miền giao thông; Stage 2 đo khả năng học XWOD và chuyển miền; Phase 2 đánh giá RT-DETR-L sau fine-tuning đa miền. SE chỉ có thiết kế, còn A0R và A1-DINO không có artifact trong Drive nên không được đưa vào bảng kết quả định lượng.",
    )
    replace_prefix(
        doc,
        "Theo nguồn mới nhất, cấu hình cuối được báo cáo",
        "Cấu hình cuối được xác nhận là RT-DETR-L Phase 2 ở kích thước ảnh 640. Drive lưu results.csv, best.pt và frozen-test JSON cho XWOD, BDD, DAWN và ACDC; do đó epoch, thời gian, validation và test đã có thể truy vết. Khoảng trống còn lại là merged-dataset manifest, SE, A0R/A1-DINO và citation metadata của XWOD.",
    )
    replace_prefix(
        doc,
        "Từ thiết kế trên, Chương 4 trình bày",
        "Từ thiết kế trên, Chương 4 trình bày môi trường, cấu hình Stage 1–Stage 2, metadata Phase 2, phân tích frozen test theo lớp và thời tiết, cùng giới hạn bằng chứng. Mỗi số liệu chính được truy về JSON, log, config, results.csv hoặc manifest trong Drive.",
    )
    replace_prefix(
        doc,
        "Kiến trúc hệ thống được chia thành bốn tầng",
        "Kiến trúc hệ thống được chia thành bốn tầng: dữ liệu, mô hình, huấn luyện–đánh giá và bằng chứng. Dữ liệu sau chuẩn hóa là đầu vào của quá trình huấn luyện; huấn luyện tạo checkpoint; XWOD validation được dùng để lựa chọn trước khi cấu hình được khóa và chuyển sang frozen test. Tầng bằng chứng lưu JSON, log, config, manifest và checkpoint để mọi bảng số liệu có thể truy vết độc lập.",
    )
    replace_prefix(
        doc,
        "Tầng dữ liệu đảm nhiệm việc đọc ảnh",
        "Tầng dữ liệu đảm nhiệm đọc ảnh, chuyển đổi nhãn và tổ chức train, validation, test. Tầng mô hình đóng gói YOLOv8n, YOLO11n, Faster R-CNN và RT-DETR-L trong cùng không gian sáu lớp. Tầng huấn luyện–đánh giá điều phối các giai đoạn tinh chỉnh, chọn best checkpoint bằng validation và chỉ chạy frozen test sau khi cấu hình đã cố định. Tầng bằng chứng phân tách rõ số liệu đã xác minh với thiết kế chưa có artifact, như nhánh SE.",
    )

    # Chapter 4 evidence source and environment.
    replace_prefix(
        doc,
        "Chương này sử dụng bảng tổng hợp adverse_weather_model_results_google_sheets.xlsx",
        "Chương này sử dụng artifact trong Google Drive object_detection_in_adverse_weather làm nguồn bằng chứng. Raw evaluation JSON/log được ưu tiên trước config, manifest, training log và câu chữ của bản báo cáo cũ. Các giá trị không truy được về artifact Drive được đánh dấu là chưa xác minh thay vì suy đoán.",
    )
    replace_prefix(
        doc,
        "Theo sheet Training Details",
        "Môi trường được ghi nhận sử dụng NVIDIA GeForce RTX 5090 với 32.607 MiB VRAM. Các capture nvidia-smi cho biết driver hỗ trợ CUDA tối đa 13.2, trong khi PyTorch 2.7.1+cu128 và torchvision 0.22.1+cu128 sử dụng build CUDA 12.8; Ultralytics là 8.3.159 và Python là 3.12.13. Driver xuất hiện ở hai phiên bản 595.80 và 595.84, nên báo cáo không chọn một số duy nhất.",
    )
    replace_prefix(
        doc,
        "Stage 1 sử dụng BDD100K 30k",
        "Stage 1 sử dụng BDD100K train 30.000 ảnh; Stage 2 tiếp tục từ best checkpoint Stage 1 và tinh chỉnh trên XWOD train 6.006 ảnh. Phase 2 chạy RT-DETR-L với dataset phase2_merged_yolo, đạt best validation tại epoch 13 và dừng ở epoch 33. Source manifests có tổng 37.188 ảnh train cơ sở, nhưng final merged size chưa xác minh do thiếu manifest của tập gộp.",
    )
    replace_prefix(
        doc,
        "Pipeline thực nghiệm gồm pretrained",
        "Pipeline thực nghiệm gồm pretrained → Stage 1 trên BDD100K → Stage 2 trên XWOD → Phase 2 với dataset phase2_merged_yolo. DAWN không tham gia huấn luyện và được dùng để đánh giá ngoài miền. ACDC test sau Phase 2 được diễn giải là held-out evaluation theo thiết kế đa miền, không phải zero-shot.",
    )
    replace_prefix(
        doc,
        "Khi các bộ trọng số Stage 1 được đánh giá trực tiếp",
        "Artifact Stage 1 xác nhận kết quả BDD và DAWN cho bốn kiến trúc; bundle bổ sung chỉ cung cấp XWOD/ACDC cho RT-DETR-L. Vì thiếu raw Stage 1 XWOD/ACDC của ba mô hình còn lại, báo cáo không xếp hạng đầy đủ bốn kiến trúc trên hai tập này.",
    )
    replace_prefix(
        doc,
        "Sau khi tinh chỉnh trên XWOD, cả bốn mô hình đều tăng mạnh",
        "Sau khi tinh chỉnh trên XWOD, RT-DETR-L đạt mAP50-95 0,500 trên XWOD test; YOLO11n đạt 0,454, YOLOv8n 0,452 và Faster R-CNN 0,381. Raw JSON xác nhận thứ hạng này trên frozen test, trong khi quyết định chuyển RT-DETR-L sang Phase 2 dựa trên XWOD validation.",
    )
    replace_prefix(
        doc,
        "RT-DETR-L được lựa chọn cho Phase 2 vì đạt mAP50-95 cao nhất",
        "RT-DETR-L được lựa chọn cho Phase 2 dựa trên XWOD validation: mAP50-95 đạt 0,559, so với 0,506 của YOLO11n, 0,503 của YOLOv8n và 0,421 của Faster R-CNN. XWOD/DAWN test chỉ được dùng để đánh giá frozen sau lựa chọn, không dùng để điều chỉnh cấu hình.",
    )
    replace_prefix(
        doc,
        "SE đã được triển khai như một khảo sát kiến trúc",
        "Báo cáo mô tả thiết kế chèn SE vào YOLOv8n, nhưng inventory Drive không có run, config, checkpoint hoặc evaluation artifact tương ứng. Vì vậy trạng thái thực nghiệm SE là chưa xác minh và không có bảng kết quả định lượng.",
    )
    replace_prefix(
        doc,
        "Trong phạm vi bằng chứng còn truy vết được, SE không được đưa",
        "Trong phạm vi bằng chứng Drive, SE không thuộc cấu hình RT-DETR-L Phase 2 và chưa có đủ dữ liệu để kết luận hiệu quả. CBAM chưa có artifact và chỉ được xem là hướng nghiên cứu tiếp theo.",
    )
    replace_prefix(
        doc,
        "Phase 2 khởi tạo từ bộ trọng số RT-DETR-L Stage 2",
        "Phase 2 khởi tạo từ best checkpoint RT-DETR-L Stage 2 và sử dụng dataset phase2_merged_yolo. Trên BDD test, mAP50-95 phục hồi từ 0,183 lên 0,354, gần mức 0,362 của Stage 1. Mức phục hồi phù hợp với mục tiêu duy trì miền giao thông của thiết kế đa miền, nhưng merged manifest còn thiếu nên báo cáo không quy toàn bộ hiệu quả cho một thành phần dữ liệu riêng lẻ.",
    )
    replace_prefix(
        doc,
        "Trên ACDC test, mAP50-95 tăng từ 0,128",
        "Trên ACDC test, mAP50-95 tăng từ 0,128 lên 0,247 ở Phase 2. Theo thiết kế đa miền, đây là held-out evaluation sau thích nghi chứ không phải zero-shot. XWOD tăng 0,010, còn DAWN giảm 0,011 so với Stage 2; vì vậy Phase 2 không cải thiện đồng đều trên mọi miền.",
    )
    replace_prefix(
        doc,
        "Ở XWOD test Stage 2, RT-DETR-L đạt mAP50-95 cao nhất",
        "Ở XWOD test Stage 2, RT-DETR-L đạt mAP50-95 0,500 với 33,7 FPS; YOLOv8n đạt 0,452 với 112,3 FPS; YOLO11n đạt 0,454 với 95,5 FPS; Faster R-CNN đạt 0,381 với 61,2 FPS. Các phép đo dùng cùng GPU và batch 1, nhưng Faster R-CNN chạy bằng torchvision/torchmetrics trong khi ba mô hình còn lại dùng Ultralytics. Vì khác framework và preprocessing, mAP50-95 là chỉ số accuracy chính; Precision/Recall và latency/FPS không được diễn giải như benchmark kernel-equivalent.",
    )
    replace_prefix(
        doc,
        "Trong thực nghiệm so sánh, RT-DETR-L đạt độ chính xác cao nhất",
        "RT-DETR-L được chọn bằng XWOD validation và sau đó đạt mAP50-95 frozen-test lần lượt 0,510 trên XWOD, 0,354 trên BDD, 0,515 trên DAWN và 0,247 trên ACDC. Kết quả cho thấy Phase 2 phục hồi BDD và cải thiện ACDC, nhưng không cải thiện đồng đều mọi miền vì DAWN giảm nhẹ so với Stage 2.",
    )
    replace_prefix(
        doc,
        "Kết quả cho thấy Phase 2 tạo lợi ích lớn nhất",
        "Kết quả Phase 2 thay đổi nhiều nhất trên BDD và ACDC, trong khi XWOD tăng nhẹ và DAWN giảm nhẹ. Vì merged manifest còn thiếu, kết luận được giới hạn ở biến động hiệu năng của toàn cấu hình Phase 2, không quy nguyên nhân cho riêng replay hay oversampling.",
    )
    replace_prefix(
        doc,
        "Phase 2 đặc biệt hữu ích cho khả năng phục hồi BDD",
        "Phase 2 phục hồi đáng kể kết quả BDD và tăng mAP50-95 trên ACDC. Tuy nhiên, DAWN giảm nhẹ, ACDC recall chỉ đạt 0,384, còn bicycle, motorcycle và cảnh ban đêm vẫn là hạn chế. Đây là quan sát trên toàn cấu hình Phase 2, không phải bằng chứng nhân quả cho từng thành phần dữ liệu.",
    )
    replace_prefix(
        doc,
        "Workbook chưa ghi epoch hoàn thành",
        "Metadata Phase 2 đã được khôi phục từ args.yaml, results.csv và checkpoint archive. Các khoảng trống còn lại gồm manifest của phase2_merged_yolo, artifact SE, A0R/A1-DINO và nguồn trích dẫn XWOD; vì vậy báo cáo không đưa tổng sau oversampling hoặc kết luận ablation chưa được xác minh.",
    )
    replace_prefix(
        doc,
        "Hướng phát triển phù hợp là hoàn thiện nhật ký Phase 2",
        "Hướng phát triển ưu tiên là lưu merged-dataset manifest và provenance, đánh giá có kiểm soát cho SE/A0R/A1-DINO, bổ sung dữ liệu thật cho lớp hiếm và khảo sát riêng điều kiện ban đêm. Chỉ khi có artifact validation và frozen test đầy đủ mới có thể kết luận về các biến thể này.",
    )

    # Rename the SE section to reflect the actual evidence state.
    replace_prefix(doc, "4.3. Kết quả khảo sát cơ chế chú ý SE", "4.3. Trạng thái bằng chứng của khảo sát SE", expected=2)
    replace_prefix(doc, "4.3.1. Phạm vi bằng chứng", "4.3.1. Thiết kế khảo sát", expected=2)
    replace_prefix(doc, "4.3.2. Kết luận sử dụng trong báo cáo", "4.3.2. Giới hạn bằng chứng", expected=2)

    # Add verified Phase 2 validation metadata immediately before the test analysis.
    target = next(p for p in doc.paragraphs if p.text.startswith("Phase 2 khởi tạo từ best checkpoint"))
    target.insert_paragraph_before(
        "Run Phase 2 dùng tối đa 50 epoch, batch 16, ảnh 640×640, AdamW, lr0 = 5×10⁻⁵, patience 20 và seed 42. Best checkpoint xuất hiện tại epoch 13 với P = 0,827, R = 0,810, mAP50 = 0,848 và mAP50-95 = 0,562 trên XWOD validation. Huấn luyện dừng sau epoch 33, kéo dài 8,803 giờ; best.pt có kích thước 66.156.359 byte.",
        style=target.style,
    )

    # Captions and appendix prose.
    replace_prefix(
        doc,
        "Bảng dưới đây tổng hợp cấu hình tham chiếu được workbook xác nhận",
        "Bảng dưới đây tổng hợp cấu hình RT-DETR-L Phase 2 được xác nhận trực tiếp từ args.yaml, results.csv và checkpoint archive trong Drive.",
    )
    replace_prefix(
        doc,
        "Các hình trong phụ lục minh họa kết quả phát hiện",
        "Các hình trong phụ lục chỉ có vai trò minh họa. Drive không lưu metadata nối từng hình với checkpoint và sample ID, nên chúng không được dùng làm bằng chứng định lượng; các kết luận số chỉ dựa trên raw evaluation JSON/log.",
    )
    replace_prefix(doc, "Bảng 4.1. Môi trường thực nghiệm được xác nhận trong workbook", "Bảng 4.1. Môi trường thực nghiệm được xác nhận từ artifact Drive", expected=2)

    # Remove the placeholder XWOD bibliography item instead of inventing metadata.
    xwod_ref = next(p for p in doc.paragraphs if p.text.startswith("[15] Bộ dữ liệu XWOD"))
    remove_paragraph(xwod_ref)
    orphan_appendix = next(p for p in doc.paragraphs if p.style.name.lower().startswith("toc") and p.text.startswith("Phụ lục D."))
    remove_paragraph(orphan_appendix)
    replace_prefix(doc, "[16] A. Paszke", "[15] A. Paszke và cộng sự, \"PyTorch: An Imperative Style, High-Performance Deep Learning Library,\" trong Advances in Neural Information Processing Systems (NeurIPS), 2019, tr. 8026–8037.")
    replace_prefix(doc, "[17] Streamlit Inc.", "[16] Streamlit Inc., \"Streamlit: The fastest way to build and share data apps,\" 2019. [Trực tuyến]. Địa chỉ: https://streamlit.io")

    # Update tables.
    tables = doc.tables
    fill_table(
        tables[2],
        [
            ["Bộ dữ liệu", "Vai trò", "Giai đoạn và phân hoạch"],
            ["BDD100K", "Thích nghi miền giao thông", "Stage 1: train 30.000; val 3.000; test 7.000"],
            ["XWOD", "Miền thời tiết bất lợi chính", "Stage 2: train 6.006; val 1.001; test 3.003"],
            ["DAWN", "Đánh giá ngoài miền", "Không train; val 308; test 718"],
            ["ACDC", "Đánh giá chéo và thích nghi đa miền", "Train 1.182; val 197; test 591"],
        ],
    )
    fill_table(
        tables[5],
        [
            ["Bộ dữ liệu", "Phân hoạch chính thức", "Vai trò", "Tập đánh giá"],
            ["BDD100K", "30.000 / 3.000 / 7.000", "Stage 1; nguồn miền giao thông", "BDD test"],
            ["XWOD", "6.006 / 1.001 / 3.003", "Stage 2; nguồn miền đích", "XWOD val/test"],
            ["DAWN", "— / 308 / 718", "Đánh giá ngoài miền; không huấn luyện", "DAWN test"],
            ["ACDC", "1.182 / 197 / 591", "Đánh giá chéo; thích nghi đa miền", "ACDC test"],
        ],
    )
    fill_table(
        tables[7],
        [
            ["Thành phần", "Số ảnh nguồn đã xác minh", "Ghi chú"],
            ["XWOD train", "6.006", "Manifest nguồn"],
            ["ACDC train", "1.182", "Manifest nguồn"],
            ["BDD100K train", "30.000", "Manifest nguồn"],
            ["Tổng cơ sở", "37.188", "Chưa phải final loader size"],
            ["Sau lấy mẫu", "Chưa xác minh", "Thiếu merged manifest/provenance"],
        ],
    )
    fill_table(
        tables[9],
        [
            ["Nhóm", "Cấu hình", "Biến khảo sát", "Tập đánh giá", "Trạng thái"],
            ["Stage 1", "Pretrained → BDD100K", "Họ mô hình", "BDD/DAWN", "Có raw artifact"],
            ["Stage 2", "Stage 1 → XWOD", "Họ mô hình", "XWOD val; frozen tests", "Có raw artifact"],
            ["SE", "YOLOv8n + SE", "Cơ chế chú ý", "Dự kiến XWOD/DAWN", "Thiếu artifact"],
            ["Phase 2", "RT-DETR-L + phase2_merged_yolo", "Fine-tuning đa miền", "XWOD/BDD/DAWN/ACDC", "Có run và frozen tests"],
            ["A0R/A1-DINO", "Post-Phase2 ablation", "Random vs retrieval", "Validation/frozen test", "Không tìm thấy artifact"],
        ],
    )
    fill_table(
        tables[10],
        [
            ["Tập", "Quy mô", "Vai trò", "Lưu ý"],
            ["XWOD val", "1.001 ảnh", "Early stopping/model selection", "Không dùng thay test"],
            ["XWOD test", "3.003 ảnh", "Frozen test cùng miền", "Chỉ đánh giá sau lựa chọn"],
            ["DAWN test", "718 ảnh", "External/OOD evaluation", "Không tham gia huấn luyện"],
            ["ACDC test trước Phase 2", "591 ảnh", "Đánh giá chéo dữ liệu", "Chưa học ACDC train"],
            ["ACDC test sau Phase 2", "591 ảnh", "Held-out evaluation", "Không còn zero-shot"],
        ],
    )
    fill_table(
        tables[11],
        [
            ["Thành phần", "Giá trị xác nhận"],
            ["GPU", "NVIDIA GeForce RTX 5090, 32.607 MiB"],
            ["Driver / CUDA support", "595.80–595.84 / tối đa CUDA 13.2"],
            ["PyTorch", "2.7.1+cu128 (runtime/build CUDA 12.8)"],
            ["torchvision", "0.22.1+cu128"],
            ["Ultralytics", "8.3.159"],
            ["Python", "3.12.13"],
            ["Giao thức tốc độ", "batch 1; ảnh 640; cùng RTX 5090"],
        ],
    )
    configuration_table = replace_table(
        doc,
        tables[12],
        [
            ["Mô hình", "Giai đoạn", "Epoch max/run", "Best/dừng", "Cấu hình chính"],
            ["YOLOv8n", "Stage 1", "50/50", "Không dừng sớm", "AdamW; lr0=0,001; batch=16"],
            ["YOLO11n", "Stage 1", "50/50", "Không dừng sớm", "AdamW; lr0=0,001; batch=16"],
            ["RT-DETR-L", "Stage 1", "50/50", "Không dừng sớm", "AdamW; lr0=10⁻⁴; batch=4"],
            ["Faster R-CNN", "Stage 1 official", "10/10", "Best official", "SGD; lr0=0,005; batch=4"],
            ["YOLOv8n", "Stage 2", "50/49", "Best ep34; patience=15", "AdamW; lr0=5×10⁻⁴; batch=16"],
            ["YOLO11n", "Stage 2", "50/50", "Không dừng sớm", "AdamW; lr0=5×10⁻⁴; batch=16"],
            ["RT-DETR-L", "Stage 2", "50/50", "Không dừng sớm", "AdamW; lr0=5×10⁻⁴; batch=4"],
            ["Faster R-CNN", "Stage 2", "30/30", "Best ep21", "torchvision; SGD"],
            ["RT-DETR-L", "Phase 2", "50/33", "Best ep13; patience=20", "AdamW; lr0=5×10⁻⁵; batch=16"],
        ],
    )
    set_table_font(configuration_table, 10.5)

    weather_rows = [
        ["Bộ dữ liệu", "Điều kiện", "Số ảnh", "mAP50-95"],
        ["XWOD", "flooding", "1.549", "0,548"],
        ["XWOD", "fog", "91", "0,387"],
        ["XWOD", "rain", "200", "0,243"],
        ["XWOD", "sand", "160", "0,344"],
        ["XWOD", "snow", "362", "0,505"],
        ["XWOD", "tornado", "351", "0,280"],
        ["XWOD", "wildfire", "290", "0,656"],
        ["ACDC", "fog", "147", "0,388"],
        ["ACDC", "night", "147", "0,121"],
        ["ACDC", "rain", "148", "0,250"],
        ["ACDC", "snow", "149", "0,234"],
        ["DAWN", "fog", "215", "0,500"],
        ["DAWN", "rain", "138", "0,648"],
        ["DAWN", "sand", "224", "0,503"],
        ["DAWN", "snow", "141", "0,616"],
    ]
    weather_table = replace_table(doc, tables[19], weather_rows)
    set_table_font(weather_table, 9)

    # The table collection changed after replacing the weather table; locate the appendix table by content.
    appendix_table = next(t for t in doc.tables if t.cell(0, 0).text.strip() == "Thành phần" and "Mô hình Phase 2" in t.cell(1, 0).text)
    fill_table(
        appendix_table,
        [
            ["Thành phần", "Giá trị theo artifact Drive"],
            ["Mô hình Phase 2", "RT-DETR-L từ Stage 2 best.pt"],
            ["Dataset pointer", "phase2_merged_yolo/dataset.yaml"],
            ["Epoch tối đa / hoàn thành", "50 / 33"],
            ["Best epoch", "13 (mAP50-95 val = 0,562)"],
            ["Early stopping", "patience = 20"],
            ["Kích thước ảnh", "640 × 640"],
            ["Batch size / seed", "16 / 42"],
            ["Optimizer / lr0", "AdamW / 5×10⁻⁵"],
            ["Thời gian", "31.689,8 giây (8,803 giờ)"],
            ["Best checkpoint", "66.156.359 byte"],
        ],
    )

    for table in doc.tables:
        shade_header(table)

    # Keep the final FPS row of Table 2.3 with the rest of the compact table.
    set_table_font(doc.tables[3], 10.5)

    replace_drive_tree(doc)
    remove_empty_paragraphs_before_headings(doc)

    if assets_dir is not None:
        replace_caption_image(doc, "Hình 3.1.", assets_dir / "system_architecture.png")
        replace_caption_image(doc, "Hình 3.2.", assets_dir / "training_pipeline.png")

    # Basic pagination safeguards for headings/captions/tables.
    for paragraph in doc.paragraphs:
        name = paragraph.style.name
        if name.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.keep_together = True
        if name.startswith("Caption"):
            paragraph.paragraph_format.keep_with_next = True
    for table in doc.tables:
        for row in table.rows:
            tr_pr = row._tr.get_or_add_trPr()
            cant_split = tr_pr.find(qn("w:cantSplit"))
            if cant_split is None:
                tr_pr.append(OxmlElement("w:cantSplit"))

    # Update fields on opening in Word/LibreOffice.
    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--assets-dir", type=Path)
    args = parser.parse_args()
    revise(args.input, args.output, args.assets_dir)


if __name__ == "__main__":
    main()
