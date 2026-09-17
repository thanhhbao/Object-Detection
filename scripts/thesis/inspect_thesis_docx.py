#!/usr/bin/env python3
"""Inspect Word block order, styles, and media around thesis Chapter 3."""

from pathlib import Path
import sys

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph


def iter_blocks(document):
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, document)
        elif child.tag.endswith("}tbl"):
            yield Table(child, document)


def main(path_str: str) -> None:
    path = Path(path_str)
    doc = Document(path)
    blocks = list(iter_blocks(doc))
    print(f"path={path}")
    print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} inline_shapes={len(doc.inline_shapes)}")
    print("styles:")
    for style in doc.styles:
        if style.type == 1 and ("Heading" in style.name or style.name in {"Normal", "Caption", "Title"}):
            print(f"  {style.name!r}")
    print("chapter-3 block range:")
    in_ch3 = False
    for index, block in enumerate(blocks):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if text.startswith("CHƯƠNG 3") and block.style.name == "Heading 1":
                in_ch3 = True
            if in_ch3:
                print(f"P {index:04d} style={block.style.name!r} text={text[:180]!r}")
            if in_ch3 and text.startswith("CHƯƠNG 4"):
                break
        elif in_ch3:
            rows = len(block.rows)
            cols = len(block.columns)
            preview = " | ".join(cell.text.replace("\n", " / ") for cell in block.rows[0].cells) if rows else ""
            print(f"T {index:04d} style={block.style.name!r} size={rows}x{cols} first={preview[:180]!r}")


if __name__ == "__main__":
    main(sys.argv[1])
