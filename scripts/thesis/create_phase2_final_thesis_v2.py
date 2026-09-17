#!/usr/bin/env python3
"""Create the clean Phase-2-final thesis narrative without overwriting the source."""

from __future__ import annotations

import argparse
import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


def set_text(paragraph, text: str) -> None:
    paragraph.text = text


def find_one(doc: Document, prefix: str, *, style_prefix: str | None = None):
    matches = [
        p
        for p in doc.paragraphs
        if p.text.strip().startswith(prefix)
        and (style_prefix is None or p.style.name.lower().startswith(style_prefix.lower()))
    ]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph {prefix!r}, found {len(matches)}")
    return matches[0]


def paragraph_index(doc: Document, target) -> int:
    return next(i for i, paragraph in enumerate(doc.paragraphs) if paragraph._p is target._p)


def find_after(doc: Document, anchor, prefix: str):
    start = paragraph_index(doc, anchor)
    matches = [p for p in doc.paragraphs[start + 1 :] if p.text.strip().startswith(prefix)]
    if not matches:
        raise RuntimeError(f"No paragraph after anchor starts with {prefix!r}")
    return matches[0]


def replace_all_prefix(doc: Document, prefix: str, replacement: str) -> None:
    matches = [p for p in doc.paragraphs if p.text.strip().startswith(prefix)]
    if not matches:
        raise RuntimeError(f"No paragraph starts with {prefix!r}")
    for paragraph in matches:
        set_text(paragraph, replacement)


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def remove_range(doc: Document, start_prefix: str, end_prefix: str) -> None:
    paragraphs = doc.paragraphs
    start = next(
        i
        for i, p in enumerate(paragraphs)
        if p.style.name.startswith("Heading") and p.text.strip().startswith(start_prefix)
    )
    end = next(
        i
        for i, p in enumerate(paragraphs)
        if i > start and p.style.name.startswith("Heading") and p.text.strip().startswith(end_prefix)
    )
    for paragraph in list(paragraphs[start:end]):
        for blip in paragraph._p.xpath(".//a:blip"):
            rid = blip.get(qn("r:embed"))
            if rid in doc.part.rels:
                del doc.part.rels[rid]
        remove_paragraph(paragraph)


def insert_before(target, text: str, style: str):
    return target.insert_paragraph_before(text, style=style)


def copy_paragraph_properties(target, template) -> None:
    target_ppr = target._p.pPr
    if target_ppr is not None:
        target._p.remove(target_ppr)
    if template._p.pPr is not None:
        target._p.insert(0, deepcopy(template._p.pPr))


def insert_toc_before(target, text: str, template):
    paragraph = target.insert_paragraph_before(text, style=template.style)
    copy_paragraph_properties(paragraph, template)
    return paragraph


def set_cell(cell, text: str, *, bold: bool = False, size: float = 11) -> None:
    cell.text = text
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        for run in paragraph.runs:
            run.bold = bold
            run.font.name = "Times New Roman"
            run.font.size = Pt(size)
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def replace_table(doc: Document, old_table, rows: list[list[str]], size: float = 11):
    new_table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    if old_table.style:
        new_table.style = old_table.style
    new_table.autofit = True
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            set_cell(new_table.cell(row_index, column_index), value, bold=row_index == 0, size=size)
    old_table._element.addprevious(new_table._element)
    old_table._element.getparent().remove(old_table._element)
    return new_table


def shade_header(table) -> None:
    for cell in table.rows[0].cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = tc_pr.find(qn("w:shd"))
        if shd is None:
            shd = OxmlElement("w:shd")
            tc_pr.append(shd)
        shd.set(qn("w:fill"), "D9EAF7")


def set_table_font(table, size: float) -> None:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(size)
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def replace_image_before_caption(doc: Document, caption_prefix: str, image_path: Path) -> None:
    caption = find_one(doc, caption_prefix, style_prefix="caption")
    paragraphs = doc.paragraphs
    index = next(i for i, p in enumerate(paragraphs) if p._p is caption._p)
    for paragraph in reversed(paragraphs[max(0, index - 3) : index]):
        blips = paragraph._p.xpath(".//a:blip")
        if blips:
            rid = blips[0].get(qn("r:embed"))
            doc.part.related_parts[rid]._blob = image_path.read_bytes()
            return
    raise RuntimeError(f"No image before {caption_prefix}")


def replace_appendix_tree(doc: Document) -> None:
    start = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "object_detection_in_adverse_weather/")
    end = next(i for i, p in enumerate(doc.paragraphs) if i > start and p.text.startswith("Phụ lục B."))
    old = list(doc.paragraphs[start:end])
    lines = [
        "object_detection_in_adverse_weather/",
        "├── data/",
        "│   ├── bdd100k/",
        "│   ├── xwod/",
        "│   ├── acdc/",
        "│   └── dawn/",
        "├── configs/                  # Cấu hình dữ liệu và huấn luyện",
        "├── training/",
        "│   ├── stage1_bdd100k/",
        "│   ├── stage2_xwod/",
        "│   └── phase2_multidomain/",
        "├── evaluation/               # Frozen test, per-class và weather",
        "└── models/",
        "    └── rtdetr_l_phase2/       # Cấu hình cuối của đề tài",
    ]
    for index, line in enumerate(lines):
        if index < len(old):
            set_text(old[index], line)
        else:
            old[-1].insert_paragraph_before(line, style=old[-1].style)
    for paragraph in old[len(lines) :]:
        remove_paragraph(paragraph)


def revise(source: Path, output: Path, assets: Path) -> None:
    doc = Document(source)
    original_tables = list(doc.tables)

    # Front abstract and Chapter 1: research objectives and scope.
    replace_all_prefix(
        doc,
        "Khóa luận thực hiện đề tài",
        "Khóa luận thực hiện đề tài \"Thiết kế và phát triển mô hình học sâu phát hiện người và phương tiện trong điều kiện thời tiết bất lợi phục vụ giám sát giao thông thông minh\". Nghiên cứu xây dựng quy trình học chuyển giao nhiều giai đoạn, so sánh YOLOv8n, YOLO11n, Faster R-CNN và RT-DETR-L, sau đó lựa chọn RT-DETR-L bằng XWOD validation để tiếp tục fine-tuning đa miền ở Phase 2. RT-DETR-L Phase 2 là mô hình cuối cùng trong phạm vi thực nghiệm của khóa luận.",
    )
    replace_all_prefix(
        doc,
        "Xuất phát từ thực tiễn đó, khóa luận thực hiện đề tài",
        "Xuất phát từ thực tiễn đó, khóa luận xây dựng một quy trình thực nghiệm có kiểm soát cho bài toán phát hiện người và phương tiện trong điều kiện thời tiết bất lợi. Bốn kiến trúc YOLOv8n, YOLO11n, Faster R-CNN và RT-DETR-L được so sánh trên cùng không gian sáu lớp; mô hình được chọn bằng XWOD validation trước khi bước vào Phase 2.",
    )
    replace_all_prefix(
        doc,
        "Đề tài áp dụng học chuyển giao theo chuỗi",
        "Quy trình nghiên cứu gồm COCO pretrained → Stage 1 trên BDD100K → Stage 2 trên XWOD → lựa chọn bằng XWOD validation → Phase 2 multi-domain fine-tuning trên XWOD, ACDC và BDD100K. Sau khi cấu hình được cố định, mô hình được đánh giá trên XWOD, BDD, DAWN và ACDC.",
    )
    replace_all_prefix(
        doc,
        "Mặt khác, dù hiện đã có nhiều hướng tiếp cận",
        "Mặt khác, các hướng phát hiện một giai đoạn, hai giai đoạn và dựa trên Transformer có đặc tính khác nhau về độ chính xác, tốc độ và chi phí tính toán. Việc đánh giá chúng theo cùng giao thức trên dữ liệu thời tiết bất lợi là cần thiết để lựa chọn kiến trúc phù hợp, đồng thời làm rõ ảnh hưởng của quá trình thích nghi miền và mất cân bằng lớp.",
    )
    replace_all_prefix(
        doc,
        "Tốc độ phát triển của hạ tầng giao thông",
        "Sự phát triển của hạ tầng và số lượng phương tiện làm gia tăng nhu cầu giám sát giao thông tự động. Các hệ thống thị giác máy tính có thể phân tích dữ liệu từ camera để hỗ trợ thống kê lưu lượng, phát hiện tình huống bất thường và cảnh báo rủi ro. Trong đó, phát hiện đối tượng là tác vụ nền tảng vì cung cấp vị trí, lớp và độ tin cậy của người cùng các phương tiện trong khung hình.",
    )
    replace_all_prefix(
        doc,
        "Nhờ sự tiến bộ của học sâu",
        "Các mô hình học sâu đạt hiệu năng cao trên nhiều bộ dữ liệu chuẩn, nhưng kết quả thường suy giảm khi ảnh chịu tác động của mưa, sương mù, tuyết, cát bụi, ngập nước hoặc thiếu sáng. Những điều kiện này làm giảm tương phản, che khuất chi tiết và thay đổi phân phối ảnh so với dữ liệu huấn luyện. Vì vậy, khả năng thích nghi với thời tiết bất lợi là yêu cầu quan trọng đối với hệ thống giám sát giao thông vận hành ngoài thực tế.",
    )
    replace_all_prefix(
        doc,
        "Vì vậy, khi gặp những bức ảnh suy giảm chất lượng",
        "Mỗi điều kiện thời tiết tạo ra một dạng suy giảm khác nhau, nên một mô hình chỉ thích nghi tốt với một miền dữ liệu chưa chắc duy trì hiệu năng trên miền khác. Điều này đặt ra yêu cầu đánh giá đa miền và tách rõ dữ liệu dùng để huấn luyện, lựa chọn mô hình và kiểm thử cuối.",
    )
    replace_all_prefix(
        doc,
        "Chính những vấn đề trên là lý do đề tài được lựa chọn",
        "Từ các vấn đề trên, đề tài tập trung xây dựng và đánh giá có hệ thống một quy trình phát hiện đối tượng nhiều giai đoạn cho dữ liệu giao thông trong điều kiện thời tiết bất lợi.",
    )

    objectives = [p for i, p in enumerate(doc.paragraphs) if p.style.name == "List Paragraph" and 260 < i < 280]
    objective_texts = [
        "Xây dựng quy trình học chuyển giao nhiều giai đoạn COCO pretrained → BDD100K → XWOD cho bài toán phát hiện người và phương tiện trong điều kiện thời tiết bất lợi.",
        "So sánh bốn kiến trúc YOLOv8n, YOLO11n, Faster R-CNN và RT-DETR-L trên cùng không gian sáu lớp.",
        "Đánh giá mô hình bằng Precision, Recall, mAP50, mAP50-95 và latency/FPS, đồng thời nêu rõ giới hạn của so sánh tốc độ giữa các framework.",
        "Lựa chọn mô hình bằng XWOD validation và tiếp tục tối ưu mô hình được chọn bằng Phase 2 multi-domain fine-tuning trên XWOD, ACDC và BDD100K.",
        "Phân tích mô hình cuối theo bộ dữ liệu, lớp đối tượng, điều kiện thời tiết, đánh đổi độ chính xác–tốc độ và khả năng duy trì tri thức giữa các miền.",
    ]
    if len(objectives) != 5:
        raise RuntimeError(f"Expected five objective bullets, found {len(objectives)}")
    for paragraph, text in zip(objectives, objective_texts):
        set_text(paragraph, text)

    replace_all_prefix(
        doc,
        "Về dữ liệu, BDD100K được dùng cho Stage 1",
        "Về dữ liệu, BDD100K được dùng ở Stage 1 để thích nghi miền giao thông và được đưa lại vào Phase 2 nhằm duy trì tri thức miền này. XWOD là miền thời tiết bất lợi chính ở Stage 2; DAWN chỉ dùng để đánh giá ngoài miền. ACDC test là đánh giá chéo trước Phase 2, còn ACDC train tham gia fine-tuning đa miền và ACDC test được giữ lại để đánh giá sau thích nghi.",
    )
    replace_all_prefix(
        doc,
        "Về mô hình, đề tài khảo sát YOLOv8n",
        "Về mô hình, đề tài khảo sát YOLOv8n, YOLO11n, Faster R-CNN và RT-DETR-L, đại diện cho các hướng một giai đoạn, hai giai đoạn và Transformer detector. RT-DETR-L được lựa chọn bằng XWOD validation và RT-DETR-L Phase 2 được xác định là mô hình cuối cùng của đề tài.",
    )
    replace_all_prefix(
        doc,
        "Đề tài được thực hiện theo phương pháp kết hợp",
        "Đề tài kết hợp nghiên cứu lý thuyết, xây dựng quy trình thực nghiệm và đánh giá định lượng. Phần lý thuyết tập trung vào các hướng phát hiện một giai đoạn, hai giai đoạn và dựa trên Transformer, cùng các kỹ thuật nền tảng gồm CNN, biểu diễn đặc trưng đa tỉ lệ, FPN, học chuyển giao và thích nghi miền.",
    )
    replace_all_prefix(
        doc,
        "Chương 3 trình bày yêu cầu bài toán",
        "Chương 3 trình bày yêu cầu bài toán, kiến trúc hệ thống, chuẩn hóa dữ liệu, bốn mô hình benchmark, chiến lược huấn luyện nhiều giai đoạn và giao thức thực nghiệm.",
    )
    replace_all_prefix(
        doc,
        "Chương 4 trình bày môi trường",
        "Chương 4 trình bày môi trường, cấu hình huấn luyện, benchmark bốn mô hình, Phase 2 của RT-DETR-L, phân tích kết quả và các hướng phát triển bám theo hạn chế quan sát được.",
    )
    replace_all_prefix(
        doc,
        "Phần tài liệu tham khảo liệt kê",
        "Phần tài liệu tham khảo liệt kê các công trình và tài liệu được sử dụng. Phần phụ lục cung cấp cấu trúc dự án, cấu hình Phase 2 và các hình minh họa kết quả phát hiện đối tượng.",
    )

    # Chapter 2: remove the former attention subsection and renumber transfer learning.
    remove_range(doc, "2.3.3. Cơ chế chú ý", "2.3.4. Học chuyển giao")
    replace_all_prefix(doc, "2.3.4. Học chuyển giao", "2.3.3. Học chuyển giao (Transfer Learning)")

    # Chapter 3: clean data narrative and reorganize the training strategy.
    replace_all_prefix(
        doc,
        "Kiến trúc hệ thống được chia thành bốn tầng",
        "Kiến trúc hệ thống gồm bốn khối liên kết: dữ liệu và tiền xử lý; các mô hình benchmark; huấn luyện–lựa chọn mô hình; Phase 2 và đánh giá cuối. Dữ liệu được chuẩn hóa về cùng sáu lớp trước khi đi qua Stage 1 và Stage 2. XWOD validation được dùng để lựa chọn RT-DETR-L; cấu hình này tiếp tục fine-tuning đa miền và chỉ được đưa sang frozen test sau khi đã cố định.",
    )
    replace_all_prefix(
        doc,
        "Tầng dữ liệu đảm nhiệm đọc ảnh",
        "Khối dữ liệu đảm nhiệm đọc ảnh, chuyển đổi nhãn và tổ chức train, validation, test. Khối mô hình gồm YOLOv8n, YOLO11n, Faster R-CNN và RT-DETR-L. Khối huấn luyện điều phối COCO pretrained → BDD100K → XWOD và tách rõ validation khỏi test. Khối cuối sử dụng RT-DETR-L Phase 2 để đánh giá frozen trên XWOD, BDD, DAWN và ACDC.",
    )
    replace_all_prefix(
        doc,
        "Protocol chính thức ghi nhận XWOD",
        "XWOD gồm 6.006 ảnh train, 1.001 ảnh validation và 3.003 ảnh test. Các ảnh và nhãn được kiểm tra khả năng giải mã, tính hợp lệ của hộp giới hạn và sự nhất quán của class ID trước khi tạo các split dùng trong huấn luyện và đánh giá.",
    )
    replace_all_prefix(
        doc,
        "Run Phase 2 sử dụng RT-DETR-L",
        "Phase 2 sử dụng RT-DETR-L được khởi tạo từ best checkpoint Stage 2. Ba nguồn train gồm XWOD 6.006 ảnh, ACDC 1.182 ảnh và BDD100K 30.000 ảnh, tạo tổng cơ sở 37.188 ảnh cho joint multi-domain fine-tuning.",
    )
    replace_all_prefix(
        doc,
        "Báo cáo sử dụng số đếm chính xác",
        "XWOD duy trì chuyên biệt hóa đối với thời tiết bất lợi; ACDC mở rộng điều kiện môi trường và ánh sáng; BDD100K giúp duy trì tri thức giao thông đã học ở Stage 1. Ba nguồn được tổ chức theo cùng không gian sáu lớp.",
    )
    replace_all_prefix(
        doc,
        "Thiết kế hiện tại mô tả lấy mẫu tăng cường",
        "Do car và person xuất hiện phổ biến hơn, phân tích theo lớp được dùng để theo dõi riêng bicycle, motorcycle và bus. Việc duy trì thống kê theo lớp giúp tránh trường hợp chỉ số trung bình che khuất hạn chế trên các lớp ít xuất hiện.",
    )
    replace_all_prefix(doc, "3.4. Lựa chọn mô hình và thiết kế quy trình huấn luyện", "3.4. Các mô hình benchmark")
    replace_all_prefix(doc, "3.4.1. Các kiến trúc được đưa vào benchmark", "3.4.1. Bốn kiến trúc so sánh")
    replace_all_prefix(
        doc,
        "Re-evaluation trên XWOD validation",
        "Bốn mô hình sử dụng cùng không gian sáu lớp và cùng vai trò dữ liệu ở Stage 1–Stage 2. XWOD validation là nguồn duy nhất dùng cho quyết định lựa chọn kiến trúc trước Phase 2; các tập test chỉ được dùng sau khi cấu hình đã cố định.",
    )
    stage_heading = find_one(doc, "3.4.2. Quy trình học chuyển giao", style_prefix="heading")
    set_text(stage_heading, "3.5. Chiến lược huấn luyện nhiều giai đoạn")
    stage_heading.style = doc.styles["Heading 2"]
    stage1 = find_one(doc, "Các mô hình không được huấn luyện từ đầu")
    insert_before(stage1, "3.5.1. Stage 1 — thích nghi miền giao thông", "Heading 33")
    set_text(
        stage1,
        "Các mô hình được khởi tạo từ trọng số COCO pretrained và tinh chỉnh trên BDD100K train 30.000 ảnh. Stage 1 giúp mô hình thích nghi với cảnh đường phố và phân bố sáu lớp mục tiêu trước khi tiếp xúc với miền thời tiết bất lợi.",
    )
    stage2 = find_one(doc, "Stage 2 sử dụng XWOD train")
    insert_before(stage2, "3.5.2. Stage 2 — thích nghi miền thời tiết bất lợi", "Heading 33")
    set_text(
        stage2,
        "Stage 2 tiếp tục từ best checkpoint Stage 1 và tinh chỉnh trên XWOD train 6.006 ảnh. XWOD validation phục vụ early stopping và lựa chọn checkpoint; XWOD test không tham gia tối ưu và chỉ được dùng cho frozen evaluation.",
    )
    phase2 = find_one(doc, "Sau khi được lựa chọn bằng XWOD validation")
    insert_before(phase2, "3.5.3. Lựa chọn mô hình bằng XWOD validation", "Heading 33")
    insert_before(
        phase2,
        "RT-DETR-L đạt mAP50-95 0,559 trên XWOD validation, cao hơn YOLO11n 0,506, YOLOv8n 0,503 và Faster R-CNN 0,421. Vì vậy RT-DETR-L được lựa chọn để tiếp tục Phase 2.",
        "Normal",
    )
    insert_before(phase2, "3.5.4. Phase 2 — joint multi-domain fine-tuning", "Heading 33")
    set_text(
        phase2,
        "Phase 2 khởi tạo từ best checkpoint RT-DETR-L Stage 2 và fine-tune đồng thời trên XWOD train, ACDC train và BDD100K train. XWOD duy trì khả năng phát hiện trong thời tiết bất lợi; ACDC bổ sung điều kiện môi trường và ánh sáng; BDD100K hỗ trợ hạn chế quên miền giao thông.",
    )
    remove_paragraph(find_one(doc, "3.4.3. Kiểm soát quá trình huấn luyện", style_prefix="heading"))
    replace_all_prefix(
        doc,
        "Trong mỗi lần huấn luyện, best checkpoint",
        "Best checkpoint của từng giai đoạn được chọn theo validation. Phase 2 chạy tối đa 50 epoch, đạt mAP50-95 validation cao nhất 0,562 tại epoch 13 và dừng sau epoch 33 theo patience 20; tổng thời gian huấn luyện là 8,803 giờ.",
    )
    remove_range(doc, "3.5. Thiết kế các hướng cải tiến", "3.6. Thiết kế thực nghiệm")
    replace_all_prefix(doc, "3.6. Thiết kế thực nghiệm và bảo đảm khả năng tái lập", "3.6. Giao thức thực nghiệm")
    replace_all_prefix(doc, "3.6.1. Ma trận thực nghiệm", "3.6.1. Phân tách train, validation và test")
    replace_all_prefix(
        doc,
        "Các nhóm thực nghiệm trả lời những câu hỏi riêng",
        "Train được dùng để tối ưu tham số; validation phục vụ early stopping, chọn best checkpoint và lựa chọn RT-DETR-L; test chỉ được dùng cho đánh giá frozen sau khi cấu hình đã cố định. Bảng 3.6 tóm tắt vai trò của ba giai đoạn thực nghiệm chính.",
    )
    replace_all_prefix(doc, "Bảng 3.6. Ma trận các nhóm thực nghiệm", "Bảng 3.6. Ma trận các giai đoạn thực nghiệm")
    replace_all_prefix(doc, "3.6.2. Giao thức đánh giá và diễn giải kết quả", "3.6.2. Độ đo và nguyên tắc diễn giải")
    replace_all_prefix(doc, "3.6.3. Quản lý kết quả và giới hạn bằng chứng", "3.6.3. Khả năng tái lập và tính công bằng")
    replace_all_prefix(
        doc,
        "Mỗi kết quả được gắn với tên run",
        "Mỗi lần huấn luyện được lưu cùng cấu hình dữ liệu, kích thước ảnh, seed, bộ tối ưu, lịch sử theo epoch và checkpoint tương ứng. Các bảng kết quả sử dụng cùng thứ tự sáu lớp, ghi rõ split và tách riêng validation khỏi frozen test.",
    )
    replace_all_prefix(
        doc,
        "Cấu hình cuối được xác nhận là RT-DETR-L Phase 2",
        "Cấu hình cuối là RT-DETR-L Phase 2 ở kích thước ảnh 640. YOLO và RT-DETR được triển khai bằng Ultralytics, còn Faster R-CNN dùng torchvision/torchmetrics; vì vậy mAP50-95 là chỉ số chính khi so sánh accuracy, còn Precision/Recall và latency/FPS giữa framework được dùng ở mức tham khảo.",
    )
    replace_all_prefix(
        doc,
        "Từ thiết kế trên, Chương 4 trình bày",
        "Từ thiết kế trên, Chương 4 trình bày môi trường và cấu hình, benchmark bốn mô hình, Phase 2 của RT-DETR-L, frozen evaluation, phân tích theo lớp và thời tiết, cùng các hạn chế và hướng phát triển.",
    )
    replace_all_prefix(
        doc,
        "Quy trình huấn luyện được xây dựng theo hướng lũy tiến",
        "Quy trình huấn luyện được xây dựng theo hướng lũy tiến: COCO pretrained → BDD100K → XWOD. RT-DETR-L được lựa chọn bằng XWOD validation và tiếp tục bằng joint multi-domain fine-tuning trên XWOD, ACDC và BDD100K. DAWN không tham gia huấn luyện hoặc lựa chọn mô hình.",
    )
    replace_all_prefix(
        doc,
        "Ảnh được thay đổi kích thước bằng letterbox",
        "Ảnh được thay đổi kích thước bằng letterbox. Phép biến đổi giữ nguyên tỉ lệ khung hình, thu phóng ảnh theo cạnh giới hạn rồi đệm phần còn thiếu; nhãn hộp được biến đổi theo cùng hệ số và độ lệch đệm. Cách làm này tránh kéo giãn người hoặc phương tiện. Kích thước ảnh 640 được sử dụng cho benchmark và RT-DETR-L Phase 2.",
    )

    # Chapter 4: clean scientific narrative and make Phase 2 the final stage.
    replace_all_prefix(
        doc,
        "Chương này sử dụng artifact",
        "Chương này trình bày môi trường triển khai, cấu hình huấn luyện và kết quả đánh giá của quy trình nhiều giai đoạn. RT-DETR-L được lựa chọn bằng XWOD validation, tiếp tục ở Phase 2 và được xem là mô hình cuối cùng trong phạm vi thực nghiệm của đề tài.",
    )
    replace_all_prefix(
        doc,
        "Pipeline thực nghiệm gồm pretrained",
        "Pipeline thực nghiệm gồm COCO pretrained → Stage 1 trên BDD100K → Stage 2 trên XWOD → XWOD validation → Phase 2 multi-domain fine-tuning. DAWN không tham gia huấn luyện; ACDC test sau Phase 2 là held-out same-source evaluation; cả bốn tập test chỉ được sử dụng sau khi cấu hình đã được cố định.",
    )
    replace_all_prefix(
        doc,
        "Môi trường được ghi nhận sử dụng",
        "Môi trường thực nghiệm sử dụng NVIDIA GeForce RTX 5090 với khoảng 32 GB VRAM. PyTorch 2.7.1+cu128 và torchvision 0.22.1+cu128 sử dụng CUDA build/runtime 12.8; driver thuộc nhánh 595.x và hỗ trợ CUDA tối đa 13.2. Ultralytics có phiên bản 8.3.159 và Python là 3.12.13.",
    )
    replace_all_prefix(doc, "Bảng 4.1. Môi trường thực nghiệm được xác nhận", "Bảng 4.1. Môi trường phần cứng và phần mềm")
    replace_all_prefix(
        doc,
        "Stage 1 sử dụng BDD100K train",
        "Stage 1 sử dụng BDD100K train 30.000 ảnh; Stage 2 tiếp tục từ best checkpoint Stage 1 và tinh chỉnh trên XWOD train 6.006 ảnh. Phase 2 dùng RT-DETR-L và joint multi-domain fine-tuning trên XWOD 6.006 ảnh, ACDC 1.182 ảnh và BDD100K 30.000 ảnh, với tổng cơ sở 37.188 ảnh.",
    )
    replace_all_prefix(
        doc,
        "Artifact Stage 1 xác nhận",
        "Kết quả Stage 1 trên BDD test cho phép so sánh trực tiếp bốn kiến trúc. Các kết quả ngoài miền được dùng để mô tả khả năng chuyển miền của từng cấu hình, không thay thế vai trò của validation trong lựa chọn mô hình.",
    )
    replace_all_prefix(
        doc,
        "Sau khi tinh chỉnh trên XWOD, RT-DETR-L đạt",
        "Sau khi tinh chỉnh trên XWOD, RT-DETR-L đạt mAP50-95 0,500 trên XWOD test; YOLO11n đạt 0,454, YOLOv8n 0,452 và Faster R-CNN 0,381. Đây là frozen-test benchmark; quyết định chuyển RT-DETR-L sang Phase 2 dựa trên XWOD validation.",
    )
    remove_range(doc, "4.3. Trạng thái bằng chứng", "4.4. Phase 2")
    replace_all_prefix(doc, "4.4. Phase 2 và kết quả của mô hình cuối cùng", "4.3. Phase 2 và mô hình cuối cùng")
    replace_all_prefix(doc, "4.4.1. Ảnh hưởng của dữ liệu gộp và phát lại BDD", "4.3.1. Thiết kế multi-domain fine-tuning")
    phase2_design = find_one(doc, "Run Phase 2 dùng tối đa")
    set_text(
        phase2_design,
        "Phase 2 khởi tạo từ best checkpoint RT-DETR-L Stage 2 và fine-tune đồng thời trên XWOD, ACDC và BDD100K. Cách tổ chức này nhằm duy trì chuyên biệt hóa trên XWOD, mở rộng điều kiện môi trường bằng ACDC và hạn chế suy giảm tri thức giao thông bằng BDD100K.",
    )
    phase2_bdd = find_after(doc, phase2_design, "Phase 2 khởi tạo từ best checkpoint RT-DETR-L Stage 2")
    insert_before(phase2_bdd, "4.3.2. Cấu hình huấn luyện", "Heading 33")
    insert_before(
        phase2_bdd,
        "RT-DETR-L Phase 2 sử dụng ảnh 640×640, batch 16, AdamW, lr0 = 5×10⁻⁵, patience 20 và seed 42. Run có tối đa 50 epoch, đạt best epoch 13 với P = 0,827, R = 0,810, mAP50 = 0,848 và mAP50-95 = 0,562 trên XWOD validation; huấn luyện dừng ở epoch 33 sau 8,803 giờ.",
        "Normal",
    )
    insert_before(phase2_bdd, "4.3.3. Frozen-test evaluation", "Heading 33")
    set_text(
        phase2_bdd,
        "Mô hình cuối cùng đạt P/R/mAP50/mAP50-95 lần lượt 0,806/0,745/0,779/0,510 trên XWOD test; 0,698/0,558/0,594/0,354 trên BDD test; 0,835/0,766/0,795/0,515 trên DAWN test; và 0,600/0,384/0,407/0,247 trên ACDC test.",
    )
    replace_all_prefix(
        doc,
        "Trên ACDC test, mAP50-95 tăng từ 0,128",
        "So với Stage 2, Phase 2 đưa BDD mAP50-95 từ 0,183 lên 0,354 và ACDC từ 0,128 lên 0,247; XWOD tăng nhẹ từ 0,500 lên 0,510, còn DAWN giảm từ 0,526 xuống 0,515. Vì vậy Phase 2 không cải thiện đồng đều trên mọi miền.",
    )
    replace_all_prefix(doc, "4.4.2. So sánh RT-DETR-L qua các giai đoạn", "4.4.1. So sánh RT-DETR-L qua các giai đoạn")
    replace_all_prefix(doc, "4.4.3. Phân tích theo lớp", "4.4.2. Phân tích theo lớp")
    replace_all_prefix(doc, "4.4.4. Phân tích theo điều kiện thời tiết", "4.4.3. Phân tích theo điều kiện thời tiết")
    analysis_heading = find_one(doc, "4.4.1. So sánh RT-DETR-L qua các giai đoạn", style_prefix="heading")
    insert_before(analysis_heading, "4.4. Phân tích kết quả", "Heading 2")
    replace_all_prefix(
        doc,
        "Kết quả Phase 2 thay đổi nhiều nhất",
        "Qua ba giai đoạn, Stage 1 tạo nền tảng miền giao thông; Stage 2 tăng chuyên biệt hóa trên XWOD nhưng làm BDD giảm mạnh; Phase 2 phục hồi BDD gần mức Stage 1 và cải thiện ACDC. XWOD tăng nhẹ, trong khi DAWN giảm nhẹ so với Stage 2.",
    )
    replace_all_prefix(
        doc,
        "Trên ACDC test, bicycle và motorcycle",
        "Trên ACDC test, bicycle và motorcycle có mAP50-95 lần lượt 0,116 và 0,119, thấp nhất trong sáu lớp; bus đạt 0,190. Những lớp này thường có ít mẫu hơn, kích thước nhỏ, dễ bị che khuất và chịu ảnh hưởng mạnh của suy giảm ảnh. Trên XWOD, bus đạt 0,396 và vẫn còn dư địa cải thiện dù bicycle đạt 0,652.",
    )
    replace_all_prefix(doc, "4.6.1. Các phát hiện chính và lựa chọn mô hình", "4.6.1. Các phát hiện chính")
    replace_all_prefix(
        doc,
        "RT-DETR-L được chọn bằng XWOD validation",
        "RT-DETR-L đạt mAP50-95 cao nhất trên XWOD validation và được lựa chọn cho Phase 2. YOLOv8n và YOLO11n có tốc độ cao hơn nhưng mAP50-95 thấp hơn. Stage 2 cải thiện thích nghi thời tiết bất lợi nhưng gây suy giảm BDD; Phase 2 phục hồi phần lớn mức giảm này, cải thiện ACDC, giữ XWOD ổn định và giảm nhẹ trên DAWN. RT-DETR-L Phase 2 là mô hình cuối cùng của đề tài.",
    )
    replace_all_prefix(
        doc,
        "Phase 2 phục hồi đáng kể kết quả BDD",
        "Kết quả cho thấy fine-tuning đa miền tạo ra sự cân bằng tốt hơn giữa miền giao thông và miền thời tiết bất lợi, nhưng chưa giải quyết đồng đều mọi điều kiện. Kết luận của đề tài vì vậy tập trung vào hiệu quả của cấu hình RT-DETR-L Phase 2 thay vì quy nguyên nhân cho một nguồn dữ liệu đơn lẻ.",
    )
    replace_all_prefix(doc, "4.6.2. Hạn chế và hướng phát triển", "4.6.2. Hạn chế")
    limitations1 = find_one(doc, "Metadata Phase 2 đã được khôi phục")
    set_text(
        limitations1,
        "Mô hình cuối vẫn có recall thấp trên ACDC, đặc biệt trong điều kiện ban đêm; bicycle, motorcycle và bus còn nhạy với số lượng mẫu, kích thước nhỏ, che khuất và chênh lệch miền. DAWN giảm nhẹ sau Phase 2, cho thấy cân bằng nhiều miền vẫn là một bài toán khó.",
    )
    limitations2 = find_one(doc, "Hướng phát triển ưu tiên")
    set_text(
        limitations2,
        "RT-DETR-L đạt độ chính xác cao nhất trong benchmark nhưng có tốc độ suy luận thấp hơn các mô hình YOLO nano. Vì vậy, chi phí tính toán và khả năng triển khai trên thiết bị biên là hạn chế cần được xem xét cùng với độ chính xác.",
    )
    references_heading = find_one(doc, "TÀI LIỆU THAM KHẢO", style_prefix="heading")
    insert_before(references_heading, "4.6.3. Hướng phát triển", "Heading 33")
    future_paragraphs = [
        "Cải thiện lớp hiếm. Một hướng tiếp theo là khai thác phần BDD100K train chưa sử dụng để bổ sung có kiểm soát các ảnh chứa bicycle, motorcycle và bus. A0R có thể được xây dựng như một đối chứng random rare-class data-selection: xác lập candidate pool, chọn ngẫu nhiên một ngân sách ảnh cố định chứa lớp hiếm và fine-tune lại mô hình. Đối chứng này giúp tách ảnh hưởng của việc tăng dữ liệu khỏi ảnh hưởng của chiến lược lựa chọn mẫu. Class-aware sampling và object-level augmentation cũng có thể được khảo sát trong cùng giao thức.",
        "Lựa chọn dữ liệu dựa trên failure case. Thay vì chọn ngẫu nhiên, A1-DINO có thể sử dụng RT-DETR-L Phase 2 để xác định các mẫu khó, tập trung vào lớp hiếm hoặc bị bỏ sót, rồi tìm ảnh tương đồng trong candidate pool BDD100K đã có nhãn. Đây là failure-driven similarity retrieval, không phải active learning theo nghĩa nghiêm ngặt.",
        "DINOv2 có thể đóng vai trò cơ chế biểu diễn và truy hồi hình ảnh cho A1-DINO. Embedding tổng quát của các mẫu khó và candidate pool được so sánh bằng cosine similarity để xếp hạng ảnh nguồn; sau đó một ngân sách cố định được chọn để fine-tune lại RT-DETR-L. Cách phân vai này cho phép RT-DETR xác định mô hình thất bại ở đâu, còn DINOv2 xác định ảnh nguồn nào có đặc trưng gần với failure case. So sánh A0R với A1-DINO sẽ cho biết lựa chọn có chủ đích có tạo lợi ích vượt quá việc chỉ tăng dữ liệu lớp hiếm hay không.",
        "Cải thiện điều kiện ban đêm. Có thể bổ sung dữ liệu night thật, áp dụng photometric nighttime augmentation, tạo dữ liệu ban đêm tổng hợp hoặc fine-tune theo miền ánh sáng. Các phương án cần được đánh giá riêng để phân biệt lợi ích của lượng dữ liệu với lợi ích của kiểu biến đổi ảnh.",
        "Giảm quên miền và tối ưu triển khai. Các hướng phù hợp gồm replay có kiểm soát, balanced multi-domain sampling, regularization và continual/domain-adaptive fine-tuning. Đối với triển khai, có thể nghiên cứu distillation, quantization, TensorRT, biến thể RT-DETR nhẹ hơn và đánh giá trên thiết bị biên.",
    ]
    for text in future_paragraphs:
        insert_before(references_heading, text, "Normal")

    # References: remove the former attention-only citation and close numbering gap.
    remove_paragraph(find_one(doc, "[7] J. Hu"))
    for old in range(8, 17):
        paragraph = find_one(doc, f"[{old}] ")
        paragraph.text = re.sub(rf"^\[{old}\]", f"[{old - 1}]", paragraph.text)

    # Appendix: presentation-oriented structure and prose.
    replace_all_prefix(
        doc,
        "Cấu trúc thư mục chính của dự án",
        "Cấu trúc dự án được tổ chức theo các nhóm dữ liệu, cấu hình, huấn luyện, đánh giá và mô hình cuối. Cách phân tách này hỗ trợ tái lập từng giai đoạn và tránh trộn lẫn kết quả validation với frozen test.",
    )
    replace_appendix_tree(doc)
    replace_all_prefix(
        doc,
        "Bảng dưới đây tổng hợp cấu hình RT-DETR-L Phase 2",
        "Bảng dưới đây tổng hợp cấu hình huấn luyện của RT-DETR-L Phase 2, là mô hình cuối cùng của đề tài.",
    )
    replace_all_prefix(
        doc,
        "Các hình trong phụ lục chỉ có vai trò minh họa",
        "Các hình trong phụ lục minh họa đầu ra phát hiện của mô hình trên XWOD, DAWN và ACDC. Kết luận định lượng được trình bày trong các bảng đánh giá ở Chương 4.",
    )

    # Tables: abbreviations, Phase 2 data, protocol, environment and appendix config.
    abbreviation_rows = [[c.text for c in row.cells] for row in original_tables[0].rows if row.cells[0].text.strip() != "CBAM"]
    replace_table(doc, original_tables[0], abbreviation_rows, size=11)
    table34 = replace_table(
        doc,
        original_tables[7],
        [
            ["Nguồn train", "Số ảnh", "Vai trò trong Phase 2"],
            ["XWOD", "6.006", "Duy trì chuyên biệt hóa thời tiết bất lợi"],
            ["ACDC", "1.182", "Mở rộng điều kiện môi trường và ánh sáng"],
            ["BDD100K", "30.000", "Duy trì tri thức miền giao thông"],
            ["Tổng cơ sở", "37.188", "Joint multi-domain fine-tuning"],
        ],
        size=10.5,
    )
    table36 = replace_table(
        doc,
        original_tables[9],
        [
            ["Giai đoạn", "Dữ liệu train", "Validation", "Đánh giá sau khi khóa cấu hình"],
            ["Stage 1", "BDD100K train", "BDD100K val", "BDD test; đánh giá chuyển miền"],
            ["Stage 2", "XWOD train", "XWOD val", "XWOD/BDD/DAWN/ACDC test"],
            ["Phase 2", "XWOD + ACDC + BDD100K train", "XWOD val", "XWOD/BDD/DAWN/ACDC frozen test"],
        ],
        size=10.5,
    )
    replace_table(
        doc,
        original_tables[8],
        [
            ["Mô hình", "Hướng tiếp cận", "Vai trò", "Đánh đổi chính"],
            ["YOLOv8n", "Một giai đoạn", "Mốc tốc độ", "Nhanh, gọn; mAP thấp hơn RT-DETR-L"],
            ["YOLO11n", "Một giai đoạn", "YOLO thế hệ mới", "mAP nhỉnh hơn YOLOv8n; FPS thấp hơn"],
            ["Faster R-CNN", "Hai giai đoạn", "Đối chứng proposal", "Recall cao ở một số tập; precision/mAP thấp hơn"],
            ["RT-DETR-L", "Transformer", "Mô hình chọn cho Phase 2", "Accuracy cao nhất; tốc độ thấp hơn YOLO nano"],
        ],
        size=9.5,
    )
    environment = [[c.text for c in row.cells] for row in original_tables[11].rows]
    environment[0][1] = "Giá trị"
    environment[1][1] = "NVIDIA GeForce RTX 5090, khoảng 32 GB"
    environment[2][1] = "Driver 595.x / hỗ trợ CUDA tối đa 13.2"
    replace_table(doc, original_tables[11], environment, size=10.5)
    set_table_font(original_tables[18], 10)
    replace_table(
        doc,
        original_tables[20],
        [
            ["Thành phần", "Cấu hình Phase 2"],
            ["Mô hình", "RT-DETR-L từ Stage 2 best checkpoint"],
            ["Nguồn train", "XWOD + ACDC + BDD100K"],
            ["Epoch tối đa / hoàn thành", "50 / 33"],
            ["Best epoch", "13 (mAP50-95 val = 0,562)"],
            ["Early stopping", "patience = 20"],
            ["Kích thước ảnh", "640 × 640"],
            ["Batch size / seed", "16 / 42"],
            ["Optimizer / lr0", "AdamW / 5×10⁻⁵"],
            ["Thời gian", "8,803 giờ"],
        ],
        size=10.5,
    )

    # Manual TOC/list cleanup. Page numbers are repaired from the final render later.
    forbidden_toc = (
        "2.3.3. Cơ chế chú ý",
        "3.5.1. Khảo sát cơ chế chú ý",
        "3.5. Thiết kế các hướng cải tiến",
        "4.3. Trạng thái bằng chứng",
        "4.3.1. Thiết kế khảo sát",
        "4.3.2. Giới hạn bằng chứng",
        "Hình 2.5.",
        "Hình 3.3.",
    )
    for paragraph in list(doc.paragraphs):
        if paragraph.style.name.lower().startswith("toc") and paragraph.text.strip().startswith(forbidden_toc):
            remove_paragraph(paragraph)
    toc_map = {
        "2.3.4. Học chuyển giao": "2.3.3. Học chuyển giao (Transfer Learning)",
        "3.4. Lựa chọn mô hình và thiết kế quy trình huấn luyện": "3.4. Các mô hình benchmark",
        "3.4.1. Các kiến trúc được đưa vào benchmark": "3.4.1. Bốn kiến trúc so sánh",
        "3.4.2. Quy trình học chuyển giao nhiều giai đoạn": "3.5. Chiến lược huấn luyện nhiều giai đoạn",
        "3.4.3. Kiểm soát quá trình huấn luyện": "3.5.1. Stage 1 — thích nghi miền giao thông",
        "3.5.2. Mở rộng dữ liệu Phase 2 và lấy mẫu lớp hiếm": "3.5.2. Stage 2 — thích nghi miền thời tiết bất lợi",
        "3.6. Thiết kế thực nghiệm và bảo đảm khả năng tái lập": "3.6. Giao thức thực nghiệm",
        "3.6.1. Ma trận thực nghiệm": "3.6.1. Phân tách train, validation và test",
        "3.6.2. Giao thức đánh giá và diễn giải kết quả": "3.6.2. Độ đo và nguyên tắc diễn giải",
        "3.6.3. Quản lý kết quả và giới hạn bằng chứng": "3.6.3. Khả năng tái lập và tính công bằng",
        "4.4. Phase 2 và kết quả của mô hình cuối cùng": "4.3. Phase 2 và mô hình cuối cùng",
        "4.4.1. Ảnh hưởng của dữ liệu gộp và phát lại BDD": "4.3.1. Thiết kế multi-domain fine-tuning",
        "4.4.2. So sánh RT-DETR-L qua các giai đoạn": "4.4.1. So sánh RT-DETR-L qua các giai đoạn",
        "4.4.3. Phân tích theo lớp": "4.4.2. Phân tích theo lớp",
        "4.4.4. Phân tích theo điều kiện thời tiết": "4.4.3. Phân tích theo điều kiện thời tiết",
        "4.6.1. Các phát hiện chính và lựa chọn mô hình": "4.6.1. Các phát hiện chính",
        "4.6.2. Hạn chế và hướng phát triển": "4.6.2. Hạn chế",
        "Bảng 4.1. Môi trường thực nghiệm được xác nhận từ artifact Drive": "Bảng 4.1. Môi trường phần cứng và phần mềm",
        "Bảng 3.6. Ma trận các nhóm thực nghiệm": "Bảng 3.6. Ma trận các giai đoạn thực nghiệm",
    }
    for paragraph in doc.paragraphs:
        if not paragraph.style.name.lower().startswith("toc"):
            continue
        title = paragraph.text.rsplit("\t", 1)[0]
        for old, new in toc_map.items():
            if title.startswith(old):
                paragraph.text = new
                break
    toc_stage2 = find_one(doc, "3.5.2. Stage 2", style_prefix="toc")
    toc_stage = find_one(doc, "3.5. Chiến lược huấn luyện nhiều giai đoạn", style_prefix="toc")
    toc_stage.style = doc.styles["toc 2"]
    copy_paragraph_properties(toc_stage, find_one(doc, "3.4. Các mô hình benchmark", style_prefix="toc"))
    following = doc.paragraphs[paragraph_index(doc, toc_stage2) + 1]
    insert_toc_before(following, "3.5.3. Lựa chọn mô hình bằng XWOD validation", toc_stage2)
    insert_toc_before(following, "3.5.4. Phase 2 — joint multi-domain fine-tuning", toc_stage2)
    toc_phase2 = find_one(doc, "4.3.1. Thiết kế multi-domain fine-tuning", style_prefix="toc")
    following = doc.paragraphs[paragraph_index(doc, toc_phase2) + 1]
    insert_toc_before(following, "4.3.2. Cấu hình huấn luyện", toc_phase2)
    insert_toc_before(following, "4.3.3. Frozen-test evaluation", toc_phase2)
    toc_analysis = find_one(doc, "4.4.1. So sánh RT-DETR-L qua các giai đoạn", style_prefix="toc")
    insert_toc_before(toc_analysis, "4.4. Phân tích kết quả", find_one(doc, "4.3. Phase 2 và mô hình cuối cùng", style_prefix="toc"))
    toc_limit = find_one(doc, "4.6.2. Hạn chế", style_prefix="toc")
    following = doc.paragraphs[paragraph_index(doc, toc_limit) + 1]
    insert_toc_before(following, "4.6.3. Hướng phát triển", toc_limit)

    # Replace the two thesis-specific main diagrams.
    replace_image_before_caption(doc, "Hình 3.1.", assets / "system_architecture.png")
    replace_image_before_caption(doc, "Hình 3.2.", assets / "training_pipeline.png")

    # Start the long configuration table on a fresh page so it is not split awkwardly.
    find_one(doc, "Bảng 3.5. Đặc điểm và vai trò", style_prefix="caption").paragraph_format.page_break_before = True
    find_one(doc, "Bảng 4.2. Tóm tắt cấu hình", style_prefix="caption").paragraph_format.page_break_before = True

    # Final pagination safeguards.
    for table in doc.tables:
        shade_header(table)
        for row in table.rows:
            tr_pr = row._tr.get_or_add_trPr()
            if tr_pr.find(qn("w:cantSplit")) is None:
                tr_pr.append(OxmlElement("w:cantSplit"))
    for paragraph in doc.paragraphs:
        if paragraph.style.name.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.keep_together = True
        if paragraph.style.name.startswith("Caption"):
            paragraph.paragraph_format.keep_with_next = True

    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--assets", type=Path, required=True)
    args = parser.parse_args()
    revise(args.source, args.output, args.assets)


if __name__ == "__main__":
    main()
