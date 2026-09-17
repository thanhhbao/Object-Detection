#!/usr/bin/env python3
"""Reorder front-matter lists and add preliminary entries to the cached TOC."""

from __future__ import annotations

import argparse
import copy
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph


FRONT_TOC = [
    ("LỜI MỞ ĐẦU", "i"),
    ("LỜI CẢM ƠN", "iii"),
    ("NHẬN XÉT CỦA GIẢNG VIÊN HƯỚNG DẪN", "iv"),
    ("NHẬN XÉT CỦA GIẢNG VIÊN PHẢN BIỆN", "v"),
    ("MỤC LỤC", "vi"),
    ("DANH MỤC BẢNG BIỂU", "x"),
    ("DANH MỤC HÌNH", "xi"),
    ("DANH MỤC CÁC TỪ VIẾT TẮT", "xii"),
]


def element_text(element) -> str:
    return "".join(element.xpath(".//w:t/text()")).strip()


def find_child(children, exact_text: str) -> int:
    for idx, child in enumerate(children):
        if child.tag == qn("w:p") and element_text(child) == exact_text:
            return idx
    raise ValueError(f"Paragraph not found: {exact_text}")


def replace_paragraph_text(doc: Document, old: str, new: str) -> Paragraph:
    paragraph = next(p for p in doc.paragraphs if p.text.strip() == old)
    paragraph.text = new
    paragraph.alignment = 1
    paragraph.paragraph_format.page_break_before = True
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.left_indent = Cm(0)
    paragraph.paragraph_format.first_line_indent = Cm(0)
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(16)
        run.font.bold = True
    return paragraph


def configure_toc_entry(paragraph: Paragraph, text: str, page: str, toc_style) -> None:
    paragraph.text = f"{text}\t{page}"
    paragraph.style = toc_style
    fmt = paragraph.paragraph_format
    fmt.tab_stops.clear_all()
    fmt.tab_stops.add_tab_stop(Cm(15.45), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    fmt.keep_together = True
    fmt.widow_control = True
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(13)


def insert_before(reference, new_element) -> None:
    reference.addprevious(new_element)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--backup", type=Path)
    args = parser.parse_args()

    if args.backup:
        args.backup.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(args.source, args.backup)

    doc = Document(args.source)
    body = doc._element.body
    children = list(body)
    i_abbrev = find_child(children, "DANH MỤC CHỮ VIẾT TẮT")
    i_figures = find_child(children, "DANH MỤC HÌNH")
    i_tables = find_child(children, "DANH MỤC BẢNG")
    i_chapter = next(
        i for i, el in enumerate(children)
        if i > i_tables and el.tag == qn("w:p") and element_text(el).startswith("CHƯƠNG 1")
    )

    abbrev_block = children[i_abbrev:i_figures]
    figures_block = children[i_figures:i_tables]
    tables_block = children[i_tables:i_chapter]
    chapter_element = children[i_chapter]

    # Preserve the section break before Chapter 1, but attach it to the final
    # preliminary block after reordering rather than carrying it with tables.
    section_properties = None
    for element in tables_block:
        if element.tag != qn("w:p"):
            continue
        ppr = element.find(qn("w:pPr"))
        sect = ppr.find(qn("w:sectPr")) if ppr is not None else None
        if sect is not None:
            section_properties = copy.deepcopy(sect)
            ppr.remove(sect)
    if section_properties is None:
        raise RuntimeError("Section break before Chapter 1 was not found")

    for element in abbrev_block + figures_block + tables_block:
        body.remove(element)

    for element in tables_block + figures_block + abbrev_block:
        insert_before(chapter_element, element)

    # The abbreviation block ends with a blank paragraph following its table.
    final_prelim_paragraph = next((el for el in reversed(abbrev_block) if el.tag == qn("w:p")), None)
    if final_prelim_paragraph is None:
        final_prelim_paragraph = OxmlElement("w:p")
        insert_before(chapter_element, final_prelim_paragraph)
    ppr = final_prelim_paragraph.find(qn("w:pPr"))
    if ppr is None:
        ppr = OxmlElement("w:pPr")
        final_prelim_paragraph.insert(0, ppr)
    ppr.append(section_properties)

    replace_paragraph_text(doc, "DANH MỤC BẢNG", "DANH MỤC BẢNG BIỂU")
    replace_paragraph_text(doc, "DANH MỤC CHỮ VIẾT TẮT", "DANH MỤC CÁC TỪ VIẾT TẮT")

    # Insert the requested preliminary entries before the first chapter entry.
    toc_title = next(p for p in doc.paragraphs if p.text.strip() == "MỤC LỤC")
    first_body_entry = next(
        p for p in doc.paragraphs
        if p.style.name.startswith("toc") and p.text.strip().startswith("CHƯƠNG 1")
    )
    toc_style = doc.styles["toc 1"]
    for text, page in FRONT_TOC:
        new_p = OxmlElement("w:p")
        insert_before(first_body_entry._p, new_p)
        configure_toc_entry(Paragraph(new_p, first_body_entry._parent), text, page, toc_style)

    # Mark all fields for refresh when opened in Word.
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")

    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
