#!/usr/bin/env python3
"""Create the V3 format/structure thesis without overwriting the V2 source."""

from __future__ import annotations

import argparse
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def find_exact(doc: Document, text: str):
    matches = [p for p in doc.paragraphs if p.text.strip() == text]
    if len(matches) != 1:
        raise RuntimeError(f"Expected exactly one paragraph {text!r}, found {len(matches)}")
    return matches[0]


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def paragraph_index(doc: Document, target) -> int:
    return next(i for i, paragraph in enumerate(doc.paragraphs) if paragraph._p is target._p)


def add_field(paragraph, instruction: str, placeholder: str = "Nhấn F9 để cập nhật trường") -> None:
    """Append a complex Word field to a paragraph."""
    begin_run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    begin_run._r.append(begin)

    instr_run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    instr_run._r.append(instr)

    sep_run = paragraph.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    sep_run._r.append(sep)

    if placeholder:
        paragraph.add_run(placeholder)

    end_run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)


def add_hidden_tc(paragraph, text: str, level: int = 1) -> None:
    """Append a hidden TC field without changing the visible review-page heading."""
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), f' TC "{text}" \\f C \\l {level} ')
    field.set(qn("w:dirty"), "true")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    vanish = OxmlElement("w:vanish")
    rpr.append(vanish)
    run.append(rpr)
    field.append(run)
    paragraph._p.append(field)


def ensure_custom_style(doc: Document, name: str, base: str):
    if name in [style.name for style in doc.styles]:
        return doc.styles[name]
    style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = doc.styles[base]
    return style


def replace_manual_block_with_field(doc: Document, title: str, next_title: str, instruction: str) -> None:
    title_p = find_exact(doc, title)
    next_p = find_exact(doc, next_title)
    body = doc._body._body
    children = list(body)
    start = children.index(title_p._p)
    end = children.index(next_p._p)
    for element in list(children[start + 1 : end]):
        if element.tag == qn("w:p"):
            body.remove(element)
    field_p = OxmlElement("w:p")
    title_p._p.addnext(field_p)
    from docx.text.paragraph import Paragraph

    paragraph = Paragraph(field_p, title_p._parent)
    paragraph.style = doc.styles["toc 1"] if "toc 1" in [s.name for s in doc.styles] else doc.styles["Normal"]
    paragraph.paragraph_format.space_after = Pt(0)
    add_field(paragraph, instruction)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def set_no_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_cell_text(cell, text: str, *, bold: bool, alignment) -> None:
    cell.text = text
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in cell.paragraphs:
        p.alignment = alignment
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        for run in p.runs:
            run.bold = bold
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")


def replace_abbreviation_table(doc: Document) -> None:
    old = doc.tables[0]
    values = [(row.cells[0].text.strip(), row.cells[1].text.strip()) for row in old.rows[1:]]
    table = doc.add_table(rows=len(values) + 1, cols=3)
    if old.style:
        table.style = old.style
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (Cm(1.5), Cm(3.6), Cm(10.2))
    headers = ("STT", "Từ viết tắt", "Ghi chú")
    for j, header in enumerate(headers):
        set_cell_text(table.cell(0, j), header, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        table.cell(0, j).width = widths[j]
    set_repeat_table_header(table.rows[0])
    for i, (abbr, meaning) in enumerate(values, 1):
        set_cell_text(table.cell(i, 0), str(i), bold=False, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(table.cell(i, 1), abbr, bold=False, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(table.cell(i, 2), meaning, bold=False, alignment=WD_ALIGN_PARAGRAPH.LEFT)
        for j, width in enumerate(widths):
            table.cell(i, j).width = width
    for row in table.rows:
        set_no_split(row)
    old._element.addprevious(table._element)
    old._element.getparent().remove(old._element)


def insert_conclusion(doc: Document) -> None:
    refs = find_exact(doc, "TÀI LIỆU THAM KHẢO")
    heading = refs.insert_paragraph_before("KẾT LUẬN", style="Heading 1")
    heading.paragraph_format.page_break_before = True
    heading.paragraph_format.keep_with_next = True
    heading.paragraph_format.keep_together = True
    paragraphs = [
        "Khóa luận đã xây dựng và đánh giá quy trình phát hiện người và phương tiện trong điều kiện thời tiết bất lợi thông qua benchmark bốn kiến trúc YOLOv8n, YOLO11n, Faster R-CNN và RT-DETR-L. Trên cơ sở kết quả XWOD validation, RT-DETR-L được lựa chọn để tiếp tục huấn luyện ở Phase 2.",
        "Phase 2 thực hiện fine-tuning đa miền với dữ liệu XWOD, ACDC và phần dữ liệu BDD100K dùng để phát lại. Kết quả tổng hợp cho thấy RT-DETR-L Phase 2 phục hồi đáng kể hiệu năng trên BDD, cải thiện trên ACDC và duy trì tương đối ổn định trên XWOD, dù kết quả trên DAWN giảm nhẹ.",
        "RT-DETR-L Phase 2 vì vậy được xác định là mô hình cuối cùng trong phạm vi khóa luận. Tuy nhiên, hiệu năng trong điều kiện ban đêm và trên một số lớp khó vẫn còn hạn chế, cho thấy mô hình chưa thích nghi đồng đều với mọi miền dữ liệu và điều kiện quan sát.",
        "Các hướng tiếp tục cải thiện dữ liệu lớp hiếm, lựa chọn mẫu theo failure case, tăng cường điều kiện ban đêm, giảm quên miền và tối ưu triển khai đã được trình bày tại mục 4.6.3. Đây là cơ sở để mở rộng nghiên cứu trong các vòng thực nghiệm tiếp theo.",
    ]
    for text in paragraphs:
        p = refs.insert_paragraph_before(text, style="Normal")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(1.27)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(0)
        for run in p.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(13)
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")


def normalize_body(doc: Document) -> None:
    start = paragraph_index(doc, find_exact(doc, "CHƯƠNG 1: TỔNG QUAN VỀ ĐỀ TÀI"))
    end = paragraph_index(doc, find_exact(doc, "TÀI LIỆU THAM KHẢO"))
    for p in doc.paragraphs[start:end]:
        if p.style.name not in ("Normal", "Normal (Web)") or not p.text.strip():
            continue
        if p._p.xpath(".//m:oMath | .//m:oMathPara | .//w:drawing | .//w:pict | .//w:numPr"):
            continue
        if p.alignment in (WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT):
            continue
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(1.27)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        for run in p.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(13)
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")


def normalize_headings_and_captions(doc: Document) -> None:
    style_names = [s.name for s in doc.styles]
    if "Heading 3" not in style_names:
        old = doc.styles["Heading 33"]
        new_element = deepcopy(old._element)
        new_element.set(qn("w:styleId"), "Heading3")
        new_element.attrib.pop(qn("w:customStyle"), None)
        new_element.find(qn("w:name")).set(qn("w:val"), "heading 3")
        if new_element.find(qn("w:qFormat")) is None:
            new_element.insert(3, OxmlElement("w:qFormat"))
        if new_element.find(qn("w:unhideWhenUsed")) is None:
            new_element.insert(3, OxmlElement("w:unhideWhenUsed"))
        old._element.addprevious(new_element)
    for p in doc.paragraphs:
        if p.style.name == "Heading 33":
            p.style = doc.styles["heading 3"]
        if p.style.name.startswith("Heading"):
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.keep_together = True
    if "Heading 33" in [s.name for s in doc.styles]:
        stale = doc.styles["Heading 33"]._element
        stale.getparent().remove(stale)

    table_caption = ensure_custom_style(doc, "Table Caption", "Caption4")
    figure_caption = ensure_custom_style(doc, "Figure Caption", "Caption4")
    for p in doc.paragraphs:
        text = p.text.strip()
        if p.style.name.startswith("Caption") or p.style.name in ("Table Caption", "Figure Caption"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_together = True
            if text.startswith("Bảng "):
                p.style = table_caption
                p.paragraph_format.keep_with_next = True
            elif text.startswith("Hình "):
                p.style = figure_caption
                p.paragraph_format.keep_with_next = False
                previous = p._p.getprevious()
                if previous is not None and previous.tag == qn("w:p"):
                    ppr = previous.get_or_add_pPr()
                    keep_next = ppr.find(qn("w:keepNext"))
                    if keep_next is None:
                        ppr.append(OxmlElement("w:keepNext"))


def audit_tables(doc: Document) -> None:
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        if table.rows:
            set_repeat_table_header(table.rows[0])
        for row in table.rows:
            set_no_split(row)
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def configure_automatic_lists(doc: Document) -> None:
    front_style = ensure_custom_style(doc, "Front Matter Heading", "Normal")
    for text in (
        "LỜI MỞ ĐẦU",
        "LỜI CẢM ƠN",
        "MỤC LỤC",
        "DANH MỤC BẢNG BIỂU",
        "DANH MỤC HÌNH",
        "DANH MỤC CÁC TỪ VIẾT TẮT",
    ):
        find_exact(doc, text).style = front_style
    add_hidden_tc(find_exact(doc, "NHẬN XÉT CỦA GIẢNG VIÊN HƯỚNG DẪN"), "NHẬN XÉT CỦA GIẢNG VIÊN HƯỚNG DẪN")
    add_hidden_tc(find_exact(doc, "NHẬN XÉT CỦA GIẢNG VIÊN PHẢN BIỆN"), "NHẬN XÉT CỦA GIẢNG VIÊN PHẢN BIỆN")

    replace_manual_block_with_field(
        doc,
        "MỤC LỤC",
        "DANH MỤC BẢNG BIỂU",
        ' TOC \\o "1-3" \\h \\z \\f C \\t "Front Matter Heading,1" ',
    )
    replace_manual_block_with_field(
        doc,
        "DANH MỤC BẢNG BIỂU",
        "DANH MỤC HÌNH",
        ' TOC \\h \\z \\t "Table Caption,1" ',
    )
    replace_manual_block_with_field(
        doc,
        "DANH MỤC HÌNH",
        "DANH MỤC CÁC TỪ VIẾT TẮT",
        ' TOC \\h \\z \\t "Figure Caption,1" ',
    )


def enable_field_updates(doc: Document) -> None:
    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def revise(source: Path, output: Path) -> None:
    doc = Document(source)
    normalize_headings_and_captions(doc)
    insert_conclusion(doc)
    normalize_body(doc)
    replace_abbreviation_table(doc)
    audit_tables(doc)
    configure_automatic_lists(doc)
    enable_field_updates(doc)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    revise(args.source, args.output)


if __name__ == "__main__":
    main()
