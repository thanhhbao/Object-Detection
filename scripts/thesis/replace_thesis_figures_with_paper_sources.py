#!/usr/bin/env python3
"""Replace theory diagrams with verified figures from the cited original papers."""

from __future__ import annotations

import argparse
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt
from docx.shape import InlineShape
from docx.text.paragraph import Paragraph
from PIL import Image


FIGURES = {
    "Hình 2.2. Quy trình tổng quát của mô hình một giai đoạn": {
        "asset": "yolo_fig1.png",
        "source": 'Nguồn: Redmon et al., "You Only Look Once: Unified, Real-Time Object Detection", CVPR 2016, Hình 1.',
        "width_cm": 15.5,
    },
    "Hình 2.3. Quy trình tổng quát của mô hình hai giai đoạn": {
        "asset": "faster_rcnn_fig2.png",
        "source": 'Nguồn: Ren et al., "Faster R-CNN: Towards Real-Time Object Detection with Region Proposal Networks", NeurIPS 2015, Hình 2.',
        "width_cm": 11.5,
    },
    "Hình 2.4. Quy trình tổng quát của mô hình dựa trên Transformer": {
        "asset": "detr_fig1.png",
        "source": 'Nguồn: Carion et al., "End-to-End Object Detection with Transformers", ECCV 2020, Hình 1.',
        "width_cm": 15.5,
    },
    "Hình 2.5. Cấu trúc mô-đun chú ý Squeeze-and-Excitation (SE)": {
        "asset": "senet_fig1.png",
        "source": 'Nguồn: Hu et al., "Squeeze-and-Excitation Networks", CVPR 2018, Hình 1.',
        "width_cm": 15.5,
    },
}


def paragraph_after(paragraph: Paragraph) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def preceding_drawing(paragraphs: list[Paragraph], caption_index: int) -> Paragraph:
    for idx in range(caption_index - 1, max(-1, caption_index - 5), -1):
        if paragraphs[idx]._p.xpath(".//w:drawing"):
            return paragraphs[idx]
        if paragraphs[idx].text.strip():
            break
    raise RuntimeError(f"No drawing found before caption: {paragraphs[caption_index].text}")


def replace_image_part(doc: Document, drawing_paragraph: Paragraph, asset: Path, width_cm: float) -> None:
    rids = drawing_paragraph._p.xpath(".//@r:embed")
    if len(rids) != 1:
        raise RuntimeError(f"Expected one embedded image, found {rids}")
    image_part = doc.part.related_parts[rids[0]]
    image_part._blob = asset.read_bytes()

    with Image.open(asset) as image:
        ratio = image.height / image.width
    width = Cm(width_cm)
    height = int(width * ratio)
    inline = drawing_paragraph._p.xpath(".//wp:inline")
    if len(inline) != 1:
        raise RuntimeError("Expected one inline drawing")
    shape = InlineShape(inline[0])
    shape.width = width
    shape.height = height
    drawing_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    drawing_paragraph.paragraph_format.keep_with_next = True
    drawing_paragraph.paragraph_format.keep_together = True


def set_source(paragraph: Paragraph, text: str, source_style) -> None:
    paragraph.text = text
    paragraph.style = source_style
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_together = True
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)
        run.font.italic = True


def add_missing_figure_list_entries(doc: Document) -> None:
    paragraphs = doc.paragraphs
    list_title = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "DANH MỤC HÌNH")
    list_end = next(
        i for i, p in enumerate(paragraphs[list_title + 1 :], list_title + 1)
        if p.text.strip().startswith("DANH MỤC ")
    )
    first = next(p for p in paragraphs[list_title + 1 :] if p.text.strip().startswith("Hình 2.1."))
    existing = {p.text.split("\t", 1)[0].strip() for p in paragraphs[list_title + 1 : list_end]}
    entries = [
        "Hình 2.2. Quy trình tổng quát của mô hình một giai đoạn",
        "Hình 2.3. Quy trình tổng quát của mô hình hai giai đoạn",
        "Hình 2.4. Quy trình tổng quát của mô hình dựa trên Transformer",
    ]
    current = first
    style = doc.styles["toc 1"]
    for text in entries:
        if text in existing:
            continue
        new_p = paragraph_after(current)
        new_p.text = f"{text}\t?"
        new_p.style = style
        new_p.paragraph_format.tab_stops.clear_all()
        new_p.paragraph_format.tab_stops.add_tab_stop(Cm(15.45), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        current = new_p


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("assets", type=Path)
    parser.add_argument("--backup", type=Path)
    args = parser.parse_args()

    if args.backup:
        args.backup.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(args.source, args.backup)

    doc = Document(args.source)
    paragraphs = doc.paragraphs
    source_style = doc.styles["Figure Source"]
    for caption_text, spec in FIGURES.items():
        caption_index = next(i for i, p in enumerate(paragraphs) if p.text.strip() == caption_text)
        caption = paragraphs[caption_index]
        drawing = preceding_drawing(paragraphs, caption_index)
        replace_image_part(doc, drawing, args.assets / spec["asset"], spec["width_cm"])

        next_nonempty = next(
            (paragraphs[i] for i in range(caption_index + 1, min(len(paragraphs), caption_index + 4)) if paragraphs[i].text.strip()),
            None,
        )
        if next_nonempty is not None and next_nonempty.text.strip().startswith("Nguồn:"):
            set_source(next_nonempty, spec["source"], source_style)
        else:
            set_source(paragraph_after(caption), spec["source"], source_style)

    add_missing_figure_list_entries(doc)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
