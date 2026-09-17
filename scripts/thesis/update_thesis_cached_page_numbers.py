#!/usr/bin/env python3
"""Update cached TOC, list-of-figures and list-of-tables page numbers from PDF."""

from __future__ import annotations

import argparse
import re
import subprocess
import tempfile
import unicodedata
from pathlib import Path

from docx import Document
from docx.shared import Pt


FRONT_TITLES = [
    "LỜI MỞ ĐẦU",
    "LỜI CẢM ƠN",
    "NHẬN XÉT CỦA GIẢNG VIÊN HƯỚNG DẪN",
    "NHẬN XÉT CỦA GIẢNG VIÊN PHẢN BIỆN",
    "MỤC LỤC",
    "DANH MỤC BẢNG BIỂU",
    "DANH MỤC HÌNH",
    "DANH MỤC CÁC TỪ VIẾT TẮT",
]


def norm(text: str) -> str:
    text = unicodedata.normalize("NFC", text).replace("–", "-").replace("−", "-")
    return re.sub(r"\s+", " ", text).strip().casefold()


def roman(value: int) -> str:
    pairs = ((1000, "m"), (900, "cm"), (500, "d"), (400, "cd"), (100, "c"), (90, "xc"), (50, "l"), (40, "xl"), (10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i"))
    result = []
    for number, token in pairs:
        while value >= number:
            result.append(token)
            value -= number
    return "".join(result)


def read_pdf(pdf: Path) -> tuple[list[str], list[list[str]]]:
    with tempfile.NamedTemporaryFile(suffix=".txt") as tmp:
        subprocess.run(["pdftotext", "-layout", str(pdf), tmp.name], check=True)
        raw_pages = Path(tmp.name).read_text(errors="ignore").split("\f")
    collapsed = [norm(page) for page in raw_pages]
    lines = [[norm(line) for line in page.splitlines() if norm(line)] for page in raw_pages]
    return collapsed, lines


def find_page(pages: list[str], text: str, start: int = 1) -> int:
    needle = norm(text)
    if needle == "phụ lục":
        needle = "phụ lục a."
    for idx, page in enumerate(pages[start - 1 :], start):
        if needle in page:
            return idx
    prefix = " ".join(needle.split()[:8])
    for idx, page in enumerate(pages[start - 1 :], start):
        if prefix in page:
            return idx
    raise RuntimeError(f"Could not find rendered page for: {text}")


def find_exact_line(lines: list[list[str]], text: str) -> int:
    needle = norm(text)
    for idx, page_lines in enumerate(lines, 1):
        if needle in page_lines:
            return idx
    raise RuntimeError(f"Could not find exact front-matter title: {text}")


def set_entry(paragraph, title: str, page: str) -> None:
    paragraph.text = f"{title}\t{page}"
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(13)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("docx", type=Path)
    parser.add_argument("pdf", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    pages, lines = read_pdf(args.pdf)
    doc = Document(args.docx)
    body_heading = next(
        p.text.strip() for p in doc.paragraphs
        if p.style.name == "Heading 1" and p.text.strip().startswith("CHƯƠNG 1:")
    )
    body_physical = find_exact_line(lines, body_heading)
    prelim_physical = find_exact_line(lines, "LỜI MỞ ĐẦU")

    heading_pages: dict[str, int] = {}
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if paragraph.style.name in {"Heading 1", "Heading 2", "Heading 3", "Heading 33"} and text:
            heading_pages[text] = find_page(pages, text, body_physical) - body_physical + 1
    for extra in ("TÀI LIỆU THAM KHẢO", "PHỤ LỤC"):
        heading_pages[extra] = find_page(pages, extra, body_physical) - body_physical + 1

    front_pages = {
        title: roman(find_exact_line(lines, title) - prelim_physical + 1)
        for title in FRONT_TITLES
    }

    in_figure_list = False
    in_table_list = False
    for paragraph in doc.paragraphs:
        raw = paragraph.text.strip()
        title = raw.split("\t", 1)[0].strip()
        if raw == "DANH MỤC HÌNH":
            in_figure_list, in_table_list = True, False
            continue
        if raw == "DANH MỤC CÁC TỪ VIẾT TẮT":
            in_figure_list = in_table_list = False
        if raw == "DANH MỤC BẢNG BIỂU":
            in_figure_list, in_table_list = False, True
            continue
        if raw == "DANH MỤC HÌNH":
            in_figure_list, in_table_list = True, False

        if in_figure_list and title.startswith("Hình "):
            page = find_page(pages, title, body_physical) - body_physical + 1
            set_entry(paragraph, title, str(page))
            continue
        if in_table_list and title.startswith("Bảng "):
            page = find_page(pages, title, body_physical) - body_physical + 1
            set_entry(paragraph, title, str(page))
            continue

        if not paragraph.style.name.startswith("toc"):
            continue
        if title in front_pages:
            set_entry(paragraph, title, front_pages[title])
        elif title in heading_pages:
            set_entry(paragraph, title, str(heading_pages[title]))

    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    print(f"body physical page={body_physical}; headings={len(heading_pages)}; front={front_pages}")


if __name__ == "__main__":
    main()
